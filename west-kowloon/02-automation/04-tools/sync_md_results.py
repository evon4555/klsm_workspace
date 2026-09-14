# -*- coding: utf-8 -*-
"""
sync_md_results.py — mirror the execution-result columns (M-S) from
test-cases-registration-login.xlsx back into the .md source, so the two stay
consistent after tools/update_evidence.py has run.

xlsx-driven and re-runnable: every test-case row that has a Status in the
workbook gets its M-S columns copied into the matching .md table row. The
.md Screenshots column gets the evidence-strip path (the xlsx itself carries
the embedded image).
"""
from pathlib import Path
import os

import openpyxl

WESTK_ROOT = Path(os.getenv("QA_WESTK_ROOT", Path(__file__).resolve().parents[2])).resolve()
PKG = (
    WESTK_ROOT
    / "01-requirements"
    / "02-subprojects"
    / "02-website"
    / "02-modules"
    / "01-login-registration"
    / "2026-06-16"
)
XLSX = PKG / "03-test-design" / "test-cases-registration-login.xlsx"
MD = PKG / "03-test-design" / "test-cases-registration-login.md"
EVID_DATE = "2026-05-23"


def main():
    ws = openpyxl.load_workbook(XLSX)["Test Cases"]

    results = {}
    for r in range(2, ws.max_row + 1):
        cid = str(ws.cell(row=r, column=1).value or "")
        status = ws.cell(row=r, column=17).value
        if not cid or not status:
            continue

        def cell(col):
            v = ws.cell(row=r, column=col).value
            return "" if v is None else \
                str(v).replace("\n", " ").replace("|", "/").strip()

        tag = cid.split("-")[-1]
        shot = (f"05-execution/evidence/registration-TC{tag}-"
                f"{str(status).lower()}-{EVID_DATE}.png")
        # M N O P Q R  +  S (screenshot path)
        results[cid] = [cell(c) for c in range(13, 19)] + [shot]

    lines = MD.read_text(encoding="utf-8").splitlines(keepends=True)
    updated = 0
    for i, line in enumerate(lines):
        for cid, ms in results.items():
            if not line.startswith(f"| {cid} "):
                continue
            nl = "\n" if line.endswith("\n") else ""
            cells = line.rstrip("\n").split("|")
            if len(cells) < 21:                  # 19 columns -> 21 split parts
                break
            for j, value in enumerate(ms):       # cells[13..19] == columns M..S
                cells[13 + j] = f" {value} "
            lines[i] = "|".join(cells) + nl
            updated += 1
            break

    MD.write_text("".join(lines), encoding="utf-8")
    print(f"synced {updated} rows from {XLSX.name} into {MD.name}")


if __name__ == "__main__":
    main()
