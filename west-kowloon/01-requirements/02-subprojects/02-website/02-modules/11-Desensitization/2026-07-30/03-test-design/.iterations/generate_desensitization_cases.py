from __future__ import annotations

import math
import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


PACKAGE = Path(r"D:\Workspace\west-kowloon\01-requirements\02-subprojects\02-website\02-modules\11-Desensitization\2026-07-30")
TEMPLATE = Path(r"D:\Workspace\west-kowloon\01-requirements\01-source-documents\02-templates\TestCase_Template.xlsx")
MD_OUT = PACKAGE / "03-test-design" / "test-cases-desensitization.md"
XLSX_OUT = PACKAGE / "03-test-design" / "test-cases-desensitization.xlsx"

HEADERS = [
    "Label\n(SIT/UAT/PROD/NA)",
    "Test Case ID",
    "Module/Feature",
    "Priority\n(High/Medium/Low)",
    "Severity\n(High/Medium/Low)",
    "Collected from",
    "Test Scenario",
    "Test Case Description",
    "Preconditions\n(if any)",
    "Test Steps",
    "Test Data\n(if any)",
    "Expected Result",
    "Test Case Owner",
    "Environment\n(SIT/UAT/PROD)",
    "Execution Date",
    "Executed By",
    "Actual Result",
    "Status\n(Pass/ Fail)",
    "Comments/Remarks",
    "Screenshots",
]

SOURCE_RULES = "ZenTao 4319；checking rules #1-7；签字范围 F1-F20"
FULL_PROFILE = (
    "手机号 +8613800008000；邮箱 user@example.com；姓名 张建国；"
    "身份证 320583198601100090；地址 河南省南阳市新野县朝阳路100号；"
    "生日 1996-08-23；性别 男"
)

ALL_PII_EXPECTED = [
    "手机号显示为 +861*****000。",
    "邮箱显示为 us****@example.com。",
    "姓名显示为 张*国。",
    "身份证显示为 320583********0090。",
    "地址显示为 河南省南阳市新野县******。",
    "生日显示为 08-23。",
    "性别显示为 *。",
    "页面不显示上述字段的完整原值。",
]

cases: list[dict[str, str]] = []


def add(
    module: str,
    priority: str,
    severity: str,
    collected_from: str,
    scenario: str,
    description: str,
    preconditions: str,
    steps: list[str],
    data: str,
    expected: list[str],
) -> None:
    case_id = f"SIT-TC-WK-PII-{len(cases) + 1:03d}"
    cases.append(
        {
            "Label": "SIT",
            "Test Case ID": case_id,
            "Module/Feature": module,
            "Priority": priority,
            "Severity": severity,
            "Collected from": collected_from,
            "Test Scenario": scenario,
            "Test Case Description": description,
            "Preconditions": preconditions,
            "Test Steps": "\n".join(f"{i}. {step}" for i, step in enumerate(steps, 1)),
            "Test Data": data,
            "Expected Result": "\n".join(f"{i}. {item}" for i, item in enumerate(expected, 1)),
            "Test Case Owner": "Antank QA Team",
            "Environment": "",
            "Execution Date": "",
            "Executed By": "",
            "Actual Result": "",
            "Status": "",
            "Comments/Remarks": "",
            "Screenshots": "",
        }
    )


# Website / Personal Center: rule-level coverage for all seven fields.
pc_pre = "SIT 已部署本次脱敏功能；测试账号可登录 Website；个人资料已保存指定测试数据。"
pc_steps = ["使用指定测试账号登录 Website。", "打开个人中心。", "打开个人资料页面。"]

add(
    "Website / 个人中心 / 登录账号",
    "High",
    "High",
    "ZenTao 4319 前端场景；checking rules #1",
    "中国内地手机号脱敏",
    "验证个人中心登录账号区域按需求示例脱敏显示中国内地手机号。",
    pc_pre,
    ["使用指定测试账号登录 Website。", "打开个人中心。"],
    "登录手机号：+8613800008000",
    ["登录手机号显示为 +861*****000。", "页面不显示完整手机号 +8613800008000。"],
)
add(
    "Website / 个人中心 / 登录账号",
    "High",
    "High",
    "ZenTao 4319 前端场景；checking rules #1",
    "国际手机号脱敏",
    "验证个人中心登录账号区域按需求示例脱敏显示香港手机号。",
    pc_pre,
    ["使用香港手机号测试账号登录 Website。", "打开个人中心。"],
    "登录手机号：+852 9123-4567",
    ["登录手机号显示为 +85****567。", "页面不显示完整手机号 +852 9123-4567。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "High",
    "High",
    "ZenTao 4319 前端场景；checking rules #2",
    "邮箱脱敏",
    "验证个人资料页面按规则脱敏显示邮箱地址。",
    pc_pre,
    pc_steps,
    "邮箱：user@example.com",
    ["邮箱显示为 us****@example.com。", "邮箱域名 example.com 保持完整。", "页面不显示完整邮箱 user@example.com。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #3",
    "三字中文姓名脱敏",
    "验证个人资料页面按规则脱敏显示长度大于二的中文姓名。",
    pc_pre,
    pc_steps,
    "姓名：张建国",
    ["姓名显示为 张*国。", "页面不显示完整姓名 张建国。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #3",
    "两字中文姓名脱敏",
    "验证个人资料页面按规则脱敏显示长度小于等于二的中文姓名。",
    pc_pre,
    pc_steps,
    "姓名：张荣",
    ["姓名显示为 张*。", "页面不显示完整姓名 张荣。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #3",
    "英文姓名脱敏",
    "验证个人资料页面按需求示例脱敏显示英文姓名。",
    pc_pre,
    pc_steps,
    "姓名：Chuchu Lin",
    ["姓名显示为 C********N。", "页面不显示完整姓名 Chuchu Lin。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #4",
    "八位证件号码脱敏",
    "验证个人资料页面按规则脱敏显示长度小于等于八的证件号码。",
    pc_pre,
    pc_steps,
    "证件号码：12345678",
    ["证件号码显示为 12****78。", "页面不显示完整证件号码 12345678。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #4",
    "一般长证件号码脱敏",
    "验证个人资料页面按需求示例脱敏显示带括号的长证件号码。",
    pc_pre,
    pc_steps,
    "证件号码：X123456(7)",
    ["证件号码显示为 X123****6(7)。", "页面不显示完整证件号码 X123456(7)。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "High",
    "High",
    "ZenTao 4319 前端场景；checking rules #4",
    "十八位身份证脱敏",
    "验证个人资料页面按规则脱敏显示十八位身份证号码。",
    pc_pre,
    pc_steps,
    "身份证：320583198601100090",
    ["身份证显示为 320583********0090。", "页面不显示完整身份证 320583198601100090。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #5",
    "地址脱敏",
    "验证个人资料页面保留省市区并脱敏显示详细地址。",
    pc_pre,
    pc_steps,
    "地址：河南省南阳市新野县朝阳路100号",
    ["地址显示为 河南省南阳市新野县******。", "页面不显示朝阳路100号。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #6",
    "完整生日脱敏",
    "验证个人资料页面隐藏年份并保留生日月日。",
    pc_pre,
    pc_steps,
    "生日：1996-08-23",
    ["生日显示为 08-23。", "页面不显示年份 1996。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Low",
    "Medium",
    "ZenTao 4319 前端场景；checking rules #6",
    "年月生日脱敏",
    "验证个人资料页面对仅含年月的生日隐藏年份并保留月份。",
    pc_pre,
    pc_steps,
    "生日：1996-08",
    ["生日显示为 08。", "页面不显示年份 1996。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Medium",
    "High",
    "ZenTao 4319 前端场景；checking rules #7",
    "性别脱敏",
    "验证个人资料页面使用星号脱敏显示性别。",
    pc_pre,
    pc_steps,
    "性别：男",
    ["性别显示为 *。", "页面不显示原性别 男。"],
)
add(
    "Website / 个人中心 / 个人资料",
    "Low",
    "Medium",
    "签字范围 U-2",
    "空值保持为空",
    "验证个人资料中的空白个人信息字段保持为空。",
    "SIT 已部署本次脱敏功能；测试账号的地址、生日和性别为空。",
    pc_steps,
    "地址：空；生日：空；性别：空",
    ["地址、生日和性别字段保持为空。", "空字段不显示星号占位。", "页面正常加载。"],
)

# Website / Membership Card.
add(
    "Website / 会员卡 / 激活信息",
    "High",
    "High",
    "ZenTao 4319 会员卡码和激活信息展示；checking rules #1-7",
    "会员卡激活信息脱敏",
    "验证 Website 会员卡激活信息按各字段规则脱敏显示个人信息。",
    "SIT 已存在完成激活的会员卡；激活人资料包含指定 PII。",
    ["登录 Website。", "打开个人中心。", "打开会员卡页面。", "选择已激活的会员卡。", "打开激活信息。"],
    FULL_PROFILE,
    ALL_PII_EXPECTED,
)
add(
    "Website / 会员卡 / 详情",
    "High",
    "High",
    "ZenTao 4319 会员卡列表及详情；checking rules #1-7",
    "会员卡详情个人信息脱敏",
    "验证 Website 会员卡详情按各字段规则脱敏显示持卡人信息。",
    "SIT 已存在绑定指定测试用户的有效会员卡。",
    ["登录 Website。", "打开个人中心。", "打开会员卡页面。", "选择指定会员卡。"],
    FULL_PROFILE,
    ALL_PII_EXPECTED,
)
add(
    "Website / 会员卡 / 详情",
    "Medium",
    "Medium",
    "签字决策 Q-8",
    "会员卡号保持原值",
    "验证 Website 会员卡详情在脱敏个人信息时保持会员卡号完整显示。",
    "SIT 已存在绑定指定测试用户的有效会员卡。",
    ["登录 Website。", "打开个人中心。", "打开会员卡页面。", "选择指定会员卡。"],
    "会员卡号：WKMC20260730001；手机号：+8613800008000；邮箱：user@example.com",
    ["会员卡号显示为 WKMC20260730001。", "手机号显示为 +861*****000。", "邮箱显示为 us****@example.com。"],
)


def add_business_surface(
    module: str,
    collected: str,
    scenario: str,
    description: str,
    preconditions: str,
    steps: list[str],
) -> None:
    add(module, "High", "High", collected, scenario, description, preconditions, steps, FULL_PROFILE, ALL_PII_EXPECTED)


# Website / Orders, explicitly split by ticket type.
add_business_surface(
    "Website / 门票订单 / 订单详情",
    "ZenTao 4319 实名制订单展示；checking rules #1-7",
    "门票订单详情脱敏",
    "验证 Website 门票订单详情按各字段规则脱敏显示实名用户信息。",
    "SIT 已存在包含指定 PII 的已完成实名制门票订单。",
    ["登录 Website。", "打开订单列表。", "选择指定门票订单。", "打开订单详情。"],
)
add_business_surface(
    "Website / 门票 / 票夹",
    "ZenTao 4319 票夹展示；checking rules #1-7",
    "门票票夹脱敏",
    "验证 Website 门票票夹按各字段规则脱敏显示持票人信息。",
    "SIT 已存在包含指定 PII 的有效门票。",
    ["登录 Website。", "打开票夹。", "选择指定门票。"],
)
add_business_surface(
    "Website / 门票 / 入场码",
    "ZenTao 4319 入场码展示；checking rules #1-7",
    "门票入场码页面脱敏",
    "验证 Website 门票入场码页面按各字段规则脱敏显示持票人信息。",
    "SIT 已存在包含指定 PII 且可展示入场码的有效门票。",
    ["登录 Website。", "打开票夹。", "选择指定门票。", "打开入场码。"],
)
add_business_surface(
    "Website / 座票订单 / 订单详情",
    "ZenTao 4319 实名制订单展示；checking rules #1-7",
    "座票订单详情脱敏",
    "验证 Website 座票订单详情按各字段规则脱敏显示实名用户信息。",
    "SIT 已存在包含指定 PII 的已完成实名制座票订单。",
    ["登录 Website。", "打开订单列表。", "选择指定座票订单。", "打开订单详情。"],
)
add_business_surface(
    "Website / 座票 / 票夹",
    "ZenTao 4319 票夹展示；checking rules #1-7",
    "座票票夹脱敏",
    "验证 Website 座票票夹按各字段规则脱敏显示持票人信息。",
    "SIT 已存在包含指定 PII 的有效座票。",
    ["登录 Website。", "打开票夹。", "选择指定座票。"],
)
add_business_surface(
    "Website / 座票 / 入场码",
    "ZenTao 4319 入场码展示；checking rules #1-7",
    "座票入场码页面脱敏",
    "验证 Website 座票入场码页面按各字段规则脱敏显示持票人信息。",
    "SIT 已存在包含指定 PII 且可展示入场码的有效座票。",
    ["登录 Website。", "打开票夹。", "选择指定座票。", "打开入场码。"],
)

# PDA, split by ticket type.
add_business_surface(
    "PDA / 门票扫码 / 用户信息",
    "ZenTao 4319 PDA 扫码场景；checking rules #1-7",
    "PDA 门票扫码用户信息脱敏",
    "验证 PDA 扫描门票后按各字段规则脱敏显示用户信息。",
    "SIT PDA 可用；存在包含指定 PII 的有效门票二维码。",
    ["登录 SIT PDA。", "进入验票页面。", "扫描指定门票二维码。", "打开用户信息区域。"],
)
add_business_surface(
    "PDA / 座票扫码 / 用户信息",
    "ZenTao 4319 PDA 扫码场景；checking rules #1-7",
    "PDA 座票扫码用户信息脱敏",
    "验证 PDA 扫描座票后按各字段规则脱敏显示用户信息。",
    "SIT PDA 可用；存在包含指定 PII 的有效座票二维码。",
    ["登录 SIT PDA。", "进入验票页面。", "扫描指定座票二维码。", "打开用户信息区域。"],
)

# Backend / Member 360 and orders.
add_business_surface(
    "后台 / 会员360 / 列表与搜索结果",
    "ZenTao 4319 后台会员360；checking rules #1-7",
    "会员360列表脱敏",
    "验证后台会员360列表和搜索结果按各字段规则脱敏显示会员信息。",
    "SIT 后台存在指定测试会员；测试账号可访问会员360。",
    ["登录 SIT 后台。", "打开会员360。", "使用会员编号搜索指定会员。"],
)
add_business_surface(
    "后台 / 会员360 / 详情",
    "ZenTao 4319 后台会员360；checking rules #1-7",
    "会员360详情脱敏",
    "验证后台会员360详情按各字段规则脱敏显示会员信息。",
    "SIT 后台存在指定测试会员；测试账号可访问会员360。",
    ["登录 SIT 后台。", "打开会员360。", "搜索指定会员。", "打开会员详情。"],
)
add_business_surface(
    "后台 / 门票订单 / 列表",
    "ZenTao 4319 后台门票订单列表；checking rules #1-7",
    "后台门票订单列表脱敏",
    "验证后台门票订单列表按各字段规则脱敏显示购票人信息。",
    "SIT 后台存在包含指定 PII 的门票订单。",
    ["登录 SIT 后台。", "打开门票订单列表。", "使用订单编号搜索指定订单。"],
)
add_business_surface(
    "后台 / 门票订单 / 详情",
    "ZenTao 4319 后台门票订单详情；checking rules #1-7",
    "后台门票订单详情脱敏",
    "验证后台门票订单详情按各字段规则脱敏显示购票人信息。",
    "SIT 后台存在包含指定 PII 的门票订单。",
    ["登录 SIT 后台。", "打开门票订单列表。", "搜索指定订单。", "打开订单详情。"],
)
add_business_surface(
    "后台 / 座票订单 / 列表",
    "ZenTao 4319 后台选座订单列表；checking rules #1-7",
    "后台座票订单列表脱敏",
    "验证后台座票订单列表按各字段规则脱敏显示购票人信息。",
    "SIT 后台存在包含指定 PII 的座票订单。",
    ["登录 SIT 后台。", "打开座票订单列表。", "使用订单编号搜索指定订单。"],
)
add_business_surface(
    "后台 / 座票订单 / 详情",
    "ZenTao 4319 后台选座订单详情；checking rules #1-7",
    "后台座票订单详情脱敏",
    "验证后台座票订单详情按各字段规则脱敏显示购票人信息。",
    "SIT 后台存在包含指定 PII 的座票订单。",
    ["登录 SIT 后台。", "打开座票订单列表。", "搜索指定订单。", "打开订单详情。"],
)
add_business_surface(
    "后台 / 会员卡订单 / 列表",
    "ZenTao 4319 后台会员卡订单；checking rules #1-7",
    "会员卡订单列表脱敏",
    "验证后台会员卡订单列表按各字段规则脱敏显示购买人信息。",
    "SIT 后台存在包含指定 PII 的会员卡订单。",
    ["登录 SIT 后台。", "打开会员卡订单列表。", "使用订单编号搜索指定订单。"],
)
add_business_surface(
    "后台 / 会员卡订单 / 详情",
    "ZenTao 4319 后台会员卡订单；checking rules #1-7",
    "会员卡订单详情脱敏",
    "验证后台会员卡订单详情按各字段规则脱敏显示购买人信息。",
    "SIT 后台存在包含指定 PII 的会员卡订单。",
    ["登录 SIT 后台。", "打开会员卡订单列表。", "搜索指定订单。", "打开订单详情。"],
)

# Backend / Membership cards, with the signed non-PII card-number decision.
add(
    "后台 / 会员卡 / 列表",
    "High",
    "High",
    "ZenTao 4319 后台会员卡列表；checking rules #1-7；签字决策 Q-8",
    "会员卡列表脱敏",
    "验证后台会员卡列表按照字段类型正确显示会员卡号和持卡人信息。",
    "SIT 后台存在绑定指定测试用户的有效会员卡。",
    ["登录 SIT 后台。", "打开会员卡列表。", "使用会员卡号搜索指定会员卡。"],
    f"会员卡号 WKMC20260730001；{FULL_PROFILE}",
    ["会员卡号显示为 WKMC20260730001。", *ALL_PII_EXPECTED],
)
add(
    "后台 / 会员卡 / 详情",
    "High",
    "High",
    "ZenTao 4319 后台会员卡详情；checking rules #1-7；签字决策 Q-8",
    "会员卡详情脱敏",
    "验证后台会员卡详情按照字段类型正确显示会员卡号和持卡人信息。",
    "SIT 后台存在绑定指定测试用户的有效会员卡。",
    ["登录 SIT 后台。", "打开会员卡列表。", "搜索指定会员卡。", "打开会员卡详情。"],
    f"会员卡号 WKMC20260730001；{FULL_PROFILE}",
    ["会员卡号显示为 WKMC20260730001。", *ALL_PII_EXPECTED],
)

# Backend reports and BI custom reports.
add_business_surface(
    "后台 / 报表 / 页面结果",
    "ZenTao 4319 后台报表；checking rules #1-7；签字决策 Q-9",
    "后台报表页面脱敏",
    "验证后台报表查询结果按各字段规则脱敏显示个人信息。",
    "SIT 后台报表包含指定测试用户数据。",
    ["登录 SIT 后台。", "打开包含个人信息的报表。", "选择包含指定用户的查询条件。", "执行查询。"],
)
add_business_surface(
    "后台 / 报表 / 预览",
    "ZenTao 4319 后台报表；checking rules #1-7；签字决策 Q-9",
    "后台报表预览脱敏",
    "验证后台报表预览按各字段规则脱敏显示个人信息。",
    "SIT 后台报表包含指定测试用户数据。",
    ["登录 SIT 后台。", "打开包含个人信息的报表。", "执行查询。", "打开报表预览。"],
)
add_business_surface(
    "后台 / 报表 / 下载",
    "ZenTao 4319 后台报表；checking rules #1-7；签字决策 Q-9",
    "后台报表下载文件脱敏",
    "验证后台报表下载文件按各字段规则脱敏保存个人信息。",
    "SIT 后台报表包含指定测试用户数据；测试账号可下载报表。",
    ["登录 SIT 后台。", "打开包含个人信息的报表。", "执行查询。", "下载报表。", "打开下载文件。"],
)
add_business_surface(
    "后台 / 报表 / 导出",
    "ZenTao 4319 后台报表；checking rules #1-7；签字决策 Q-9",
    "后台报表导出文件脱敏",
    "验证后台报表导出文件按各字段规则脱敏保存个人信息。",
    "SIT 后台报表包含指定测试用户数据；测试账号可导出报表。",
    ["登录 SIT 后台。", "打开包含个人信息的报表。", "执行查询。", "执行报表导出。", "打开导出文件。"],
)
add_business_surface(
    "BI / 定制报表 / 页面结果",
    "ZenTao 4319 BI 定制报表；checking rules #1-7；签字决策 Q-10",
    "BI 定制报表页面脱敏",
    "验证 BI 定制报表页面按各字段规则脱敏显示个人信息。",
    "SIT BI 存在包含指定测试用户数据的定制报表。",
    ["登录 SIT BI。", "打开指定定制报表。", "选择包含指定用户的查询条件。", "执行查询。"],
)
add_business_surface(
    "BI / 定制报表 / 下载",
    "ZenTao 4319 BI 定制报表；checking rules #1-7；签字决策 Q-9/Q-10",
    "BI 定制报表下载文件脱敏",
    "验证 BI 定制报表下载文件按各字段规则脱敏保存个人信息。",
    "SIT BI 存在包含指定测试用户数据的定制报表；测试账号可下载。",
    ["登录 SIT BI。", "打开指定定制报表。", "执行查询。", "下载报表。", "打开下载文件。"],
)
add_business_surface(
    "BI / 定制报表 / 导出",
    "ZenTao 4319 BI 定制报表；checking rules #1-7；签字决策 Q-9/Q-10",
    "BI 定制报表导出文件脱敏",
    "验证 BI 定制报表导出文件按各字段规则脱敏保存个人信息。",
    "SIT BI 存在包含指定测试用户数据的定制报表；测试账号可导出。",
    ["登录 SIT BI。", "打开指定定制报表。", "执行查询。", "执行报表导出。", "打开导出文件。"],
)


COLUMNS = [
    "Label",
    "Test Case ID",
    "Module/Feature",
    "Priority",
    "Severity",
    "Collected from",
    "Test Scenario",
    "Test Case Description",
    "Preconditions",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Test Case Owner",
    "Environment",
    "Execution Date",
    "Executed By",
    "Actual Result",
    "Status",
    "Comments/Remarks",
    "Screenshots",
]


def md_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>")


def write_markdown() -> None:
    lines = [
        "# 脱敏测试用例",
        "",
        "## 范围摘要",
        "",
        "- 需求：ZenTao story 4319。",
        "- 签字范围：手机号、邮箱、姓名、身份证/实名制 ID、地址、生日、性别 7 类字段。",
        "- 模块：Website、PDA、后台、BI 定制报表，全部放在同一个 Test Cases Sheet。",
        "- 排除：BI Dashboard、订单权限/审计、KMS/集中存储、复杂格式边界、多语言、旧用例回写。",
        "- Ticket type：门票和座票相关页面分别覆盖。",
        "",
        "## 测试用例列表",
        "",
        "| " + " | ".join(h.replace("\n", " ") for h in HEADERS) + " |",
        "|" + "|".join("---" for _ in HEADERS) + "|",
    ]
    for case in cases:
        lines.append("| " + " | ".join(md_cell(case[col]) for col in COLUMNS) + " |")
    lines.extend(
        [
            "",
            "## 覆盖说明",
            "",
            f"- 用例总数：{len(cases)}。",
            "- 7 类字段规则均有独立用例；业务页面用例再验证模块内组合展示。",
            "- Website 门票与座票、PDA 门票与座票、后台门票与座票均分开编写。",
            "- 报表覆盖页面、预览、下载和导出；BI 仅覆盖定制报表。",
            "- 设计阶段执行字段保持为空。",
            "",
            "## 已确认排除项",
            "",
            "- BI Dashboard。",
            "- 订单权限和审计。",
            "- KMS、集中存储和密钥相关技术测试。",
            "- 未在 checking rules 中给出的复杂边界和格式化组合。",
            "- 多语言、重复脱敏和标准产品回流。",
            "",
        ]
    )
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")


def estimated_row_height(values: list[str]) -> float:
    chars_per_line = {6: 28, 7: 24, 8: 26, 9: 28, 10: 22, 11: 28, 12: 23}
    estimates = []
    for col_idx, width in chars_per_line.items():
        value = values[col_idx - 1] or ""
        logical = 0
        for line in str(value).splitlines() or [""]:
            logical += max(1, math.ceil(len(line) / width))
        estimates.append(logical)
    return float(min(260, max(70, 18 * max(estimates) + 12)))


def write_workbook() -> None:
    shutil.copy2(TEMPLATE, XLSX_OUT)
    wb = load_workbook(XLSX_OUT)
    ws = wb["Test Cases"]

    sample_styles = []
    for col in range(1, 21):
        cell = ws.cell(2, col)
        sample_styles.append(
            {
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
            }
        )

    max_target_row = len(cases) + 1
    for row in range(2, max(ws.max_row, max_target_row) + 1):
        for col in range(1, 21):
            ws.cell(row, col).value = None

    for row_idx, case in enumerate(cases, start=2):
        values = [case[col] for col in COLUMNS]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx)
            style = sample_styles[col_idx - 1]
            cell.font = copy(style["font"])
            cell.fill = copy(style["fill"])
            cell.border = copy(style["border"])
            cell.alignment = copy(style["alignment"])
            cell.number_format = style["number_format"]
            cell.protection = copy(style["protection"])
            cell.value = value or None
        ws.row_dimensions[row_idx].height = estimated_row_height(values)

    # Remove unused trailing template rows while keeping the header and all real cases.
    if ws.max_row > max_target_row:
        ws.delete_rows(max_target_row + 1, ws.max_row - max_target_row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:T{max_target_row}"
    ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(XLSX_OUT)


if __name__ == "__main__":
    write_markdown()
    write_workbook()
    print(f"Generated {len(cases)} cases")
    print(MD_OUT)
    print(XLSX_OUT)
