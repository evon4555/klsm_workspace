"""Clean up email/mobile residue in cases that were already split but text leaked.
Affects AUTH-103 / 105 / 106 / 107 / 117.
"""
import openpyxl, io, sys
from openpyxl.styles import PatternFill
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX = 'test-cases-registration-login_2026-06-12.xlsx'
YELLOW = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')

UPDATES = {
    'SIT-TC-WEB-AUTH-103': {
        # contact name is optional — position is relative to whichever contact field is rendered
        'Test Case Description': 'Verify the Order page change: contact-name field is optional and repositioned after the contact field (email or mobile, whichever the user picked); orders submit successfully with the name left empty.',
        'Test Steps': '1. On the order page, pick a contact method (email OR mobile) and fill it.\n'
                      '2. Locate the contact-name input.\n'
                      '3. Verify it is positioned AFTER the picked contact field (not before).\n'
                      '4. Verify it has no required indicator.\n'
                      '5. Leave contact name empty, fill the rest, submit.',
    },
    'SIT-TC-WEB-AUTH-105': {
        # email-only branch — drop "(or mobile)" residue
        'Test Steps': '1. Pick Email; enter an unregistered email; click elsewhere to blur.\n'
                      '2. Observe the area below the email field.\n'
                      '3. Verify the consent checkbox + description text appear in the active UI language.\n'
                      '4. Switch UI language to zh-HK and en; verify the corresponding text.\n'
                      '5. Verify the checkbox is unchecked by default.',
    },
    'SIT-TC-WEB-AUTH-106': {
        # email-only branch — drop "(or mobile)" residue
        'Test Steps': '1. Pick Email; enter a registered email; blur.\n'
                      '2. Observe the area below the email field.',
    },
    'SIT-TC-WEB-AUTH-107': {
        # email-only branch — clean expected wording
        'Expected Result': '1. The guest temp account is converted to a registered account.\n'
                           '2. The EMAIL is stored on the account; the mobile field is empty per PRD.\n'
                           '3. Conversion tag is set on the account.\n'
                           '4. Order is linked to the converted account.\n'
                           '5. Welcome reward pack (if configured) is granted.',
    },
    'SIT-TC-WEB-AUTH-117': {
        # mobile-only branch — replace {邮箱/手机} placeholder with the mobile-specific {手机号}
        'Test Case Description': 'Verify the payment-success page when conversion happened via MOBILE: the 3-language registration-success message renders with the {手机号} placeholder filled in; the "Back to Home" button routes to the Homepage. (Updated 2026-06-12: button changed from "Go to My Account" -> "Back to Home".)',
        'Expected Result': 'Same as AUTH-109 expectations, but the {手机号} placeholder resolves to the bound mobile (not {邮箱}); button label and destination identical (Back to Home -> Homepage).',
    },
}

wb = openpyxl.load_workbook(XLSX)
ws = wb['Test Cases']
hdr = {str(ws.cell(1, c).value).split('\n')[0].strip(): c for c in range(1, ws.max_column+1) if ws.cell(1,c).value}
id_col = hdr['Test Case ID']
row_by_id = {ws.cell(rn, id_col).value: rn for rn in range(2, ws.max_row+1) if ws.cell(rn, id_col).value}

modified = 0
for tcid, fields in UPDATES.items():
    rn = row_by_id.get(tcid)
    if not rn:
        print(f'WARN: {tcid} not found')
        continue
    for col_name, val in fields.items():
        c = hdr[col_name]
        ws.cell(rn, c).value = val
        ws.cell(rn, c).fill = YELLOW
        modified += 1

wb.save(XLSX)
print(f'Cleaned {modified} cells')
