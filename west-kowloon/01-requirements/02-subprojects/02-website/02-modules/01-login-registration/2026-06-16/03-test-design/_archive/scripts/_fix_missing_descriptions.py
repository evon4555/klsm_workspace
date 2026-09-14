"""Fix the missing Test Case Description column for AUTH-089~109.

These 21 cells were left empty because my generator script omitted the
'Test Case Description' key for those rows. The Test Scenario column has a
title, but Description needs a 1-sentence "what does this verify" sentence.
"""
import openpyxl, io, sys
from openpyxl.styles import PatternFill
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX = 'test-cases-registration-login_2026-06-12.xlsx'
YELLOW = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')

DESCRIPTIONS = {
    'SIT-TC-WEB-AUTH-089': 'Verify that the Change Password flow rejects a wrong current password — submission is blocked with a localized "current password incorrect" error and the password is not changed.',
    'SIT-TC-WEB-AUTH-090': 'Verify that the Change Password flow rejects a wrong or expired OTP — submission is blocked, the localized OTP error is shown, and the password is not changed.',
    'SIT-TC-WEB-AUTH-091': 'Verify the OTP-channel routing rule in Change Password — email-only account receives OTP on email; mobile-only on mobile; account with both defaults to email.',
    'SIT-TC-WEB-AUTH-092': 'Verify that the new-password field flags rule violations inline on focus-out (not only on submit) and that Submit is blocked while the inline error is showing.',
    'SIT-TC-WEB-AUTH-093': 'Verify that the image CAPTCHA is mandatory in Change Password — missing or wrong CAPTCHA does NOT unlock the OTP step.',

    'SIT-TC-WEB-AUTH-094': 'Verify the happy-path Bind-Email flow from Personal Center for a mobile-only account: format check, uniqueness check (twice), OTP send + verify, DB unique-constraint, success, page refresh.',
    'SIT-TC-WEB-AUTH-095': 'Verify the happy-path Bind-Mobile flow from Personal Center for an email-only account: same gating as AUTH-094, applied to the mobile field.',
    'SIT-TC-WEB-AUTH-096': 'Verify Bind uniqueness check 1 (pre-OTP-send): when the entered contact is already bound to another account, OTP is NOT sent and the 3-language error is shown.',
    'SIT-TC-WEB-AUTH-097': 'Verify Bind uniqueness check 2 (post-OTP-verify): when another account takes the same contact during the OTP wait, the second check catches the race and bind fails.',
    'SIT-TC-WEB-AUTH-098': 'Verify Bind OTP-error handling: a wrong OTP returns the OTP-incorrect error; an expired OTP returns the OTP-expired error; bind fails in both cases.',
    'SIT-TC-WEB-AUTH-099': 'Verify Bind entry-point visibility: the "Bind Now" CTA appears only when the corresponding contact field is empty on the account.',

    'SIT-TC-WEB-AUTH-100': 'Verify the Project Detail page changes for guest-to-member: primary CTA "Quick Buy Without Login" + Cloudflare bot check + silent guest login; secondary "Buy tickets as member" routes to the existing logged-in Buy Tickets flow.',
    'SIT-TC-WEB-AUTH-101': 'Verify the Order page contact-method picker: exactly one of Email or Mobile is required at a time (controlled by the picker); the non-chosen method has no required marker; submit blocked when chosen is empty.',
    'SIT-TC-WEB-AUTH-102': 'Verify the Order page data-submission rule: when both Email and Mobile are filled, only the user-chosen method is submitted in the payload and recorded on the order; the other field is dropped.',
    'SIT-TC-WEB-AUTH-103': 'Verify the Order page change: contact-name field is optional and repositioned after the contact email/mobile inputs; orders submit successfully with the name left empty.',
    'SIT-TC-WEB-AUTH-104': 'Verify that the Order page enforces an OTP step for guest users (gates submission), while logged-in users continue to skip OTP entirely.',
    'SIT-TC-WEB-AUTH-105': 'Verify the unregistered branch of the blur-time uniqueness check: when the entered contact is unregistered, the 3-language opt-in checkbox + description text appear and are unchecked by default.',
    'SIT-TC-WEB-AUTH-106': 'Verify the registered branch of the blur-time uniqueness check: when the entered contact is already registered, NO opt-in checkbox/description is shown and the guest purchase continues normally.',
    'SIT-TC-WEB-AUTH-107': 'Verify the upgrade branch: opt-in checked + payment success converts the guest temp account into a registered account (storing only the chosen contact field), tags it as "guest-purchase-conversion", links the order, and grants the welcome reward pack if configured.',
    'SIT-TC-WEB-AUTH-108': 'Verify the no-upgrade branch: opt-in unchecked + payment success leaves the guest temp account as a guest; no registered account is created; order is recorded under the guest temp account per the existing flow.',
    'SIT-TC-WEB-AUTH-109': 'Verify the payment-success page: when conversion has occurred, the 3-language registration-success message is shown with the chosen contact placeholder filled in, and the "Go to My Account" button routes to Personal Center.',
}

wb = openpyxl.load_workbook(XLSX)
ws = wb['Test Cases']
hdr = {str(ws.cell(1, c).value).split('\n')[0].strip(): c for c in range(1, ws.max_column+1) if ws.cell(1,c).value}
id_col = hdr['Test Case ID']
desc_col = hdr['Test Case Description']

fixed = 0
for rn in range(2, ws.max_row+1):
    tcid = ws.cell(rn, id_col).value
    if tcid in DESCRIPTIONS:
        cell = ws.cell(rn, desc_col)
        cell.value = DESCRIPTIONS[tcid]
        cell.fill = YELLOW
        fixed += 1

wb.save(XLSX)
print(f'Fixed {fixed} descriptions')

# Verify
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2['Test Cases']
empty_left = 0
for rn in range(2, ws2.max_row+1):
    tcid = ws2.cell(rn, id_col).value
    if not tcid or not tcid.startswith('SIT-TC-WEB-AUTH-'):
        continue
    num = int(tcid.split('-')[-1])
    if num < 88: continue
    if not ws2.cell(rn, desc_col).value:
        empty_left += 1
        print(f'  STILL EMPTY: {tcid}')
print(f'Empty descriptions remaining: {empty_left}')
