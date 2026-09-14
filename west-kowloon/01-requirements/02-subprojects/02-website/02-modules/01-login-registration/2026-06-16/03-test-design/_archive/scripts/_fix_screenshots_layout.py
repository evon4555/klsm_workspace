"""Normalize embedded screenshots in test-cases-registration-login_2026-06-12.xlsx:
  - One image per row (drop duplicates / overlapping older images, keep the latest).
  - Every kept image resized to fit within column T (~385 px).
  - Row height adjusted so the image isn't visually clipped.
"""
import openpyxl, io, sys
from collections import defaultdict
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX = 'test-cases-registration-login_2026-06-12.xlsx'

# Column T is the Screenshots column. Its width is 55 (Excel chars) which is
# ~385px with the default font. Target image width of 340 leaves a small
# margin so the picture clearly sits "inside" the cell box.
TARGET_WIDTH_PX = 340
PX_TO_EMU = 9525           # 1 px at 96 DPI = 9525 EMU
EMU_PER_PX = PX_TO_EMU
MIN_ROW_HEIGHT_PT = 50     # row height in points; 1pt ≈ 1.333 px; >= 50pt -> >= 67px

wb = openpyxl.load_workbook(XLSX)
ws = wb['Test Cases']
hdr = {str(ws.cell(1, c).value).split('\n')[0].strip(): c
       for c in range(1, ws.max_column+1) if ws.cell(1, c).value}
id_col = hdr['Test Case ID']

# 1. Group images by anchor row
groups = defaultdict(list)
for img in list(ws._images):
    f = img.anchor._from
    groups[f.row + 1].append(img)  # 1-indexed row

# 2. For each row, keep the LAST image only (most recently added wins)
keep = []
dropped_rows = []
for row, imgs in groups.items():
    if len(imgs) > 1:
        tcid = ws.cell(row, id_col).value
        dropped_rows.append((row, tcid, len(imgs)))
    keep.append(imgs[-1])

print(f'Multi-image rows collapsed: {len(dropped_rows)}')
for (row, tcid, n) in dropped_rows:
    print(f'  row {row} {tcid}: {n} images -> kept latest')

# 3. Resize each kept image so width <= TARGET_WIDTH_PX AND height <= TARGET_HEIGHT_PX,
#    preserving aspect ratio (use min scale of width-fit and height-fit).
TARGET_HEIGHT_PX = 240   # cap so portrait/full-page shots don't dominate the row
oversize_fixed = 0
for img in keep:
    ext = getattr(img.anchor, 'ext', None)
    if ext is None or ext.cx <= 0 or ext.cy <= 0:
        continue
    cur_w_px = ext.cx / EMU_PER_PX
    cur_h_px = ext.cy / EMU_PER_PX
    if cur_w_px <= TARGET_WIDTH_PX and cur_h_px <= TARGET_HEIGHT_PX:
        continue
    scale_w = TARGET_WIDTH_PX / cur_w_px if cur_w_px > TARGET_WIDTH_PX else 1.0
    scale_h = TARGET_HEIGHT_PX / cur_h_px if cur_h_px > TARGET_HEIGHT_PX else 1.0
    scale = min(scale_w, scale_h)
    ext.cx = int(ext.cx * scale)
    ext.cy = int(ext.cy * scale)
    oversize_fixed += 1

print(f'Images resized to fit column T: {oversize_fixed}')

# 4. Ensure each row hosting an image has enough vertical room for it
adjusted_rows = 0
for img in keep:
    f = img.anchor._from
    row = f.row + 1
    ext = getattr(img.anchor, 'ext', None)
    if not ext:
        continue
    needed_pt = max(MIN_ROW_HEIGHT_PT, ext.cy / EMU_PER_PX * 0.75 + 4)  # px -> pt = px / 1.333
    cur_h = ws.row_dimensions[row].height or 0
    if cur_h < needed_pt:
        ws.row_dimensions[row].height = needed_pt
        adjusted_rows += 1

print(f'Row heights raised to fit images: {adjusted_rows}')

# 5. Replace the images list with the kept (and now resized) set
ws._images = keep

wb.save(XLSX)
print(f'\nFinal image count: {len(ws._images)}')

# Verify
wb2 = openpyxl.load_workbook(XLSX)
ws2 = wb2['Test Cases']
groups2 = defaultdict(list)
for img in ws2._images:
    groups2[img.anchor._from.row + 1].append(img)
overlap = [(r, len(v)) for r, v in groups2.items() if len(v) > 1]
oversized = []
for img in ws2._images:
    ext = img.anchor.ext
    if ext.cx / EMU_PER_PX > TARGET_WIDTH_PX + 5 or ext.cy / EMU_PER_PX > 250:
        oversized.append((img.anchor._from.row + 1, ext.cx / EMU_PER_PX, ext.cy / EMU_PER_PX))
print(f'\nVerify: rows still with multi-image: {len(overlap)}')
print(f'Verify: images still oversized: {len(oversized)}')
if oversized[:5]:
    print('  examples:', oversized[:5])
