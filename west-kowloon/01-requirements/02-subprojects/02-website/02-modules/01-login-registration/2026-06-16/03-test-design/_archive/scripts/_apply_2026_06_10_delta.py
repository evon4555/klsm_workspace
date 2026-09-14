"""Apply 2026-06-10 PRD delta to test-cases-registration-login xlsx.
- Preserves all existing rows / formatting / fills.
- Modifies specific cells (yellow highlight) for PRD-driven changes.
- Appends new TCs (yellow highlight on all populated cells).
- Saves as test-cases-registration-login_2026-06-12.xlsx (new dated file).
"""
import openpyxl, io, sys, shutil
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

SRC = 'test-cases-registration-login_2026-06-02.xlsx'
DST = 'test-cases-registration-login_2026-06-12.xlsx'

YELLOW = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb['Test Cases']

# Build header map: normalized name -> column index
hdr_raw = {}
for c in range(1, ws.max_column + 1):
    v = ws.cell(1, c).value
    if v is None:
        continue
    key = str(v).split('\n')[0].strip()
    hdr_raw[key] = c

# Confirm we found all needed columns
NEEDED = ['Test Case ID', 'Test Scenario', 'Test Steps', 'Expected Result', 'Status', 'Comments/Remarks',
          'Module/Feature', 'Priority', 'Severity', 'Collected from', 'Test Case Description',
          'Preconditions', 'Test Data', 'Test Case Owner', 'Environment', 'Label']
missing = [n for n in NEEDED if n not in hdr_raw]
if missing:
    print('MISSING headers:', missing)
    print('FOUND:', list(hdr_raw.keys()))
    sys.exit(1)

# Find row by Test Case ID
id_col = hdr_raw['Test Case ID']
row_by_id = {ws.cell(rn, id_col).value: rn for rn in range(2, ws.max_row + 1) if ws.cell(rn, id_col).value}

# Reference row for style/border copy when appending
REF_ROW = row_by_id.get('SIT-TC-WEB-AUTH-001', 2)

def set_cell(rn, col_name, val, *, highlight=True):
    c = hdr_raw[col_name]
    cell = ws.cell(rn, c)
    cell.value = val
    if highlight:
        cell.fill = YELLOW

def add_row(tcid, **vals):
    """Append a new row. All populated cells get yellow fill."""
    rn = ws.max_row + 1
    # Copy formatting from REF_ROW to maintain borders/wrapping
    for c in range(1, ws.max_column + 1):
        ref = ws.cell(REF_ROW, c)
        new = ws.cell(rn, c)
        if ref.has_style:
            new.font = copy(ref.font)
            new.alignment = copy(ref.alignment)
            new.border = copy(ref.border)
            new.number_format = ref.number_format
    set_cell(rn, 'Test Case ID', tcid)
    for k, v in vals.items():
        if v is not None and v != '':
            set_cell(rn, k, v)
    return rn

# =============================================================
# Section A: Modify existing TCs (PRD-driven changes)
# =============================================================

# --- AUTH-001: Successful native registration ---
# Flow becomes gated 5-step per 2026-06-10 PRD (was a single submit)
rn = row_by_id['SIT-TC-WEB-AUTH-001']
set_cell(rn, 'Test Case Description',
         'Verify a visitor can complete native registration through the gated 5-step flow: '
         'select email/mobile -> image CAPTCHA -> OTP -> password (with focus-out rule check) -> Submit.')
set_cell(rn, 'Test Steps',
         '1. Open registration page; select Email (or Mobile) mode.\n'
         '2. Enter email/mobile + solve image CAPTCHA (动态码); confirm CAPTCHA passes before next step is unlocked.\n'
         '3. Click "Get Code"; receive OTP and enter it; confirm OTP passes before password step is unlocked.\n'
         '4. Enter password; verify focus-out triggers inline password-rule validation.\n'
         '5. Enter password again (repeat).\n'
         '6. Click Submit.')
set_cell(rn, 'Expected Result',
         '1. CAPTCHA gating: subsequent fields are unlocked only after CAPTCHA passes.\n'
         '2. OTP gating: password step is unlocked only after OTP passes.\n'
         '3. Password rule violations are flagged inline on focus-out, not only on submit.\n'
         '4. On Submit, account is created and the user is logged in.\n'
         '5. No duplicate account is created.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12 per 2026-06-10 PRD §"用户邮箱/手机+密码注册": gated 5-step flow + focus-out password rule.')

# --- AUTH-004: Duplicate account prevention - add 3-lang error text ---
rn = row_by_id['SIT-TC-WEB-AUTH-004']
set_cell(rn, 'Test Case Description',
         'Verify duplicate email/mobile registration is blocked with the exact 3-language error specified by PRD.')
set_cell(rn, 'Expected Result',
         '1. Registration is blocked.\n'
         '2. For duplicate EMAIL, the email field shows the 3-language error (red):\n'
         '   - zh-CN: 该邮箱已注册，请登录或找回密码。\n'
         '   - zh-HK: 此電子郵件已註冊，請直接登入或重設密碼。\n'
         '   - en: This email address is already registered. Please sign in or reset your password.\n'
         '3. For duplicate MOBILE, the mobile field shows:\n'
         '   - zh-CN: 该手机号已注册，请直接登录。\n'
         '   - zh-HK: 此手機號碼已註冊，請直接登入。\n'
         '   - en: This mobile number is already registered. Please sign in or reset your password.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12 per 2026-06-10 PRD: exact 3-language wording required.')

# --- AUTH-010: Invalid password login - generic error ---
rn = row_by_id['SIT-TC-WEB-AUTH-010']
set_cell(rn, 'Expected Result',
         '1. Login rejected.\n'
         '2. Generic error: "用户名或密码错误" (zh-CN) / "用戶名或密碼錯誤" (zh-HK) / "Username or password incorrect" (en).\n'
         '3. System MUST NOT differentiate between "account does not exist" and "wrong password".\n'
         '4. User remains unauthenticated.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12 per 2026-06-10 PRD: generic error - no account enumeration.')

# --- AUTH-012: Unknown account login - same generic error ---
rn = row_by_id['SIT-TC-WEB-AUTH-012']
set_cell(rn, 'Expected Result',
         '1. Login rejected with the SAME generic error as wrong-password (see AUTH-010): "用户名或密码错误" / "Username or password incorrect".\n'
         '2. NO "go register" guidance must be shown.\n'
         '3. User remains unauthenticated.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12 per 2026-06-10 PRD: unknown account no longer routed to registration; same generic error.')

# --- AUTH-023: Guest session validity countdown - 15min ---
rn = row_by_id['SIT-TC-WEB-AUTH-023']
set_cell(rn, 'Expected Result',
         '1. Countdown is displayed.\n'
         '2. Guest session validity = 15 minutes from login (per 2026-06-10 PRD; was 10 minutes in 06-02 PRD).')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12: 10min -> 15min per 2026-06-10 PRD §游客模式.')

# --- AUTH-024: Guest timeout clears state ---
rn = row_by_id['SIT-TC-WEB-AUTH-024']
set_cell(rn, 'Test Steps',
         '1. Login as guest.\n'
         '2. Wait 15 minutes for the validity to expire.\n'
         '3. Try to continue the previous guest flow (cart / order draft / filled information).')
set_cell(rn, 'Expected Result',
         '1. After 15-minute timeout, cart / order draft / filled information are all cleared.\n'
         '2. User cannot continue the previous guest flow without re-logging-in as guest.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12: 10min -> 15min.')

# --- AUTH-036, AUTH-037: Email-link registration - deprecated ---
for tcid in ['SIT-TC-WEB-AUTH-036', 'SIT-TC-WEB-AUTH-037']:
    rn = row_by_id[tcid]
    set_cell(rn, 'Comments/Remarks',
             'DEPRECATED by 2026-06-10 PRD: email-link registration flow removed; only in-page OTP flow is in scope.')

# --- AUTH-038: 5-min-remaining popup - align with 15min window ---
rn = row_by_id['SIT-TC-WEB-AUTH-038']
set_cell(rn, 'Expected Result',
         '1. With guest validity = 15 minutes, the "5 minutes remaining" popup appears at the 10-minute mark from guest login.\n'
         '2. Popup wording shown in 3 languages per back-office configuration.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12: aligned with new 15-minute guest window.')

# --- AUTH-066: Password-mode session auto-expiry = 30 days ---
rn = row_by_id['SIT-TC-WEB-AUTH-066']
set_cell(rn, 'Expected Result',
         '1. Non-guest password-mode session expires 30 days from last successful authentication (per 2026-06-10 PRD).\n'
         '2. After expiry, the user is logged out and must re-authenticate.\n'
         '3. The 30-day window is a code constant but can be tuned.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12: PRD now specifies 30-day non-guest session.')

# --- AUTH-067: OTP-mode session auto-expiry = 30 days ---
rn = row_by_id['SIT-TC-WEB-AUTH-067']
set_cell(rn, 'Expected Result',
         '1. Non-guest OTP-mode session expires 30 days from last successful authentication (per 2026-06-10 PRD).\n'
         '2. After expiry, the user is logged out and must re-authenticate.')
set_cell(rn, 'Comments/Remarks',
         'Updated 2026-06-12: PRD now specifies 30-day non-guest session.')

# =============================================================
# Section B: New TCs (added at end with yellow fill)
# =============================================================

NEW_TCS = []

# ---- Change Password (new module) ----
NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-088',
    Label='SIT/UAT',
    **{'Module/Feature': 'Change Password'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §修改密码'},
    **{'Test Scenario': 'Successful change-password flow'},
    **{'Test Case Description': 'Verify the gated change-password flow: old password -> image CAPTCHA -> OTP -> new password (focus-out rule check) + repeat -> Submit -> success page -> auto-logout in 1-3s.'},
    Preconditions='Active registered account with known current password.',
    **{'Test Steps': '1. Open Personal Center -> Account Security -> Change Password.\n'
                    '2. Enter current password.\n'
                    '3. Solve image CAPTCHA (动态码); confirm next step unlocks only after CAPTCHA passes.\n'
                    '4. Click "Get Code"; receive OTP and enter it.\n'
                    '5. Enter new password; verify focus-out triggers inline rule validation.\n'
                    '6. Repeat new password.\n'
                    '7. Click Submit.'},
    **{'Test Data': 'Current password (valid); new password meets all rules.'},
    **{'Expected Result': '1. Success page is shown with 3-language title: "密码修改成功 / 密碼修改成功 / Password Changed Successfully".\n'
                         '2. Body text in 3 languages: "为了保障您的账号安全，请使用新密码重新登录。 / 為保障您的帳號安全，請使用新密碼重新登入。 / For security reasons, please sign in again using your new password."\n'
                         '3. User is auto-logged-out after 1-3 seconds and returned to login page.\n'
                         '4. Old password no longer authenticates; new password authenticates.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT',
    Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-089',
    Label='SIT/UAT',
    **{'Module/Feature': 'Change Password'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §修改密码 step 5.a'},
    **{'Test Scenario': 'Change password - wrong current password blocks submission'},
    Preconditions='Active account.',
    **{'Test Steps': '1. Open Change Password.\n'
                    '2. Enter WRONG current password.\n'
                    '3. Pass CAPTCHA + OTP + new password.\n'
                    '4. Click Submit.'},
    **{'Expected Result': '1. Submission blocked.\n'
                         '2. Inline error: "当前密码错误，请重新输入 / 當前密碼錯誤，請重新輸入 / Current password is incorrect, please re-enter".\n'
                         '3. Password is NOT changed.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-090',
    Label='SIT/UAT',
    **{'Module/Feature': 'Change Password'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §修改密码 step 5.b'},
    **{'Test Scenario': 'Change password - wrong/expired OTP blocks submission'},
    Preconditions='Active account.',
    **{'Test Steps': '1. Open Change Password; pass current password + CAPTCHA.\n'
                    '2. Click Get Code.\n'
                    '3. Enter wrong OTP -> Submit. Repeat with expired OTP.\n'},
    **{'Expected Result': '1. Submission blocked.\n'
                         '2. Localized error indicating OTP wrong/expired.\n'
                         '3. Password is NOT changed.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-091',
    Label='SIT',
    **{'Module/Feature': 'Change Password'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §修改密码 step 3'},
    **{'Test Scenario': 'Change password - OTP routing logic (email-only / mobile-only / both)'},
    Preconditions='Three test accounts: A (email only), B (mobile only), C (both email and mobile).',
    **{'Test Steps': '1. For account A: open Change Password, click Get Code -> verify OTP is sent to email.\n'
                    '2. For account B: same -> verify OTP is sent to mobile.\n'
                    '3. For account C: same -> verify OTP defaults to email.'},
    **{'Expected Result': '1. Email-only account: OTP goes to email.\n'
                         '2. Mobile-only account: OTP goes to mobile.\n'
                         '3. Both: OTP defaults to email per PRD.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-092',
    Label='SIT',
    **{'Module/Feature': 'Change Password'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §修改密码 step 4'},
    **{'Test Scenario': 'Change password - new password fails rule check on focus-out (inline)'},
    Preconditions='Active account.',
    **{'Test Steps': '1. Open Change Password; pass current password + CAPTCHA + OTP.\n'
                    '2. Type a new password that violates rules; move focus away.\n'
                    '3. Observe inline message under the field.\n'
                    '4. Try Submit anyway.'},
    **{'Expected Result': '1. Inline rule-violation message appears under the field on focus-out.\n'
                         '2. Submit is blocked.\n'
                         '3. Password is NOT changed.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-093',
    Label='SIT',
    **{'Module/Feature': 'Change Password'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §修改密码'},
    **{'Test Scenario': 'Change password - image CAPTCHA mandatory and validates'},
    Preconditions='Active account.',
    **{'Test Steps': '1. Open Change Password.\n'
                    '2. Enter current password.\n'
                    '3. Try to skip CAPTCHA / enter wrong CAPTCHA -> observe.\n'
                    '4. Enter correct CAPTCHA -> verify OTP step unlocks.'},
    **{'Expected Result': '1. Without/with wrong CAPTCHA, the OTP step does not unlock.\n'
                         '2. With correct CAPTCHA, the flow proceeds.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

# ---- Bind mobile/email via Personal Center (new module) ----
NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-094',
    Label='SIT/UAT',
    **{'Module/Feature': 'Bind Contact'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §绑定手机/邮箱'},
    **{'Test Scenario': 'Personal Center: Bind Email to a mobile-only account (happy path)'},
    Preconditions='Account that has a mobile number but NO email.',
    **{'Test Steps': '1. Log in; open Personal Center -> Personal Information (view page).\n'
                    '2. Observe that the Email row shows "Bind Now" entry.\n'
                    '3. Click "Bind Now"; enter a never-used email; format check passes.\n'
                    '4. (Uniqueness check 1) Confirm no error.\n'
                    '5. Click Send OTP; receive code; enter and submit.\n'
                    '6. (Uniqueness check 2 + DB constraint succeeds.)\n'
                    '7. Observe success and page refresh.'},
    **{'Expected Result': '1. Format check accepts a valid email.\n'
                         '2. Uniqueness check 1 passes for unused email.\n'
                         '3. OTP send + verify succeed.\n'
                         '4. Uniqueness check 2 + DB unique constraint succeed.\n'
                         '5. Bind success message shown; page refreshed; the email now appears in Personal Information.\n'
                         '6. Bind audit log is recorded.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-095',
    Label='SIT/UAT',
    **{'Module/Feature': 'Bind Contact'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §绑定手机/邮箱'},
    **{'Test Scenario': 'Personal Center: Bind Mobile to an email-only account (happy path)'},
    Preconditions='Account that has an email but NO mobile.',
    **{'Test Steps': '1. Log in; open Personal Center -> Personal Information.\n'
                    '2. Observe that the Mobile row shows "Bind Now" entry.\n'
                    '3. Bind via mobile + OTP, same pattern as AUTH-094.'},
    **{'Expected Result': 'Same gate as AUTH-094, applied to mobile. Page refreshes and mobile now appears in Personal Information.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-096',
    Label='SIT',
    **{'Module/Feature': 'Bind Contact'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §绑定手机/邮箱'},
    **{'Test Scenario': 'Bind - uniqueness check 1 catches conflict (email already bound to another account)'},
    Preconditions='Account A: mobile-only. Account B: holds email X.',
    **{'Test Steps': '1. Log in as Account A; open Bind Email.\n'
                    '2. Enter email X (already on Account B).\n'
                    '3. Observe error before OTP send.'},
    **{'Expected Result': '1. Uniqueness check 1 fails: "该邮箱/手机已绑定其他账号" (3 languages).\n'
                         '2. OTP is NOT sent.\n'
                         '3. Account A remains unchanged.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-097',
    Label='SIT',
    **{'Module/Feature': 'Bind Contact'},
    Priority='Medium', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §绑定手机/邮箱'},
    **{'Test Scenario': 'Bind - uniqueness check 2 catches race during OTP entry'},
    Preconditions='Account A: mobile-only. Two browser sessions or a back-office script to simulate a competing bind.',
    **{'Test Steps': '1. Log in as Account A; start Bind Email with email Y.\n'
                    '2. After OTP send, before submitting, have email Y get bound to another account (race).\n'
                    '3. Submit OTP.'},
    **{'Expected Result': '1. Uniqueness check 2 fails; error: "该邮箱/手机已绑定其他账号".\n'
                         '2. Account A is NOT bound to Y.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-098',
    Label='SIT',
    **{'Module/Feature': 'Bind Contact'},
    Priority='Medium', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §绑定手机/邮箱'},
    **{'Test Scenario': 'Bind - OTP wrong/expired'},
    Preconditions='Account A: mobile-only.',
    **{'Test Steps': '1. Start Bind Email with a fresh email.\n'
                    '2. Send OTP; enter a wrong OTP; submit.\n'
                    '3. Resend OTP; wait until expiry; enter expired OTP; submit.'},
    **{'Expected Result': '1. Wrong OTP: error "OTP incorrect" (3 languages); bind fails.\n'
                         '2. Expired OTP: error "OTP expired"; bind fails.\n'
                         '3. Account A remains unbound to the new email.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-099',
    Label='SIT',
    **{'Module/Feature': 'Bind Contact'},
    Priority='Medium', Severity='Low',
    **{'Collected from': '2026-06-10 PRD §绑定手机/邮箱'},
    **{'Test Scenario': 'Bind entry visibility - only shown when the target contact field is empty'},
    Preconditions='Account with both email and mobile.',
    **{'Test Steps': '1. Log in; open Personal Information.\n'
                    '2. Observe email and mobile rows.\n'
                    '3. Verify no "Bind Now" entry for either.'},
    **{'Expected Result': '1. Both contact rows display the existing value with no "Bind Now" CTA.\n'
                         '2. The account cannot trigger a duplicate bind.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

# ---- Guest -> Register: full 10-case set per 游客购票引导注册功能.md (2026.6.3 新增) ----

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-100',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §游客购票引导注册 / 项目详情页'},
    **{'Test Scenario': 'Project detail page: "Quick Buy Without Login" + "Buy tickets as member" entries'},
    Preconditions='Visitor on project detail page; not logged in.',
    **{'Test Steps': '1. Open a project detail page as a visitor.\n'
                    '2. Locate the primary CTA — verify it reads "Quick Buy Without Login".\n'
                    '3. Verify a secondary entry below it reads "Buy tickets as member".\n'
                    '4. Click "Quick Buy Without Login"; observe Cloudflare bot check then silent guest login.\n'
                    '5. Click "Buy tickets as member"; verify it routes to the current Buy Tickets flow (login required).'},
    **{'Expected Result': '1. Primary CTA labelled "Quick Buy Without Login".\n'
                         '2. Secondary entry labelled "Buy tickets as member" visible below.\n'
                         '3. Cloudflare bot check fires on the Quick Buy path; on pass, silent login is performed and the user enters guest mode.\n'
                         '4. "Buy tickets as member" follows the existing Buy Tickets behaviour (gates to login if not authenticated).'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-101',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §订单确认页 / 调整联系邮箱/手机选择'},
    **{'Test Scenario': 'Order page: contact method picker — selected method is required, the other is optional'},
    Preconditions='Guest user reaches the order-confirmation page.',
    **{'Test Steps': '1. On the order page, observe the contact method picker (Email / Mobile).\n'
                    '2. Pick Email; verify Email becomes required (red *) and Mobile becomes optional.\n'
                    '3. Switch the picker to Mobile; verify the requirement flips.\n'
                    '4. Try submitting with the chosen method empty — submit must be blocked.'},
    **{'Expected Result': '1. Exactly one method is required at a time, controlled by the picker.\n'
                         '2. The non-chosen method has no validation error if left empty.\n'
                         '3. Submit is blocked when the chosen method is empty.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-102',
    Label='SIT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §订单确认页 step 1.b'},
    **{'Test Scenario': 'Order page: when BOTH email and mobile are filled, only the chosen one is submitted'},
    Preconditions='Guest user on order page.',
    **{'Test Steps': '1. Fill BOTH email and mobile.\n'
                    '2. Pick Email as the selected method.\n'
                    '3. Submit the order.\n'
                    '4. Inspect the submit-order payload (DevTools Network) AND the resulting order record.\n'
                    '5. Repeat with Mobile selected.'},
    **{'Expected Result': '1. With Email selected: payload + order record contain ONLY the email; mobile is omitted.\n'
                         '2. With Mobile selected: payload + order record contain ONLY the mobile; email is omitted.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-103',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='Medium', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §订单确认页 step 2'},
    **{'Test Scenario': 'Order page: contact name is now OPTIONAL and repositioned after email/mobile'},
    Preconditions='Guest user on order page.',
    **{'Test Steps': '1. On the order page, locate the contact name input.\n'
                    '2. Verify it is positioned AFTER email/mobile (not before).\n'
                    '3. Verify it has no required indicator.\n'
                    '4. Leave contact name empty, fill the rest, submit.'},
    **{'Expected Result': '1. Contact name field appears after email/mobile inputs.\n'
                         '2. No "required" marker on contact name.\n'
                         '3. Order submits successfully with contact name empty.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-104',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §订单确认页 step 3'},
    **{'Test Scenario': 'Order page: OTP verification gates order submission (guest only; logged-in users skip)'},
    Preconditions='Two scenarios — guest user AND logged-in user reaching the order page.',
    **{'Test Steps': '1. As a GUEST: fill contact method; observe an OTP field appears.\n'
                    '2. Try to submit without OTP — blocked.\n'
                    '3. Send OTP; enter wrong OTP — submit blocked with localized error.\n'
                    '4. Enter correct OTP — submit allowed.\n'
                    '5. As a LOGGED-IN user: confirm there is NO OTP field on the order page.'},
    **{'Expected Result': '1. Guest path: OTP field is shown and gates submission.\n'
                         '2. Wrong/missing OTP blocks submission with the localized error.\n'
                         '3. Logged-in path: no OTP field; submission flow unchanged.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-105',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §账号校验 (unregistered branch)'},
    **{'Test Scenario': 'Order page: blur-time uniqueness check — UNREGISTERED contact shows 3-language opt-in checkbox + description (default unchecked)'},
    Preconditions='Guest user on order page; have a known UNREGISTERED email (or mobile).',
    **{'Test Steps': '1. Pick Email; enter an unregistered email; click elsewhere to blur.\n'
                    '2. Observe the area below the email field.\n'
                    '3. Verify the consent checkbox + description text appear in the active UI language.\n'
                    '4. Switch UI language to zh-HK and en; verify the corresponding text.\n'
                    '5. Verify the checkbox is unchecked by default.'},
    **{'Expected Result': '1. Checkbox label per language:\n'
                         '   - zh-CN: 同意同步注册为西九文化区注册用户\n'
                         '   - zh-HK: 同意同步註冊為西九文化區註冊用戶\n'
                         '   - en: I agree to register as a registered user of WestK.\n'
                         '2. Description text (3 languages) appears below per PRD:\n'
                         '   - zh-CN: 完成订单支付后，系统将自动为您创建西九文化区注册账号 ...\n'
                         '   - zh-HK: 完成訂單支付後，系統將自動為您建立西九文化區註冊帳號 ...\n'
                         '   - en: After successful payment, a WestK account will be automatically created for you and this order will be linked to your account.\n'
                         '3. Checkbox is unchecked by default and not required for order submission.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-106',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §账号校验 (registered branch)'},
    **{'Test Scenario': 'Order page: blur-time uniqueness check — ALREADY REGISTERED contact hides the opt-in checkbox'},
    Preconditions='Guest user on order page; have a known REGISTERED email (or mobile).',
    **{'Test Steps': '1. Pick Email; enter a registered email; blur.\n'
                    '2. Observe the area below the email field.\n'
                    '3. Repeat for mobile.'},
    **{'Expected Result': '1. No opt-in checkbox is shown.\n'
                         '2. No opt-in description text is shown.\n'
                         '3. The user can continue the guest purchase normally.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-107',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='High',
    **{'Collected from': '2026-06-10 PRD §提交订单 (upgrade branch)'},
    **{'Test Scenario': 'Submit + payment success WITH opt-in checked → guest account upgrades to registered user'},
    Preconditions='Guest user on order page; unregistered email/mobile; checkbox ticked.',
    **{'Test Steps': '1. Tick the opt-in checkbox.\n'
                    '2. Submit the order; complete payment.\n'
                    '3. Inspect the resulting account in admin / DB:\n'
                    '   a. Confirm the guest account was upgraded to a registered account.\n'
                    '   b. Confirm the chosen field (email or mobile) is stored as the account contact; the OTHER field is empty.\n'
                    '   c. Confirm the account carries a "guest-purchase-conversion" tag.\n'
                    '   d. Confirm the order is linked to this account.\n'
                    '4. Verify the membership-reward pack (if configured) is granted.'},
    **{'Expected Result': '1. The guest temp account is converted to a registered account.\n'
                         '2. The contact field stored matches the user choice; the other field is empty per PRD.\n'
                         '3. Conversion tag is set on the account.\n'
                         '4. Order is linked to the converted account.\n'
                         '5. Welcome reward pack (if configured) is granted.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-108',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §提交订单 (non-upgrade branch)'},
    **{'Test Scenario': 'Submit + payment success WITHOUT opt-in → no upgrade; normal guest purchase'},
    Preconditions='Guest user on order page; unregistered email/mobile; checkbox left unchecked.',
    **{'Test Steps': '1. Leave the opt-in checkbox UNCHECKED.\n'
                    '2. Submit the order; complete payment.\n'
                    '3. Inspect the user state in admin / DB.'},
    **{'Expected Result': '1. No account upgrade happens.\n'
                         '2. No registered account is created.\n'
                         '3. The guest temp account remains a guest.\n'
                         '4. Order is recorded under the guest temp account per existing guest-purchase flow.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

NEW_TCS.append(dict(
    tcid='SIT-TC-WEB-AUTH-109',
    Label='SIT/UAT',
    **{'Module/Feature': 'Guest-to-Member'},
    Priority='High', Severity='Medium',
    **{'Collected from': '2026-06-10 PRD §支付成功页'},
    **{'Test Scenario': 'Payment success page: 3-language registration-success message + "Go to My Account" button'},
    Preconditions='Guest user just completed AUTH-107 (converted on payment).',
    **{'Test Steps': '1. After payment success, land on the success page.\n'
                    '2. Verify the registration-success message in the active language.\n'
                    '3. Switch to zh-HK and en; verify the corresponding messages.\n'
                    '4. Verify the {邮箱/手机} placeholder is filled with the user-chosen contact.\n'
                    '5. Click "Go to My Account" / "前往我的账号" / "前往我的帳號".'},
    **{'Expected Result': '1. Message in 3 languages per PRD:\n'
                         '   - zh-CN: 恭喜！您已成功注册成为西九文化区注册用户。本次订单已自动关联至您的账号 ...\n'
                         '   - zh-HK: 恭喜！您已成功註冊成為西九文化區註冊用戶。本次訂單已自動關聯至您的帳號 ...\n'
                         '   - en: Congratulations! You have successfully registered as a WestK user. This order has been automatically linked to your account ...\n'
                         '2. {邮箱/手机} placeholder is replaced with the actual chosen value.\n'
                         '3. Button label per language: 前往我的账号 / 前往我的帳號 / Go to My Account.\n'
                         '4. Clicking the button takes the user to their Personal Center.'},
    **{'Test Case Owner': 'Auto'}, Environment='SIT', Status='NA',
))

# ---- Apply all NEW_TCS ----
for tc in NEW_TCS:
    tcid = tc.pop('tcid')
    add_row(tcid, **tc)

# =============================================================
# Save
# =============================================================
wb.save(DST)

# Verify
wb2 = openpyxl.load_workbook(DST, data_only=True)
ws2 = wb2['Test Cases']
new_rows = ws2.max_row - 1
ids2 = [ws2.cell(rn, id_col).value for rn in range(2, ws2.max_row+1)]
print(f'Total rows: {new_rows} (was 81)')
print(f'New TCs: {[i for i in ids2 if i and i not in row_by_id]}')
yellow_cnt = 0
for rn in range(2, ws2.max_row+1):
    for c in range(1, ws2.max_column+1):
        cell = ws2.cell(rn, c)
        f = cell.fill
        if f and f.start_color and str(f.start_color.rgb).upper().endswith('FFFF00'):
            yellow_cnt += 1
print(f'Yellow-filled cells: {yellow_cnt}')
print(f'Saved: {DST}')
