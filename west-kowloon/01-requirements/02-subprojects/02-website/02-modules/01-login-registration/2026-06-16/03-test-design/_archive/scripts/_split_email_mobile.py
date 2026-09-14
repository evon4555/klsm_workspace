"""Split 8 mode-dependent new TCs into explicit Email + Mobile variants.

Modifies AUTH-088, 090, 098, 104, 105, 106, 107, 109 to make them Email-mode
explicit, and adds NEW TCs AUTH-110 ~ 117 as the mobile counterparts.

Mobile variants: Status = NA + Comments noting mainland team can't auto-test
HK SMS; flagged for QA manual verification.
"""
import openpyxl, io, sys
from copy import copy
from openpyxl.styles import PatternFill
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX = 'test-cases-registration-login_2026-06-12.xlsx'
YELLOW = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')
MOBILE_NOTE = 'NA - QA手动测 (mainland team cannot auto-test HK SMS infra; manual verification by QA owner).'

# =============================================================
# Section A: modify 8 existing TCs to be Email-mode explicit
# =============================================================
MODIFY = {
    'SIT-TC-WEB-AUTH-088': {
        'Test Scenario': 'Successful change-password flow (email OTP)',
        'Test Case Description': 'Verify the gated change-password flow on an account whose OTP routes to EMAIL: old password -> image CAPTCHA -> email OTP -> new password (focus-out rule check) + repeat -> Submit -> success page -> auto-logout in 1-3s.',
        'Test Data': 'Account: email-only OR email+mobile (defaults to email per PRD). Current password valid; new password meets all rules.',
    },
    'SIT-TC-WEB-AUTH-090': {
        'Test Scenario': 'Change password - wrong or expired email OTP blocks submission',
        'Test Case Description': 'Verify that the Change Password flow rejects a wrong or expired EMAIL OTP — submission is blocked, the localized OTP error is shown, and the password is not changed.',
        'Test Data': 'Account routed to email OTP.',
    },
    'SIT-TC-WEB-AUTH-098': {
        'Test Scenario': 'Bind Email - OTP wrong/expired blocks completion',
        'Test Case Description': 'Verify Bind-Email OTP-error handling: a wrong OTP returns the OTP-incorrect error; an expired OTP returns the OTP-expired error; bind fails in both cases. (OTP is delivered to the email being bound.)',
        'Test Data': 'Mobile-only account; binding a fresh email; OTP sent to that fresh email.',
    },
    'SIT-TC-WEB-AUTH-104': {
        'Test Scenario': 'Order page (guest, email) - email OTP gates submission; logged-in users skip',
        'Test Case Description': 'Verify that the Order page enforces an EMAIL OTP step for guest users (gates submission); logged-in users continue to skip OTP entirely.',
        'Test Data': 'Guest user using EMAIL as the contact method; logged-in user (any).',
    },
    'SIT-TC-WEB-AUTH-105': {
        'Test Scenario': 'Order page - UNREGISTERED email shows 3-language opt-in checkbox + description (default unchecked)',
        'Test Case Description': 'Verify the unregistered-email branch of the blur-time uniqueness check: when the entered EMAIL is unregistered, the 3-language opt-in checkbox + description text appear and are unchecked by default.',
        'Test Data': 'Guest user, contact method = Email, value is a fresh unregistered email.',
    },
    'SIT-TC-WEB-AUTH-106': {
        'Test Scenario': 'Order page - REGISTERED email hides the opt-in checkbox',
        'Test Case Description': 'Verify the registered-email branch of the blur-time uniqueness check: when the entered EMAIL is already registered, NO opt-in checkbox/description is shown and the guest purchase continues normally.',
        'Test Data': 'Guest user, contact method = Email, value is a known registered email.',
    },
    'SIT-TC-WEB-AUTH-107': {
        'Test Scenario': 'Submit + payment success WITH opt-in (email branch) → upgrade; email stored, mobile empty',
        'Test Case Description': 'Verify the upgrade branch when the chosen contact method is EMAIL: opt-in checked + payment success converts the guest temp account into a registered account; the chosen EMAIL is stored as the account contact, mobile field is empty; account is tagged "guest-purchase-conversion"; the order is linked; welcome reward pack (if configured) is granted.',
        'Test Data': 'Guest user, contact = unregistered email; opt-in CHECKED.',
    },
    'SIT-TC-WEB-AUTH-109': {
        'Test Scenario': 'Payment success page (email branch) - 3-language registration-success message + "Go to My Account"',
        'Test Case Description': 'Verify the payment-success page when conversion happened via EMAIL: the 3-language registration-success message renders with the {邮箱} placeholder filled in; the "Go to My Account" button routes to Personal Center.',
        'Test Data': 'Converted user from AUTH-107 (email branch).',
    },
}

# =============================================================
# Section B: 8 new mobile-variant TCs
# =============================================================
ADD = [
    dict(
        tcid='SIT-TC-WEB-AUTH-110',
        Label='SIT/UAT',
        **{'Module/Feature': 'Change Password'},
        Priority='High', Severity='High',
        **{'Collected from': '2026-06-10 PRD §修改密码 — mobile variant of AUTH-088'},
        **{'Test Scenario': 'Successful change-password flow (mobile OTP)'},
        **{'Test Case Description': 'Verify the gated change-password flow on a mobile-only account: old password -> image CAPTCHA -> mobile OTP (SMS) -> new password (focus-out rule check) + repeat -> Submit -> success page -> auto-logout in 1-3s.'},
        Preconditions='Active registered mobile-only account with known current password.',
        **{'Test Steps': '1. Open Personal Center -> Account Security -> Change Password.\n'
                        '2. Enter current password.\n'
                        '3. Solve image CAPTCHA (动态码).\n'
                        '4. Click "Get Code"; receive OTP via SMS and enter it.\n'
                        '5. Enter new password; verify focus-out triggers inline rule validation.\n'
                        '6. Repeat new password.\n'
                        '7. Click Submit.'},
        **{'Test Data': 'Mobile-only account, current password valid; new password meets rules; HK mobile + SMS.'},
        **{'Expected Result': '1. OTP is delivered via SMS to the bound mobile.\n'
                             '2. Success page is shown with the same 3-language title/body as AUTH-088.\n'
                             '3. User is auto-logged-out after 1-3 seconds.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
    dict(
        tcid='SIT-TC-WEB-AUTH-111',
        Label='SIT',
        **{'Module/Feature': 'Change Password'},
        Priority='High', Severity='Medium',
        **{'Collected from': '2026-06-10 PRD §修改密码 step 5.b — mobile variant of AUTH-090'},
        **{'Test Scenario': 'Change password - wrong or expired mobile OTP blocks submission'},
        **{'Test Case Description': 'Verify that the Change Password flow rejects a wrong or expired MOBILE OTP — submission is blocked, the localized OTP error is shown, and the password is not changed.'},
        Preconditions='Active mobile-only account.',
        **{'Test Steps': '1. Open Change Password; pass current password + CAPTCHA.\n'
                        '2. Click Get Code; OTP sent via SMS.\n'
                        '3. Enter wrong OTP -> Submit. Repeat with expired OTP.'},
        **{'Test Data': 'Mobile-only account routed to mobile OTP.'},
        **{'Expected Result': '1. Wrong OTP / expired OTP blocks submission with localized error.\n'
                             '2. Password is NOT changed.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
    dict(
        tcid='SIT-TC-WEB-AUTH-112',
        Label='SIT',
        **{'Module/Feature': 'Bind Contact'},
        Priority='Medium', Severity='Medium',
        **{'Collected from': '2026-06-10 PRD §绑定手机/邮箱 — mobile variant of AUTH-098'},
        **{'Test Scenario': 'Bind Mobile - OTP wrong/expired blocks completion'},
        **{'Test Case Description': 'Verify Bind-Mobile OTP-error handling: a wrong OTP returns the OTP-incorrect error; an expired OTP returns the OTP-expired error; bind fails in both cases. (OTP is delivered via SMS to the mobile being bound.)'},
        Preconditions='Email-only account; binding a fresh mobile.',
        **{'Test Steps': '1. Start Bind Mobile with a fresh mobile number.\n'
                        '2. Send OTP via SMS; enter wrong OTP; submit.\n'
                        '3. Resend OTP; wait until expiry; enter expired OTP; submit.'},
        **{'Test Data': 'Email-only account; fresh HK mobile number; HK SMS provider.'},
        **{'Expected Result': '1. Wrong OTP: error "OTP incorrect" (3 languages); bind fails.\n'
                             '2. Expired OTP: error "OTP expired"; bind fails.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
    dict(
        tcid='SIT-TC-WEB-AUTH-113',
        Label='SIT/UAT',
        **{'Module/Feature': 'Guest-to-Member'},
        Priority='High', Severity='High',
        **{'Collected from': '2026-06-10 PRD §订单确认页 step 3 — mobile variant of AUTH-104'},
        **{'Test Scenario': 'Order page (guest, mobile) - mobile OTP (SMS) gates submission'},
        **{'Test Case Description': 'Verify that the Order page enforces a MOBILE OTP (SMS) step for guest users using the Mobile contact method; logged-in users skip OTP.'},
        Preconditions='Guest user reaches the order page; chooses Mobile as the contact method.',
        **{'Test Steps': '1. As GUEST: choose Mobile; observe OTP field.\n'
                        '2. Send OTP via SMS; enter wrong/correct OTP.\n'
                        '3. As LOGGED-IN: confirm no OTP field on order page.'},
        **{'Test Data': 'Guest using HK mobile.'},
        **{'Expected Result': '1. Guest mobile path: OTP via SMS gates submission.\n'
                             '2. Logged-in path: no OTP field.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
    dict(
        tcid='SIT-TC-WEB-AUTH-114',
        Label='SIT/UAT',
        **{'Module/Feature': 'Guest-to-Member'},
        Priority='High', Severity='High',
        **{'Collected from': '2026-06-10 PRD §账号校验 (unregistered branch) — mobile variant of AUTH-105'},
        **{'Test Scenario': 'Order page - UNREGISTERED mobile shows 3-language opt-in checkbox + description (default unchecked)'},
        **{'Test Case Description': 'Verify the unregistered-mobile branch of the blur-time uniqueness check: when the entered MOBILE is unregistered, the 3-language opt-in checkbox + description text appear, unchecked by default.'},
        Preconditions='Guest on order page; method = Mobile; mobile value is unregistered.',
        **{'Test Steps': '1. Pick Mobile; enter an unregistered mobile; blur.\n'
                        '2. Observe checkbox + description in each of 3 languages.\n'
                        '3. Verify default unchecked.'},
        **{'Test Data': 'Guest, fresh HK mobile.'},
        **{'Expected Result': 'Same 3-language checkbox + description as AUTH-105, but trigger is unregistered mobile.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
    dict(
        tcid='SIT-TC-WEB-AUTH-115',
        Label='SIT/UAT',
        **{'Module/Feature': 'Guest-to-Member'},
        Priority='High', Severity='Medium',
        **{'Collected from': '2026-06-10 PRD §账号校验 (registered branch) — mobile variant of AUTH-106'},
        **{'Test Scenario': 'Order page - REGISTERED mobile hides the opt-in checkbox'},
        **{'Test Case Description': 'Verify the registered-mobile branch of the blur-time uniqueness check: when the entered MOBILE is already registered, NO opt-in checkbox/description is shown.'},
        Preconditions='Guest on order page; method = Mobile; value = known registered mobile.',
        **{'Test Steps': '1. Pick Mobile; enter registered mobile; blur.\n'
                        '2. Observe no checkbox.\n'
                        '3. Continue purchase normally.'},
        **{'Test Data': 'Guest, registered mobile.'},
        **{'Expected Result': '1. No opt-in checkbox.\n'
                             '2. No description text.\n'
                             '3. Guest purchase continues normally.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
    dict(
        tcid='SIT-TC-WEB-AUTH-116',
        Label='SIT/UAT',
        **{'Module/Feature': 'Guest-to-Member'},
        Priority='High', Severity='High',
        **{'Collected from': '2026-06-10 PRD §提交订单 (upgrade branch) — mobile variant of AUTH-107'},
        **{'Test Scenario': 'Submit + payment success WITH opt-in (mobile branch) → upgrade; mobile stored, email empty'},
        **{'Test Case Description': 'Verify the upgrade branch when the chosen contact method is MOBILE: opt-in checked + payment success converts the guest into a registered account; the MOBILE is stored as the account contact, the email field is empty; account is tagged "guest-purchase-conversion"; the order is linked; welcome reward pack (if configured) is granted.'},
        Preconditions='Guest on order page; method = Mobile; unregistered mobile; checkbox CHECKED.',
        **{'Test Steps': '1. Tick the opt-in checkbox.\n'
                        '2. Submit + pay.\n'
                        '3. Inspect resulting account in admin / DB:\n'
                        '   a. Mobile stored; email empty.\n'
                        '   b. Conversion tag set; order linked.\n'
                        '4. Verify reward pack issued (if configured).'},
        **{'Test Data': 'Guest, unregistered HK mobile; opt-in CHECKED.'},
        **{'Expected Result': 'Same as AUTH-107 expectations, but with mobile stored / email empty.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
    dict(
        tcid='SIT-TC-WEB-AUTH-117',
        Label='SIT/UAT',
        **{'Module/Feature': 'Guest-to-Member'},
        Priority='High', Severity='Medium',
        **{'Collected from': '2026-06-10 PRD §支付成功页 — mobile variant of AUTH-109'},
        **{'Test Scenario': 'Payment success page (mobile branch) - 3-language registration message + "Go to My Account"'},
        **{'Test Case Description': 'Verify the payment-success page when conversion happened via MOBILE: the 3-language registration-success message renders with the {手机号} placeholder filled in; the "Go to My Account" button routes to Personal Center.'},
        Preconditions='User just converted via AUTH-116 (mobile branch).',
        **{'Test Steps': '1. After payment success, verify 3-language message + the {手机号} placeholder is the actual chosen mobile.\n'
                        '2. Click "Go to My Account" / "前往我的账号" / "前往我的帳號".'},
        **{'Test Data': 'Converted user from AUTH-116.'},
        **{'Expected Result': 'Same as AUTH-109 expectations, but {邮箱/手机} resolves to the bound mobile.'},
        **{'Test Case Owner': 'QA-manual'}, Environment='SIT', Status='NA',
        **{'Comments/Remarks': MOBILE_NOTE},
    ),
]

# =============================================================
# Execute
# =============================================================
wb = openpyxl.load_workbook(XLSX)
ws = wb['Test Cases']
hdr = {str(ws.cell(1, c).value).split('\n')[0].strip(): c for c in range(1, ws.max_column+1) if ws.cell(1,c).value}
id_col = hdr['Test Case ID']
row_by_id = {ws.cell(rn, id_col).value: rn for rn in range(2, ws.max_row+1) if ws.cell(rn, id_col).value}
REF_ROW = row_by_id.get('SIT-TC-WEB-AUTH-088')

# A. modify
modified = 0
for tcid, fields in MODIFY.items():
    rn = row_by_id.get(tcid)
    if not rn:
        print(f'WARN: {tcid} not found')
        continue
    for col_name, val in fields.items():
        c = hdr[col_name]
        ws.cell(rn, c).value = val
        ws.cell(rn, c).fill = YELLOW
        modified += 1

# B. add
added = 0
for tc in ADD:
    rn = ws.max_row + 1
    # copy formatting
    for c in range(1, ws.max_column + 1):
        ref = ws.cell(REF_ROW, c)
        new = ws.cell(rn, c)
        if ref.has_style:
            new.font = copy(ref.font)
            new.alignment = copy(ref.alignment)
            new.border = copy(ref.border)
            new.number_format = ref.number_format
    tcid = tc.pop('tcid')
    c = hdr['Test Case ID']
    ws.cell(rn, c).value = tcid
    ws.cell(rn, c).fill = YELLOW
    for col_name, val in tc.items():
        if val:
            c = hdr.get(col_name)
            if c:
                ws.cell(rn, c).value = val
                ws.cell(rn, c).fill = YELLOW
    added += 1

wb.save(XLSX)
print(f'Modified {modified} cells across 8 existing TCs; added {added} new TCs (AUTH-110 ~ 117).')

# Verify
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2['Test Cases']
total = ws2.max_row - 1
new_ids = []
for rn in range(2, ws2.max_row+1):
    t = ws2.cell(rn, id_col).value
    if t and t.startswith('SIT-TC-WEB-AUTH-1') and int(t.split('-')[-1]) >= 110:
        new_ids.append(t)
print(f'Total TC rows now: {total} (was 103)')
print(f'New TCs added: {new_ids}')
