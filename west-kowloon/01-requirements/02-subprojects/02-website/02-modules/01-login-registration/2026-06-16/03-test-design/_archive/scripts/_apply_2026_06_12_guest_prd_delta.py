"""Apply 2026-06-12 guest-to-member PRD delta to xlsx.

Only functional change vs 2026-06-10 spec is the payment-success page button:
  OLD: 前往我的账号 / 前往我的帳號 / Go to My Account  (-> Personal Center)
  NEW: 返回首页 / 返回首頁 / Back to Home              (-> Homepage)

Affects AUTH-109 (email branch) and AUTH-117 (mobile variant).
"""
import openpyxl, io, sys
from openpyxl.styles import PatternFill
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX = 'test-cases-registration-login_2026-06-12.xlsx'
YELLOW = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')

UPDATES = {
    'SIT-TC-WEB-AUTH-109': {
        'Test Scenario': 'Payment success page (email branch) - 3-language registration-success message + "Back to Home" button',
        'Test Case Description': 'Verify the payment-success page when conversion happened via EMAIL: the 3-language registration-success message renders with the {邮箱} placeholder filled in; the "Back to Home" button routes to the Homepage. (Updated 2026-06-12: button changed from "Go to My Account" -> "Back to Home"; destination Personal Center -> Homepage.)',
        'Test Steps': '1. After payment success, land on the success page.\n'
                      '2. Verify the registration-success message in the active language.\n'
                      '3. Switch to zh-HK and en; verify the corresponding messages.\n'
                      '4. Verify the {邮箱} placeholder is filled with the user-chosen email.\n'
                      '5. Click "返回首页" / "返回首頁" / "Back to Home".',
        'Expected Result': '1. Message in 3 languages per PRD:\n'
                           '   - zh-CN: 恭喜！您已成功注册成为西九文化区注册用户。本次订单已自动关联至您的账号 ...\n'
                           '   - zh-HK: 恭喜！您已成功註冊成為西九文化區註冊用戶。本次訂單已自動關聯至您的帳號 ...\n'
                           '   - en: Congratulations! You have successfully registered as a WestK user. This order has been automatically linked to your account ...\n'
                           '2. {邮箱} placeholder is replaced with the actual chosen email.\n'
                           '3. Button label per language: 返回首页 / 返回首頁 / Back to Home.\n'
                           '4. Clicking the button takes the user to the Homepage (NOT Personal Center; updated 2026-06-12).',
    },
    'SIT-TC-WEB-AUTH-117': {
        'Test Scenario': 'Payment success page (mobile branch) - 3-language registration message + "Back to Home" button',
        'Test Case Description': 'Verify the payment-success page when conversion happened via MOBILE: the 3-language registration-success message renders with the {手机号} placeholder filled in; the "Back to Home" button routes to the Homepage. (Updated 2026-06-12: button changed from "Go to My Account" -> "Back to Home".)',
        'Test Steps': '1. After payment success, verify 3-language message + the {手机号} placeholder is the actual chosen mobile.\n'
                      '2. Click "返回首页" / "返回首頁" / "Back to Home".',
        'Expected Result': 'Same as AUTH-109 expectations, but {邮箱/手机} resolves to the bound mobile; button label and destination identical (Back to Home -> Homepage).',
    },
}

wb = openpyxl.load_workbook(XLSX)
ws = wb['Test Cases']
hdr = {str(ws.cell(1, c).value).split('\n')[0].strip(): c
       for c in range(1, ws.max_column+1) if ws.cell(1,c).value}
id_col = hdr['Test Case ID']
row_by_id = {ws.cell(rn, id_col).value: rn for rn in range(2, ws.max_row+1) if ws.cell(rn, id_col).value}

modified = 0
for tcid, fields in UPDATES.items():
    rn = row_by_id.get(tcid)
    if not rn:
        print(f'WARN: {tcid} not found')
        continue
    for col_name, val in fields.items():
        c = hdr[col_name]
        ws.cell(rn, c).value = val
        ws.cell(rn, c).fill = YELLOW
        modified += 1

# Also append a Comments/Remarks note about the source PRD update
for tcid in UPDATES:
    rn = row_by_id.get(tcid)
    if not rn: continue
    c = hdr['Comments/Remarks']
    cur = ws.cell(rn, c).value or ''
    note = ' Updated per 2026-06-12 PRD (guest spec rev 1): success page button now routes to Homepage.'
    if note not in cur:
        ws.cell(rn, c).value = (cur + note).strip()
        ws.cell(rn, c).fill = YELLOW
        modified += 1

wb.save(XLSX)
print(f'Modified {modified} cells for AUTH-109 + AUTH-117 (button change).')
