"""Validate deterministic review gates for test-cases-registration-login_2026-06-12.xlsx.

This is a read-only gate script. It exits non-zero if the workbook is not ready
for requirement-review use.
"""
import io
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("test-cases-registration-login_2026-06-12.xlsx")
SHEET = "Test Cases"

MANDATORY = [
    "Test Case ID",
    "Module/Feature",
    "Priority\n(High/Medium/Low)",
    "Severity\n(High/Medium/Low)",
    "Test Scenario",
    "Test Case Description",
    "Test Steps",
    "Expected Result",
    "Status\n(Pass/ Fail)",
]
VALID_LEVELS = {"High", "Medium", "Low"}
VALID_STATUS = {"Pass", "Fail", "NA", "Not Run"}
FORBIDDEN_MARKERS = [
    "Deferred",
    "deferred",
    "TBD",
    "Open Question",
    "open question",
    "[Open]",
    "pending",
    "Pending",
    "待讨论",
    "待討論",
    "待确认",
    "待確認",
    "DEPRECATED",
    "@todo",
    "NA-deferred",
]
DELETED_IDS = {
    "SIT-TC-WEB-AUTH-036",
    "SIT-TC-WEB-AUTH-037",
    "SIT-TC-WEB-AUTH-046",
    "SIT-TC-WEB-AUTH-047",
    "SIT-TC-WEB-AUTH-064",
    "SIT-TC-WEB-AUTH-065",
    "SIT-TC-WEB-AUTH-068",
    "SIT-TC-WEB-AUTH-069",
    "SIT-TC-WEB-AUTH-070",
    "SIT-TC-WEB-AUTH-071",
    "SIT-TC-WEB-AUTH-072",
    "SIT-TC-WEB-AUTH-074",
}
GUEST_TO_MEMBER_COVERAGE = {
    "project_detail_quick_buy": "Quick Buy Without Login",
    "member_entry": "Buy tickets as member",
    "email_required": "Email mode required",
    "mobile_required": "Mobile mode required",
    "only_chosen_submitted": "only the chosen one",
    "name_optional": "contact name is OPTIONAL",
    "guest_email_otp": "guest, email",
    "guest_mobile_otp": "guest, mobile",
    "unregistered_email_optin": "UNREGISTERED email",
    "registered_email_hides_optin": "REGISTERED email",
    "unregistered_mobile_optin": "UNREGISTERED mobile",
    "registered_mobile_hides_optin": "REGISTERED mobile",
    "email_upgrade": "WITH opt-in (email",
    "email_no_upgrade": "WITHOUT opt-in (email",
    "mobile_upgrade": "WITH opt-in (mobile",
    "mobile_no_upgrade": "WITHOUT opt-in (mobile",
    "email_success_page": "Payment success page (email",
    "mobile_success_page": "Payment success page (mobile",
    "delivery_notice": "delivery notice follows selected",
}
DISCUSSION_TERMS = [
    "reward pack",
    "conversion-tag",
    "礼包",
    "模板",
    "template",
    "critical point",
    "临界",
    "刷新用户",
    "refresh user",
    "待讨论",
    "pending",
    "open question",
]


def cell(row, idx, name):
    return row[idx[name]]


def main():
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[SHEET]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    issues = []
    ids = []
    for rn, row in enumerate(rows, start=2):
        tcid = cell(row, idx, "Test Case ID")
        ids.append(tcid)

        for header in MANDATORY:
            value = cell(row, idx, header)
            if value is None or (isinstance(value, str) and not value.strip()):
                issues.append(("EMPTY_MANDATORY", rn, tcid, header, ""))

        module = cell(row, idx, "Module/Feature")
        priority = cell(row, idx, "Priority\n(High/Medium/Low)")
        severity = cell(row, idx, "Severity\n(High/Medium/Low)")
        status = cell(row, idx, "Status\n(Pass/ Fail)")
        if module in {"High", "Medium", "Low", "SIT", "SIT/UAT", "NA", "mj"}:
            issues.append(("COLUMN_SHIFT_MODULE", rn, tcid, "Module/Feature", module))
        if priority not in VALID_LEVELS:
            issues.append(("BAD_PRIORITY", rn, tcid, "Priority", priority))
        if severity not in VALID_LEVELS:
            issues.append(("BAD_SEVERITY", rn, tcid, "Severity", severity))
        if status and status not in VALID_STATUS:
            issues.append(("BAD_STATUS", rn, tcid, "Status", status))

        for ci, value in enumerate(row):
            if not isinstance(value, str):
                continue
            header = headers[ci].replace("\n", " / ")
            if value != value.strip() or "\r" in value:
                issues.append(("WHITESPACE", rn, tcid, header, value[:80]))
            for marker in FORBIDDEN_MARKERS:
                if marker in value:
                    issues.append(("FORBIDDEN_MARKER", rn, tcid, header, marker))
            for deleted_id in DELETED_IDS:
                if deleted_id in value:
                    issues.append(("DELETED_ID_REFERENCE", rn, tcid, header, deleted_id))

    counts = Counter(ids)
    for tcid, count in counts.items():
        if count > 1:
            issues.append(("DUPLICATE_ID", "-", tcid, "count", count))
    for deleted_id in DELETED_IDS:
        if deleted_id in counts:
            issues.append(("DELETED_ID_PRESENT", "-", deleted_id, "Test Case ID", deleted_id))

    gtm_rows = [row for row in rows if cell(row, idx, "Module/Feature") == "Guest-to-Member"]
    for coverage_name, term in GUEST_TO_MEMBER_COVERAGE.items():
        if not any(
            term.lower()
            in " ".join(
                str(cell(row, idx, h))
                for h in [
                    "Test Scenario",
                    "Test Case Description",
                    "Test Steps",
                    "Expected Result",
                    "Collected from",
                    "Comments/Remarks",
                ]
                if cell(row, idx, h)
            ).lower()
            for row in gtm_rows
        ):
            issues.append(("MISSING_PRD_COVERAGE", "-", coverage_name, "term", term))

    for rn, row in enumerate(rows, start=2):
        if cell(row, idx, "Module/Feature") != "Guest-to-Member":
            continue
        text = " ".join(str(v) for v in row if v)
        for term in DISCUSSION_TERMS:
            if term.lower() in text.lower():
                issues.append(("DISCUSSION_TERM_IN_GTM", rn, cell(row, idx, "Test Case ID"), "term", term))

    for rn, row in enumerate(rows, start=2):
        tcid = cell(row, idx, "Test Case ID")
        for header in ["Collected from", "Test Case Description", "Expected Result", "Comments/Remarks"]:
            value = cell(row, idx, header)
            if isinstance(value, str) and ("?" in value or "？" in value):
                issues.append(("QUESTION_MARK_AUDIT", rn, tcid, header, value[:120]))

    if issues:
        print(f"FAIL: {len(issues)} gate issue(s)")
        for issue in issues:
            print(issue)
        return 1

    status_header = "Status\n(Pass/ Fail)"
    print("PASS: deterministic review gates")
    print(f"rows={len(rows)} unique_ids={len(set(ids))} status_counts={dict(Counter(cell(row, idx, status_header) for row in rows))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
