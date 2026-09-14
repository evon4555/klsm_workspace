"""Round 2: also scrub Preconditions / Scenario / Expected fields that still
combine email/mobile in cases I claimed were split.

AUTH-103 stays mode-neutral (outcome identical), but rephrased to drop "email/mobile".
AUTH-105/106/107: drop "(or mobile)" or "email/mobile" from Preconditions (these are explicit email branches).
"""
import openpyxl, io, sys
from openpyxl.styles import PatternFill
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX = 'test-cases-registration-login_2026-06-12.xlsx'
YELLOW = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')

UPDATES = {
    'SIT-TC-WEB-AUTH-103': {
        # Outcome is identical regardless of email/mobile pick — phrase mode-neutrally.
        'Test Scenario': 'Order page: contact name is OPTIONAL and repositioned after the contact field',
        'Expected Result': '1. Contact-name field appears after the chosen contact field (email or the mobile equivalent, whichever the picker selected).\n2. No "required" marker on contact name.\n3. Order submits successfully with contact name empty.',
    },
    'SIT-TC-WEB-AUTH-105': {
        'Preconditions': 'Guest user on order page; have a known UNREGISTERED email.',
    },
    'SIT-TC-WEB-AUTH-106': {
        'Preconditions': 'Guest user on order page; have a known REGISTERED email.',
    },
    'SIT-TC-WEB-AUTH-107': {
        'Preconditions': 'Guest user on order page; contact method = Email; unregistered email; checkbox ticked.',
    },
}

wb = openpyxl.load_workbook(XLSX)
ws = wb['Test Cases']
hdr = {str(ws.cell(1, c).value).split('\n')[0].strip(): c for c in range(1, ws.max_column+1) if ws.cell(1,c).value}
id_col = hdr['Test Case ID']
row_by_id = {ws.cell(rn, id_col).value: rn for rn in range(2, ws.max_row+1) if ws.cell(rn, id_col).value}

modified = 0
for tcid, fields in UPDATES.items():
    rn = row_by_id.get(tcid)
    if not rn:
        print(f'WARN: {tcid} not found')
        continue
    for col_name, val in fields.items():
        c = hdr[col_name]
        ws.cell(rn, c).value = val
        ws.cell(rn, c).fill = YELLOW
        modified += 1

wb.save(XLSX)
print(f'Round-2 cleanup: {modified} cells')

# Re-scan
import re
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2['Test Cases']
pats = [re.compile(r'email\s*[/、／]\s*mobile', re.I), re.compile(r'email\s*[/、／]\s*phone', re.I),
        re.compile(r'邮箱\s*[/、／]\s*手机', re.I), re.compile(r'mobile\s*[/、／]\s*email', re.I),
        re.compile(r'\(or mobile\)', re.I), re.compile(r'Email\s*/\s*Mobile', re.I)]
print('\nRemaining combined cases:')
for rn in range(2, ws2.max_row+1):
    tcid = ws2.cell(rn, id_col).value
    if not tcid: continue
    blob = ''
    for col in ['Test Scenario','Test Case Description','Preconditions','Test Steps','Test Data','Expected Result']:
        v = ws2.cell(rn, hdr[col]).value
        if v: blob += str(v) + ' || '
    hits = set()
    for p in pats:
        for m in p.finditer(blob): hits.add(m.group())
    if hits:
        print(f'  {tcid}: {hits}')
