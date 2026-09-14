"""Export a fixed execution-result snapshot from dashboard.db.

The dashboard remains the live view and dashboard.db remains the factual
source. This tool writes a package-local audit snapshot under
05-execution/04-execution-results so execution review can cite a stable
document instead of a mutable dashboard page.

Usage:
  python export_execution_results.py --run-id 175 --package <package-dir> \
    --scope batch-session-configuration --project "standard product"
"""
from __future__ import annotations

import argparse
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

THIS = Path(__file__).resolve().parent
REPO_ROOT = THIS.parents[1]
WORKSPACE_ROOT = REPO_ROOT.parent
DEFAULT_DB = REPO_ROOT / "02-platform" / "02-dashboard" / "01-backend" / "dashboard.db"
CASE_ID_RE = re.compile(r"\b(SIT-TC-[A-Z0-9-]+-\d{3,})\b", re.IGNORECASE)


def norm(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def case_id_from_text(*values: object) -> str | None:
    for value in values:
        match = CASE_ID_RE.search(str(value or ""))
        if match:
            return match.group(1).upper()
    return None


def markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        cells = [str(value if value is not None else "").replace("|", r"\|") for value in row]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def read_case_metadata(xlsx: Path) -> dict[str, dict[str, str]]:
    if not xlsx.is_file():
        return {}
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Test Cases"] if "Test Cases" in wb.sheetnames else wb.active
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(row=1, column=col).value
        key = re.sub(r"\s+", " ", str(raw or "").strip().lower())
        if key:
            headers[key] = col

    def col(*names: str) -> int | None:
        for name in names:
            key = re.sub(r"\s+", " ", name.strip().lower())
            for header, idx in headers.items():
                if header.startswith(key):
                    return idx
        return None

    id_col = col("Test Case ID")
    if not id_col:
        return {}
    scenario_col = col("Test Scenario")
    desc_col = col("Test Case Description")
    automation_col = col("Automation Type")
    module_col = col("Module/Feature")
    out: dict[str, dict[str, str]] = {}
    for row in range(2, ws.max_row + 1):
        case_id = case_id_from_text(ws.cell(row=row, column=id_col).value)
        if not case_id:
            continue

        def value(idx: int | None) -> str:
            return norm(ws.cell(row=row, column=idx).value) if idx else ""

        out[case_id] = {
            "case_id": case_id,
            "test_scenario": value(scenario_col),
            "test_case_description": value(desc_col),
            "module_feature": value(module_col),
            "automation_type": value(automation_col),
        }
    return out


def merge_automation_metadata(base: dict[str, dict[str, str]], automation_xlsx: Path | None) -> None:
    if not automation_xlsx or not automation_xlsx.is_file():
        return
    auto = read_case_metadata(automation_xlsx)
    for case_id, row in auto.items():
        if case_id not in base:
            base[case_id] = row
        elif row.get("automation_type"):
            base[case_id]["automation_type"] = row["automation_type"]


def feature_locations(project_root: Path) -> dict[str, dict[str, Any]]:
    features_root = project_root / "02-automation" / "01-features"
    automation_root = project_root / "02-automation"
    if not features_root.exists():
        return {}
    out: dict[str, dict[str, Any]] = {}
    for path in sorted(features_root.rglob("*.feature")):
        if "deprecated" in path.parts:
            continue
        try:
            text = path.read_text(encoding="utf-8")
        except Exception:
            text = path.read_text(encoding="gbk", errors="replace")
        pending_tags: list[str] = []
        feature = ""
        for line_no, line in enumerate(text.splitlines(), start=1):
            stripped = line.strip()
            if stripped.startswith("Feature:"):
                feature = stripped.removeprefix("Feature:").strip()
                continue
            if stripped.startswith("@"):
                pending_tags = re.findall(r"@[\w-]+", stripped)
                continue
            if not stripped.startswith("Scenario"):
                continue
            scenario = re.sub(r"^Scenario(?: Outline)?:\s*", "", stripped)
            for match in CASE_ID_RE.findall(" ".join([scenario, *pending_tags])):
                case_id = match.upper()
                out.setdefault(case_id, {
                    "file": str(path.relative_to(automation_root)).replace("\\", "/"),
                    "absolute_file": str(path).replace("\\", "/"),
                    "line": line_no,
                    "feature": feature,
                    "scenario": scenario,
                    "test_kind": "Behave Scenario",
                    "test_name": scenario,
                    "tags": pending_tags,
                })
            pending_tags = []
    return out


def screenshot_manifest(project_root: Path, artifact_scope: str, run_id: int) -> tuple[Path | None, dict[str, dict[str, Any]]]:
    manifest_path = (
        project_root / "02-automation" / "07-artifacts" / artifact_scope /
        f"case-screenshot-manifest-run-{run_id}.json"
    )
    if not manifest_path.is_file():
        return None, {}
    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    cases = data.get("cases") or {}
    out: dict[str, dict[str, Any]] = {}
    if not isinstance(cases, dict):
        return manifest_path, out
    for case_id, row in cases.items():
        if not isinstance(row, dict):
            continue
        screenshots = row.get("screenshots") or []
        if not isinstance(screenshots, list):
            screenshots = []
        out[str(case_id).upper()] = {
            "screenshot_count": len(screenshots),
            "screenshots": [
                shot for shot in screenshots
                if isinstance(shot, dict)
            ],
        }
    return manifest_path, out


def load_run(db_path: Path, run_id: int) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    con = sqlite3.connect(str(db_path))
    con.row_factory = sqlite3.Row
    try:
        run_row = con.execute(
            "SELECT id, started_at, finished_at, status, env, project_key, tags, "
            "total, passed, failed, errored, skipped FROM test_runs WHERE id = ?",
            (run_id,),
        ).fetchone()
        if not run_row:
            raise SystemExit(f"dashboard run not found: {run_id}")
        scenario_rows = con.execute(
            "SELECT id, run_id, feature, name, status, duration_s, error_msg, tags "
            "FROM test_scenarios WHERE run_id = ? ORDER BY id",
            (run_id,),
        ).fetchall()
    finally:
        con.close()
    return dict(run_row), [dict(row) for row in scenario_rows]


def build_snapshot(args: argparse.Namespace) -> dict[str, Any]:
    package = args.package.resolve()
    project_root = args.project_root or package.parents[5]
    project_root = project_root.resolve()
    out_dir = args.out_dir or package / "05-execution" / "04-execution-results"
    xlsx = args.xlsx or package / "03-test-design" / f"test-cases-{args.scope}.xlsx"
    automation_xlsx = args.automation_xlsx or (
        package / "05-execution" / "02-automation-assessment" /
        f"automation-assessment-{args.scope}.xlsx"
    )

    run, scenarios = load_run(args.db, args.run_id)
    case_meta = read_case_metadata(xlsx)
    merge_automation_metadata(case_meta, automation_xlsx)
    locations = feature_locations(project_root)
    manifest_path, screenshot_meta = screenshot_manifest(project_root, args.artifact_scope, args.run_id)

    enriched = []
    for item in scenarios:
        case_id = case_id_from_text(item.get("name"), item.get("tags"))
        meta = case_meta.get(case_id or "", {})
        loc = locations.get(case_id or "")
        shot_meta = screenshot_meta.get(case_id or "", {})
        enriched.append({
            **item,
            "case_id": case_id,
            "test_case_description": meta.get("test_case_description") or meta.get("test_scenario") or "",
            "automation_type": meta.get("automation_type") or "",
            "automation_location": loc,
            "screenshot_count": shot_meta.get("screenshot_count", 0),
            "screenshots": shot_meta.get("screenshots", []),
            "error_msg": item.get("error_msg") or "",
        })

    dashboard_url = args.dashboard_url.rstrip("/")
    api_url = args.api_url.rstrip("/")
    artifact_root = args.artifact_root or project_root / "02-automation" / "07-artifacts" / args.artifact_scope
    snapshot = {
        "exported_at": now_iso(),
        "package": str(package).replace("\\", "/"),
        "scope": args.scope,
        "project": args.project,
        "dashboard_run_id": args.run_id,
        "dashboard_url": f"{dashboard_url}/?page=test-run&project={args.project.replace(' ', '%20')}",
        "run_api_url": f"{api_url}/api/runs/{args.run_id}",
        "dashboard_db": str(args.db).replace("\\", "/"),
        "artifact_root": str(artifact_root).replace("\\", "/"),
        "runtime_screenshot_manifest": str(manifest_path).replace("\\", "/") if manifest_path else None,
        "source_test_case_workbook": str(xlsx).replace("\\", "/"),
        "automation_assessment_workbook": str(automation_xlsx).replace("\\", "/") if automation_xlsx else None,
        "manual_scope": args.manual_scope,
        "run": run,
        "summary": {
            "total": int(run.get("total") or len(enriched)),
            "passed": int(run.get("passed") or 0),
            "failed": int(run.get("failed") or 0),
            "errored": int(run.get("errored") or 0),
            "skipped": int(run.get("skipped") or 0),
        },
        "scenarios": enriched,
    }
    out_dir.mkdir(parents=True, exist_ok=True)
    return snapshot


def render_md(snapshot: dict[str, Any]) -> str:
    s = snapshot["summary"]
    run = snapshot["run"]
    scenario_rows = []
    for item in snapshot["scenarios"]:
        loc = item.get("automation_location") or {}
        location = ""
        if loc:
            location = str(loc.get("file") or "")
        scenario_rows.append([
            item.get("case_id") or "",
            item.get("test_case_description") or "",
            item.get("automation_type") or "",
            item.get("status") or "",
            f"{float(item.get('duration_s') or 0):.2f}s",
            item.get("screenshot_count", 0),
            location,
        ])

    pass_rate = "0.0%"
    total = s["total"]
    if total:
        pass_rate = f"{(s['passed'] / total) * 100:.1f}%"

    return "\n".join([
        f"# Execution Results - {snapshot['scope']}",
        "",
        "## Summary",
        "",
        markdown_table(
            ["Item", "Value"],
            [
                ["Exported At", snapshot["exported_at"]],
                ["Dashboard Run", f"`{snapshot['dashboard_run_id']}`"],
                ["Project", f"`{snapshot['project']}`"],
                ["Environment", f"`{run.get('env')}`"],
                ["Run Started", run.get("started_at")],
                ["Run Finished", run.get("finished_at")],
                ["Result", f"`{s['passed']} passed / {s['failed']} failed / {s['errored']} errored / {s['skipped']} skipped`"],
                ["Pass Rate", pass_rate],
                ["Manual Scope", snapshot["manual_scope"]],
            ],
        ),
        "",
        "## Trace Links",
        "",
        markdown_table(
            ["Item", "Location"],
            [
                ["Dashboard Page", snapshot["dashboard_url"]],
                ["Run JSON API", snapshot["run_api_url"]],
                ["SQLite Source", snapshot["dashboard_db"]],
                ["Artifact Root", snapshot["artifact_root"]],
                ["Runtime Screenshot Manifest", snapshot.get("runtime_screenshot_manifest") or ""],
                ["Test Case Workbook", snapshot["source_test_case_workbook"]],
                ["Automation Assessment Workbook", snapshot["automation_assessment_workbook"] or ""],
            ],
        ),
        "",
        "## Scenario Results",
        "",
        markdown_table(
            ["Case ID", "Test Case Description", "Automation Type", "Status", "Duration", "Screenshots", "Automation Location"],
            scenario_rows,
        ),
        "",
        "## Audit Note",
        "",
        "- `dashboard.db` is the execution fact source for live dashboard history.",
        "- Runtime screenshot availability is sourced from the per-case screenshot manifest for this run.",
        "- This file is a fixed package-local snapshot for execution review and audit.",
        "- Re-export this document only when a new execution run supersedes the current run.",
        "",
    ])


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--run-id", type=int, required=True)
    parser.add_argument("--package", type=Path, required=True)
    parser.add_argument("--scope", required=True)
    parser.add_argument("--project", default="west-kowloon")
    parser.add_argument("--project-root", type=Path, default=None)
    parser.add_argument("--xlsx", type=Path, default=None)
    parser.add_argument("--automation-xlsx", type=Path, default=None)
    parser.add_argument("--out-dir", type=Path, default=None)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--dashboard-url", default="http://127.0.0.1:5173")
    parser.add_argument("--api-url", default="http://127.0.0.1:8002")
    parser.add_argument("--artifact-root", type=Path, default=None)
    parser.add_argument("--artifact-scope", default="batch_session_configuration")
    parser.add_argument("--manual-scope", default="None")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    snapshot = build_snapshot(args)
    out_dir = args.out_dir or args.package / "05-execution" / "04-execution-results"
    json_path = out_dir / f"execution-results-{args.scope}.json"
    md_path = out_dir / f"execution-results-{args.scope}.md"
    json_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    md_path.write_text(render_md(snapshot), encoding="utf-8")
    print(f"[execution-results] wrote {md_path}")
    print(f"[execution-results] wrote {json_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
