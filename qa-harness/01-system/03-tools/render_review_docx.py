"""Render human-facing review markdown into a polished Word document.

This is intentionally separate from md_docx.py. md_docx.py is a faithful
handoff converter; this renderer is for external-share-ready review/sign-off
forms where layout quality matters.

Usage:
  python render_review_docx.py <input.md> [-o output.docx]
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "1F4E79"
BLUE_DARK = "17365D"
BLUE_LIGHT = "D9EAF7"
GREEN_LIGHT = "E2F0D9"
GRAY_LIGHT = "F7F9FB"
GRAY_MID = "D9E2F3"
WHITE = "FFFFFF"
BLACK = "1F2933"
MUTED = "5B677A"
RED_LIGHT = "FCE4D6"

LATIN_FONT = "Calibri"
EA_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, mono: bool = False) -> None:
    font_name = MONO_FONT if mono else LATIN_FONT
    ea_name = MONO_FONT if mono else EA_FONT
    run.font.name = font_name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key, value in (
        ("w:ascii", font_name),
        ("w:hAnsi", font_name),
        ("w:eastAsia", ea_name),
        ("w:cs", font_name),
    ):
        r_fonts.set(qn(key), value)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "C9D3DF", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, bottom: int = 90,
                     left: int = 120, right: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (
        ("top", top),
        ("left", left),
        ("bottom", bottom),
        ("right", right),
    ):
        elem = tc_mar.find(qn(f"w:{side}"))
        if elem is None:
            elem = OxmlElement(f"w:{side}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")


def set_table_width_pct(table, pct: int = 5000) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), str(pct))


def add_cell_text(cell, text: str, header: bool = False, size: float = 9.0) -> None:
    cell.text = ""
    lines = text.split("\n") if text else [""]
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    for idx, line in enumerate(lines):
        if idx:
            paragraph.add_run().add_break()
        add_inline_runs(
            paragraph,
            line,
            size=size,
            bold=header,
            color=WHITE if header else BLACK,
        )
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_inline_runs(paragraph, text: str, size: float = 10.0,
                    bold: bool = False, color: str = BLACK) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size - 0.5, bold=bold, color=color, mono=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color)


def normalize_cell(text: str) -> str:
    text = text.strip()
    text = text.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    text = re.sub(r"^\*\*(.*)\*\*$", r"\1", text)
    return text


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        raw = lines[idx].strip()
        cells = [normalize_cell(c) for c in raw.strip("|").split("|")]
        is_separator = all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in cells)
        if not is_separator:
            rows.append(cells)
        idx += 1
    return rows, idx


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    for style_name, size, bold, color in (
        ("Normal", 10, False, BLACK),
        ("Heading 1", 18, True, BLUE_DARK),
        ("Heading 2", 12, True, BLUE_DARK),
        ("Heading 3", 10.5, True, BLUE_DARK),
    ):
        style = styles[style_name]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:ascii"), LATIN_FONT)
        r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
        r_fonts.set(qn("w:eastAsia"), EA_FONT)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("QA Review")
    set_run_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(f"Generated review hand-off | {date.today().isoformat()}")
    set_run_font(run, size=8, color=MUTED)


def add_title_banner(doc: Document, title: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width_pct(table)
    cell = table.cell(0, 0)
    shade_cell(cell, BLUE_DARK)
    set_cell_borders(cell, BLUE_DARK, "0")
    set_cell_margins(cell, top=190, bottom=190, left=180, right=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_run_font(run, size=17, bold=True, color=WHITE)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, color=BLUE_DARK)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(p, text, size=9.8, color=BLACK)


def row_fill(row: list[str], row_idx: int) -> str:
    joined = " ".join(row)
    if row_idx == 0:
        return BLUE
    first = row[0].strip() if row else ""
    value = row[1].strip() if len(row) > 1 else ""
    value_lower = value.lower()
    if first in {"Sign-off Date", "Signature / Confirmation"} and (
        not value
        or "yyyy" in value_lower
        or "<type name or sign here>" in value_lower
        or "sign here" in value_lower
    ):
        return BLUE_LIGHT
    if "Final Status" in joined and "Signed Off" in joined:
        return GREEN_LIGHT
    if "Handoff Decision" in joined and re.search(r"\bProceed\b|Signed Off", joined):
        return GREEN_LIGHT
    # Historical "Needs Revision" rows often remain in the trail after a later
    # row closes them. Do not keep those rows red in the final hand-off; reserve
    # warning color for currently open blocking states.
    if re.search(r"\b(Open Blocker|Blocked|Current Blocker)\b", joined):
        return RED_LIGHT
    return GRAY_LIGHT if row_idx % 2 == 0 else WHITE


def add_review_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_width_pct(table)

    for r_idx, row_data in enumerate(rows):
        fill = row_fill(row_data, r_idx)
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            text = row_data[c_idx] if c_idx < len(row_data) else ""
            shade_cell(cell, fill)
            set_cell_borders(cell)
            set_cell_margins(cell)
            add_cell_text(cell, text, header=(r_idx == 0), size=8.8)
            if c_idx == 0 and r_idx != 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render(md_path: Path, docx_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    if re.search(r"(?im)^\s*<table\b", text):
        raise SystemExit(f"raw HTML table is not allowed: {md_path}")

    lines = text.splitlines()
    doc = Document()
    configure_document(doc)
    add_header_footer(doc)

    idx = 0
    title_done = False
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            idx += 1
            continue
        if line.startswith("# "):
            title = line[2:].strip()
            if not title_done:
                add_title_banner(doc, title)
                title_done = True
            else:
                add_section_heading(doc, title)
            idx += 1
            continue
        if line.startswith("## "):
            add_section_heading(doc, line[3:].strip())
            idx += 1
            continue
        if line.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_review_table(doc, rows)
            continue
        add_body_paragraph(doc, line)
        idx += 1

    doc.save(docx_path)


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path, default=None)
    args = parser.parse_args()

    output = args.output or args.input.with_suffix(".docx")
    render(args.input, output)
    print(f"[OK] rendered review docx: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
