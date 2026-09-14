"""Record test-case add/modify/remove decisions in human-visible audit trails.

This writer updates the artifacts that reviewers actually open:

- 03-test-design/test-cases-<scope>.xlsx -> Audit Trail sheet
- 04-test-case-review/test-case-review-<scope>.md/.docx -> Review Trail row
- 05-execution/02-automation-assessment/automation-assessment-<scope>.xlsx
  and automation-assessment-review-<scope>.md/.docx when they exist

It does not edit the test-case rows for you. Make the case add/modify/remove
first, then run this tool to leave the audit trail. If --old-xlsx is supplied,
modified cells and added case ID cells are highlighted yellow with comments.

Usage:
  python record_testcase_change.py --package <package-dir> --scope <scope> \
    --change-type remove --case-id SIT-TC-STD-CONFIG-015 \
    --actor "Test Manager / QA1 (AI)" \
    --reason "Test Manager confirmed this case is unnecessary." \
    --scope-impact "Current scope is 14 cases." \
    --evidence "03-test-design/CHANGE.md"
"""
from __future__ import annotations

import argparse
import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

THIS = Path(__file__).resolve().parent
sys.path.insert(0, str(THIS))
from _validator_common import CASE_ID_PREFIX_RE, cell_str, extract_case_id, find_columns, reconfigure_stdout  # noqa: E402
from diff_xlsx_versions import SIGNIFICANT_COLS, diff  # noqa: E402


AUDIT_HEADERS = [
    "Date",
    "Actor",
    "Change Type",
    "Decision / Reason",
    "Affected Case(s)",
    "Current Scope Impact",
    "Evidence / Linked Artifact",
]
AUDIT_SHEET = "Audit Trail"
TESTCASE_PREFIX = "test-cases-"
AUTOMATION_PREFIX = "automation-assessment-"
REVIEW_PREFIX = "test-case-review-"
AUTOMATION_REVIEW_PREFIX = "automation-assessment-review-"
YELLOW = PatternFill(start_color="FFFFFACD", end_color="FFFFFACD", fill_type="solid")
HEADER_FILL = PatternFill(start_color="FF1F4E79", end_color="FF1F4E79", fill_type="solid")


def norm(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def infer_scope(package: Path) -> str:
    files = [
        p for p in (package / "03-test-design").glob(TESTCASE_PREFIX + "*.xlsx")
        if p.is_file() and not p.name.startswith("~$")
    ]
    if len(files) != 1:
        raise SystemExit(
            "Scope is required when the package has zero or multiple "
            f"test-case workbooks: {[p.name for p in files]}"
        )
    return files[0].stem.removeprefix(TESTCASE_PREFIX)


def collect_case_ids(args: argparse.Namespace) -> list[str]:
    ids: list[str] = []
    for raw in args.case_id:
        for match in CASE_ID_PREFIX_RE.findall(raw):
            if match not in ids:
                ids.append(match)
    if args.old_xlsx:
        report = diff(args.old_xlsx, args.current_xlsx)
        for cid in report["added"]:
            if cid not in ids:
                ids.append(cid)
        for cid in report["removed"]:
            if cid not in ids:
                ids.append(cid)
        for item in report["modified"]:
            cid = item["id"]
            if cid not in ids:
                ids.append(cid)
    return ids


def display_change_type(raw: str) -> str:
    value = raw.strip().lower().replace("_", "-")
    mapping = {
        "add": "Case added",
        "added": "Case added",
        "modify": "Case modified",
        "modified": "Case modified",
        "update": "Case modified",
        "updated": "Case modified",
        "remove": "Case removed",
        "removed": "Case removed",
        "delete": "Case removed",
        "deleted": "Case removed",
        "scope-adjustment": "Post sign-off scope adjustment",
        "post-signoff-scope-adjustment": "Post sign-off scope adjustment",
        "post-sign-off-scope-adjustment": "Post sign-off scope adjustment",
        "mixed": "Case scope change",
    }
    return mapping.get(value, raw.strip())


def ensure_audit_sheet(wb):
    if AUDIT_SHEET in wb.sheetnames:
        ws = wb[AUDIT_SHEET]
    else:
        ws = wb.create_sheet(AUDIT_SHEET)

    for col, header in enumerate(AUDIT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col)
        if not cell.value:
            cell.value = header
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    widths = [14, 28, 28, 54, 34, 54, 60]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    return ws


def row_values(row: list[str]) -> list[str]:
    return [norm(v) for v in row]


def append_audit_row(xlsx_path: Path, row: list[str], dry_run: bool = False) -> bool:
    wb = load_workbook(xlsx_path)
    ws = ensure_audit_sheet(wb)
    desired = row_values(row)
    for ridx in range(2, ws.max_row + 1):
        existing = [norm(ws.cell(ridx, c).value) for c in range(1, len(AUDIT_HEADERS) + 1)]
        if existing == desired:
            print(f"[audit] skip duplicate row in {xlsx_path.name}")
            return False

    if dry_run:
        print(f"[audit] would append row to {xlsx_path}")
        return True

    target = ws.max_row + 1
    for col, value in enumerate(row, start=1):
        cell = ws.cell(row=target, column=col, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(xlsx_path)
    print(f"[audit] appended row to {xlsx_path}")
    return True


def case_row_index(ws, id_col: int) -> dict[str, int]:
    rows: dict[str, int] = {}
    for row in range(2, ws.max_row + 1):
        cid = extract_case_id(cell_str(ws, row, id_col))
        if cid:
            rows[cid] = row
    return rows


def add_comment(cell, text: str, actor: str) -> None:
    existing = cell.comment.text if cell.comment else ""
    content = text if not existing else f"{existing}\n\n{text}"
    cell.comment = Comment(content[:32000], actor[:54] or "QA")


def apply_diff_markers(current_xlsx: Path, old_xlsx: Path, reason: str,
                       actor: str, change_date: str, dry_run: bool = False) -> None:
    report = diff(old_xlsx, current_xlsx)
    wb = load_workbook(current_xlsx)
    ws = wb["Test Cases"] if "Test Cases" in wb.sheetnames else wb.active
    cols, _missing = find_columns(ws)
    if "id" not in cols:
        raise SystemExit(f"{current_xlsx}: cannot highlight changes without Test Case ID column")
    rows = case_row_index(ws, cols["id"])
    comment = f"{change_date} {actor}: {reason}"

    marked = 0
    for cid in report["added"]:
        row = rows.get(cid)
        if row:
            cell = ws.cell(row=row, column=cols["id"])
            cell.fill = YELLOW
            add_comment(cell, "Added case. " + comment, actor)
            marked += 1

    for item in report["modified"]:
        cid = item["id"]
        row = rows.get(cid)
        if not row:
            continue
        for col_name in item["changes"]:
            if col_name not in SIGNIFICANT_COLS or col_name not in cols:
                continue
            cell = ws.cell(row=row, column=cols[col_name])
            cell.fill = YELLOW
            add_comment(cell, f"Modified {col_name}. {comment}", actor)
            marked += 1

    if dry_run:
        print(f"[audit] would mark {marked} changed cell(s) in {current_xlsx.name}")
        return
    if marked:
        wb.save(current_xlsx)
    print(f"[audit] marked {marked} changed cell(s) in {current_xlsx.name}")


def escape_md_cell(value: str) -> str:
    return norm(value).replace("|", r"\|")


def append_review_trail(md_path: Path, row: list[str], dry_run: bool = False) -> bool:
    if not md_path.is_file():
        print(f"[review] skip missing review md: {md_path}")
        return False
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    heading = None
    for idx, line in enumerate(lines):
        if line.strip().lower() == "## 4. review trail and version record":
            heading = idx
            break
    if heading is None:
        raise SystemExit(f"{md_path}: missing '## 4. Review Trail and Version Record'")

    table_start = None
    for idx in range(heading + 1, len(lines)):
        stripped = lines[idx].strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            table_start = idx
            break
        if stripped.startswith("## "):
            break
    if table_start is None:
        raise SystemExit(f"{md_path}: missing Review Trail markdown table")

    table_end = table_start
    while table_end < len(lines) and lines[table_end].strip().startswith("|"):
        table_end += 1

    row_line = "| " + " | ".join(escape_md_cell(v) for v in row) + " |"
    if any(norm(line) == norm(row_line) for line in lines[table_start:table_end]):
        print(f"[review] skip duplicate trail row in {md_path.name}")
        return False
    if dry_run:
        print(f"[review] would append trail row to {md_path}")
        return True

    lines.insert(table_end, row_line)
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"[review] appended trail row to {md_path}")
    return True


def render_docx(md_path: Path, dry_run: bool = False) -> None:
    if not md_path.is_file():
        return
    if dry_run:
        print(f"[review] would render DOCX for {md_path}")
        return
    proc = subprocess.run([sys.executable, str(THIS / "render_review_docx.py"), str(md_path)])
    if proc.returncode != 0:
        raise SystemExit(proc.returncode)


def build_targets(package: Path, scope: str, include_automation: bool) -> tuple[list[Path], list[Path]]:
    xlsx_targets = [package / "03-test-design" / f"{TESTCASE_PREFIX}{scope}.xlsx"]
    review_targets = [package / "04-test-case-review" / f"{REVIEW_PREFIX}{scope}.md"]

    automation_xlsx = package / "05-execution" / "02-automation-assessment" / f"{AUTOMATION_PREFIX}{scope}.xlsx"
    automation_review = package / "05-execution" / "02-automation-assessment" / f"{AUTOMATION_REVIEW_PREFIX}{scope}.md"
    if include_automation and automation_xlsx.is_file():
        xlsx_targets.append(automation_xlsx)
    if include_automation and automation_review.is_file():
        review_targets.append(automation_review)
    return xlsx_targets, review_targets


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--package", type=Path, required=True, help="requirement package directory")
    parser.add_argument("--scope", default="", help="scope suffix, e.g. batch-session-configuration")
    parser.add_argument("--change-type", required=True,
                        help="add/modify/remove/scope-adjustment/mixed or a display label")
    parser.add_argument("--case-id", action="append", default=[],
                        help="affected case ID; may be repeated")
    parser.add_argument("--actor", default="QA1 (AI)")
    parser.add_argument("--date", default=date.today().isoformat())
    parser.add_argument("--reason", required=True)
    parser.add_argument("--scope-impact", required=True)
    parser.add_argument("--evidence", action="append", default=[],
                        help="linked artifact or evidence reference; may be repeated")
    parser.add_argument("--round", default="", help="Review Trail round/version label")
    parser.add_argument("--status", default="Closed")
    parser.add_argument("--closure-evidence", default="")
    parser.add_argument("--old-xlsx", type=Path, default=None,
                        help="previous test-case xlsx; enables yellow diff markers")
    parser.add_argument("--current-xlsx", type=Path, default=None,
                        help="current xlsx for --old-xlsx diff; default is package test-case workbook")
    parser.add_argument("--no-automation-assessment", action="store_true",
                        help="do not update paired automation assessment artifacts")
    parser.add_argument("--no-render-docx", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args()


def main() -> int:
    reconfigure_stdout()
    args = parse_args()
    package = args.package.resolve()
    scope = args.scope or infer_scope(package)
    test_xlsx = package / "03-test-design" / f"{TESTCASE_PREFIX}{scope}.xlsx"
    args.current_xlsx = args.current_xlsx or test_xlsx
    if not test_xlsx.is_file():
        raise SystemExit(f"test-case workbook not found: {test_xlsx}")

    case_ids = collect_case_ids(args)
    affected = "; ".join(case_ids) if case_ids else "All affected cases"
    change_type = display_change_type(args.change_type)
    evidence = "; ".join(args.evidence) if args.evidence else "Not specified"
    closure = args.closure_evidence or (
        f"Audit Trail updated in workbook(s); {args.scope_impact}"
    )
    round_label = args.round or change_type

    xlsx_targets, review_targets = build_targets(
        package, scope, include_automation=not args.no_automation_assessment
    )
    audit_row = [
        args.date,
        args.actor,
        change_type,
        args.reason,
        affected,
        args.scope_impact,
        evidence,
    ]
    review_row = [
        round_label,
        args.actor,
        args.reason,
        affected,
        closure,
        args.status,
    ]

    for target in xlsx_targets:
        if not target.is_file():
            print(f"[audit] skip missing workbook: {target}")
            continue
        append_audit_row(target, audit_row, dry_run=args.dry_run)

    if args.old_xlsx:
        apply_diff_markers(args.current_xlsx, args.old_xlsx, args.reason, args.actor, args.date,
                           dry_run=args.dry_run)

    for md_path in review_targets:
        changed = append_review_trail(md_path, review_row, dry_run=args.dry_run)
        if changed and not args.no_render_docx:
            render_docx(md_path, dry_run=args.dry_run)

    print("[audit] done")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        import traceback
        traceback.print_exc()
        raise SystemExit(2)
