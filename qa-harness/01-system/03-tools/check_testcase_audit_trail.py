"""Check human-visible audit trail consistency for test-case changes.

This validator protects the path the reviewer actually sees:

- xlsx workbook Audit Trail sheet
- 04-test-case-review Review Trail and rendered DOCX
- paired automation assessment workbook/review when present

It does not try to infer every possible unrecorded edit. It enforces that once
a test-case add/modify/remove/scope-adjustment is recorded in either the review
document or the workbook, the matching human-visible artifacts are kept in sync.

Exit:
  0  PASS
  1  FAIL
  2  usage/setup error
"""
from __future__ import annotations

import argparse
import glob
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

THIS = Path(__file__).resolve().parent
sys.path.insert(0, str(THIS))
from _paths import repo_root  # noqa: E402
from _validator_common import CASE_ID_PREFIX_RE, cell_str, expand_paths, extract_case_id, find_columns, reconfigure_stdout  # noqa: E402
from diff_xlsx_versions import SIGNIFICANT_COLS  # noqa: E402


REPO = repo_root()
WORKSPACE = REPO.parent
AUDIT_SHEET = "Audit Trail"
AUDIT_HEADERS = [
    "Date",
    "Actor",
    "Change Type",
    "Decision / Reason",
    "Affected Case(s)",
    "Current Scope Impact",
    "Evidence / Linked Artifact",
]
TESTCASE_PREFIX = "test-cases-"
REVIEW_PREFIX = "test-case-review-"
AUTOMATION_PREFIX = "automation-assessment-"
AUTOMATION_REVIEW_PREFIX = "automation-assessment-review-"
CHANGE_KEYWORDS = (
    "post sign-off",
    "post-sign-off",
    "scope adjustment",
    "case added",
    "case modified",
    "case removed",
    "case deleted",
    "added case",
    "modified case",
    "removed case",
    "deleted case",
    "audit trail",
)
YELLOW_MARKERS = ("FFFACD", "FFF2CC")


@dataclass
class AuditCheck:
    xlsx: Path
    kind: str
    review_md: Path | None
    failures: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def relpath(path: Path | None) -> str:
    if path is None:
        return "<missing>"
    try:
        return str(path.relative_to(WORKSPACE))
    except ValueError:
        return str(path)


def norm(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def norm_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def case_ids(text: str) -> set[str]:
    return set(CASE_ID_PREFIX_RE.findall(text or ""))


def scope_from_xlsx(path: Path) -> str:
    if path.stem.startswith(TESTCASE_PREFIX):
        return path.stem.removeprefix(TESTCASE_PREFIX)
    if path.stem.startswith(AUTOMATION_PREFIX):
        return path.stem.removeprefix(AUTOMATION_PREFIX)
    return path.stem


def package_from_xlsx(path: Path) -> Path | None:
    for parent in path.parents:
        if parent.name == "03-test-design":
            return parent.parent
        if parent.name == "02-automation-assessment":
            execution = parent.parent
            if execution.name == "05-execution":
                return execution.parent
    return None


def review_for_xlsx(path: Path) -> tuple[str, Path | None]:
    package = package_from_xlsx(path)
    scope = scope_from_xlsx(path)
    if package is None:
        return "unknown", None
    if path.stem.startswith(TESTCASE_PREFIX):
        return "test-case", package / "04-test-case-review" / f"{REVIEW_PREFIX}{scope}.md"
    if path.stem.startswith(AUTOMATION_PREFIX):
        return "automation-assessment", (
            package / "05-execution" / "02-automation-assessment" /
            f"{AUTOMATION_REVIEW_PREFIX}{scope}.md"
        )
    return "unknown", None


def paired_automation_xlsx(testcase_xlsx: Path) -> Path | None:
    package = package_from_xlsx(testcase_xlsx)
    if package is None:
        return None
    scope = scope_from_xlsx(testcase_xlsx)
    path = package / "05-execution" / "02-automation-assessment" / f"{AUTOMATION_PREFIX}{scope}.xlsx"
    return path if path.is_file() else None


def expand_xlsx(patterns: list[str], include_automation: bool) -> list[Path]:
    if patterns:
        files = expand_paths(patterns)
    else:
        files = [Path(p) for p in glob.glob(
            str(WORKSPACE / "*" / "01-requirements" / "**" / "test-cases-*.xlsx"),
            recursive=True,
        )]
    out: list[Path] = []
    seen: set[str] = set()
    for path in files:
        if path.suffix.lower() != ".xlsx" or path.name.startswith("~$"):
            continue
        key = str(path.resolve()).lower()
        if key not in seen:
            seen.add(key)
            out.append(path.resolve())
        if include_automation and path.stem.startswith(TESTCASE_PREFIX):
            auto = paired_automation_xlsx(path)
            if auto:
                auto_key = str(auto.resolve()).lower()
                if auto_key not in seen:
                    seen.add(auto_key)
                    out.append(auto.resolve())
    return sorted(out)


def load_audit_rows(xlsx: Path) -> tuple[bool, list[dict[str, str]], list[str]]:
    wb = load_workbook(xlsx, data_only=True)
    if AUDIT_SHEET not in wb.sheetnames:
        return False, [], []
    ws = wb[AUDIT_SHEET]
    header_map: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header_map[norm_key(ws.cell(row=1, column=col).value)] = col
    missing = [header for header in AUDIT_HEADERS if norm_key(header) not in header_map]
    rows: list[dict[str, str]] = []
    for row in range(2, ws.max_row + 1):
        values = {header: norm(ws.cell(row=row, column=header_map.get(norm_key(header), 0)).value)
                  for header in AUDIT_HEADERS if norm_key(header) in header_map}
        if any(values.values()):
            rows.append(values)
    return True, rows, missing


def workbook_case_count(xlsx: Path) -> int:
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Test Cases"] if "Test Cases" in wb.sheetnames else wb.active
    cols, _missing = find_columns(ws)
    if "id" not in cols:
        return 0
    count = 0
    for row in range(2, ws.max_row + 1):
        if extract_case_id(cell_str(ws, row, cols["id"])):
            count += 1
    return count


def review_text(md_path: Path | None) -> str:
    if md_path and md_path.is_file():
        return md_path.read_text(encoding="utf-8", errors="replace")
    return ""


def docx_text(docx_path: Path) -> str:
    doc = Document(docx_path)
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def parse_review_summary_count(text: str, kind: str) -> int | None:
    names = ["Number of Test Cases"] if kind == "test-case" else ["Source Test Case Count"]
    for name in names:
        pattern = r"\|\s*" + re.escape(name) + r"\s*\|\s*(\d+)\s*\|"
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None


def split_table_row(line: str) -> list[str]:
    return [norm(cell.replace(r"\|", "|")) for cell in line.strip().strip("|").split("|")]


def is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def review_trail_change_rows(text: str) -> list[str]:
    lines = text.splitlines()
    heading = None
    for idx, line in enumerate(lines):
        if line.strip().lower() == "## 4. review trail and version record":
            heading = idx
            break
    if heading is None:
        return []

    table_start = None
    for idx in range(heading + 1, len(lines)):
        stripped = lines[idx].strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            table_start = idx
            break
        if stripped.startswith("## "):
            return []
    if table_start is None:
        return []

    rows: list[list[str]] = []
    idx = table_start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = split_table_row(lines[idx])
        if not is_separator_row(cells):
            rows.append(cells)
        idx += 1
    if len(rows) <= 1:
        return []
    changed: list[str] = []
    for row in rows[1:]:
        row_text = " | ".join(row)
        low = row_text.lower()
        if any(keyword in low for keyword in CHANGE_KEYWORDS):
            changed.append(row_text)
    return changed


def has_yellow_or_comment_marker(xlsx: Path, cid: str) -> bool:
    wb = load_workbook(xlsx, data_only=False)
    ws = wb["Test Cases"] if "Test Cases" in wb.sheetnames else wb.active
    cols, _missing = find_columns(ws)
    if "id" not in cols:
        return False
    target_row = None
    for row in range(2, ws.max_row + 1):
        if extract_case_id(cell_str(ws, row, cols["id"])) == cid:
            target_row = row
            break
    if target_row is None:
        return False
    for col_name in ("id", *SIGNIFICANT_COLS):
        col = cols.get(col_name)
        if not col:
            continue
        cell = ws.cell(row=target_row, column=col)
        rgb = str(getattr(cell.fill.fgColor, "rgb", "") or "").upper()
        if any(marker in rgb for marker in YELLOW_MARKERS):
            return True
        if cell.comment and norm(cell.comment.text):
            return True
    return False


def validate_one(xlsx: Path, strict_yellow: bool = False) -> AuditCheck:
    kind, review_md = review_for_xlsx(xlsx)
    result = AuditCheck(xlsx=xlsx, kind=kind, review_md=review_md)

    has_audit, audit_rows, missing_headers = load_audit_rows(xlsx)
    text = review_text(review_md)
    trail_change_rows = review_trail_change_rows(text)

    if missing_headers:
        result.failures.append(f"Audit Trail sheet is missing header(s): {missing_headers}")

    if trail_change_rows and not has_audit:
        result.failures.append(
            "Review Trail records a test-case change, but workbook has no Audit Trail sheet"
        )

    if has_audit and not audit_rows:
        result.failures.append("Audit Trail sheet exists but has no audit rows")

    if audit_rows and (review_md is None or not review_md.is_file()):
        result.failures.append("Workbook Audit Trail exists, but paired review markdown is missing")

    if kind in {"test-case", "automation-assessment"} and text:
        expected_count = parse_review_summary_count(text, kind)
        actual_count = workbook_case_count(xlsx)
        if expected_count is not None and expected_count != actual_count:
            result.failures.append(
                f"review summary count ({expected_count}) does not match workbook case count ({actual_count})"
            )

    audit_text = "\n".join(" | ".join(row.values()) for row in audit_rows)
    audit_ids = case_ids(audit_text)
    trail_text = "\n".join(trail_change_rows)
    trail_ids = case_ids(trail_text)

    if trail_ids and not audit_ids.issuperset(trail_ids):
        missing = sorted(trail_ids - audit_ids)
        result.failures.append(
            f"Review Trail changed case ID(s) missing from workbook Audit Trail: {missing}"
        )

    if audit_ids and text:
        missing_from_review = sorted(cid for cid in audit_ids if cid not in text)
        if missing_from_review:
            result.failures.append(
                f"Audit Trail case ID(s) missing from review markdown: {missing_from_review}"
            )

    if audit_rows:
        for idx, row in enumerate(audit_rows, start=2):
            for header in AUDIT_HEADERS:
                if not row.get(header):
                    result.failures.append(f"Audit Trail row {idx} has blank '{header}'")
            if not case_ids(" ".join(row.values())) and "all" not in row.get("Affected Case(s)", "").lower():
                result.failures.append(f"Audit Trail row {idx} has no recognizable affected case ID")

        if review_md and review_md.is_file():
            docx_path = review_md.with_suffix(".docx")
            if not docx_path.is_file():
                result.failures.append("paired review DOCX is missing")
            else:
                docx = docx_text(docx_path)
                missing_from_docx = sorted(cid for cid in audit_ids if cid not in docx)
                if missing_from_docx:
                    result.failures.append(
                        f"Audit Trail case ID(s) missing from review DOCX: {missing_from_docx}"
                    )
                audit_keywords = sorted({
                    kw for kw in CHANGE_KEYWORDS
                    if kw in audit_text.lower() and kw not in {"audit trail"}
                })
                for keyword in audit_keywords:
                    if keyword not in docx.lower():
                        result.failures.append(
                            f"audit keyword {keyword!r} is missing from review DOCX"
                        )

    if strict_yellow:
        for row in audit_rows:
            low = " ".join(row.values()).lower()
            is_modify = ("modified" in low or "case modified" in low) and "removed" not in low
            if is_modify:
                for cid in case_ids(" ".join(row.values())):
                    if not has_yellow_or_comment_marker(xlsx, cid):
                        result.failures.append(
                            f"modified case {cid} has no yellow/comment marker in workbook"
                        )
    else:
        for row in audit_rows:
            low = " ".join(row.values()).lower()
            is_modify = ("modified" in low or "case modified" in low) and "removed" not in low
            if is_modify:
                for cid in case_ids(" ".join(row.values())):
                    if not has_yellow_or_comment_marker(xlsx, cid):
                        result.warnings.append(
                            f"modified case {cid} has no yellow/comment marker in workbook"
                        )

    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--xlsx", nargs="+", default=[],
                        help="xlsx file(s) or glob(s); default scans workspace test-cases")
    parser.add_argument("--no-automation-assessment", action="store_true",
                        help="do not auto-check paired automation assessment workbooks")
    parser.add_argument("--strict-yellow", action="store_true",
                        help="fail modified cases that lack yellow/comment markers")
    return parser.parse_args()


def main() -> int:
    reconfigure_stdout()
    args = parse_args()
    files = expand_xlsx(args.xlsx, include_automation=not args.no_automation_assessment)
    print(f"[testcase-audit-trail] scanning {len(files)} workbook(s)")
    if not files:
        print("RESULT: PASS  checked=0  failures=0")
        return 0

    results = [validate_one(path, strict_yellow=args.strict_yellow) for path in files]
    failures: list[str] = []
    warnings: list[str] = []
    for item in results:
        status = "PASS" if not item.failures else "FAIL"
        print(f"  [{status}] {relpath(item.xlsx)} review={relpath(item.review_md)}")
        for warning in item.warnings:
            warnings.append(f"{relpath(item.xlsx)}: {warning}")
        for failure in item.failures:
            failures.append(f"{relpath(item.xlsx)}: {failure}")

    if warnings:
        print(f"\n[testcase-audit-trail] {len(warnings)} warning(s):")
        for warning in warnings:
            print(f"  WARN  {warning}")

    if failures:
        print(f"\n[testcase-audit-trail] {len(failures)} failure(s):")
        for failure in failures:
            print(f"  FAIL  {failure}")
        print("\nRESULT: FAIL - test-case audit trail is inconsistent.")
        return 1

    print("\nRESULT: PASS - test-case audit trail is consistent.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        import traceback
        traceback.print_exc()
        raise SystemExit(2)
