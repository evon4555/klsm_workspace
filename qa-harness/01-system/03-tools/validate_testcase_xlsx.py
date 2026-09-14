"""validate_testcase_xlsx.py — Phase-1 validator #1.

Structural check of one or more test-cases-*.xlsx files. Validators key on
canonical semantic column names (discovered from the header row), so files
may include extra columns (e.g. 西九's 'Label') without breaking validation.

Usage:
  python validate_testcase_xlsx.py <xlsx-or-glob> [...]

Exit:  0 PASS,  1 FAIL,  2 usage/crash.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from _validator_common import (
    CANONICAL_COLS, CASE_ID_RE, EXEC_REQUIRED_WHEN_STATUS, REQUIRED_NON_EMPTY,
    VALID_PRIORITY, VALID_STATUS, cell_str, emit_result, expand_paths,
    find_columns, is_yellow, reconfigure_stdout,
)


def check_one(path: Path) -> list[str]:
    failures: list[str] = []
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        return [f"cannot open: {e!r}"]

    if "Test Cases" not in wb.sheetnames:
        return [f"missing sheet 'Test Cases' (found: {wb.sheetnames})"]
    ws = wb["Test Cases"]

    cols, missing = find_columns(ws)
    if missing:
        return [f"missing required header column(s): {missing}"]

    seen_ids: dict[str, int] = {}
    id_col = cols["id"]
    for r in range(2, ws.max_row + 1):
        cid = cell_str(ws, r, id_col)
        if not cid:
            other_filled = any(
                cell_str(ws, r, c)
                for c in cols.values() if c != id_col
            )
            if other_filled:
                failures.append(f"row {r}: blank ID but other cells filled")
            continue

        if not CASE_ID_RE.match(cid):
            failures.append(f"row {r}: id {cid!r} does not match SIT-TC-PROJ-MOD-NNN")
            continue
        if cid in seen_ids:
            failures.append(
                f"row {r}: duplicate id {cid} (first seen at row {seen_ids[cid]})"
            )
        else:
            seen_ids[cid] = r

        for key in REQUIRED_NON_EMPTY:
            if key == "id":
                continue
            if not cell_str(ws, r, cols[key]):
                failures.append(f"row {r} ({cid}): required '{key}' is empty")

        pr = cell_str(ws, r, cols["priority"])
        if pr and pr not in VALID_PRIORITY:
            failures.append(
                f"row {r} ({cid}): priority {pr!r} not in {sorted(VALID_PRIORITY-{''})}"
            )

        st = cell_str(ws, r, cols["status"])
        if st not in VALID_STATUS:
            failures.append(
                f"row {r} ({cid}): status {st!r} not in {sorted(VALID_STATUS-{''})}"
            )

        if st and st != "Not Run":
            for key in EXEC_REQUIRED_WHEN_STATUS:
                if not cell_str(ws, r, cols[key]):
                    failures.append(
                        f"row {r} ({cid}): status={st} but execution col '{key}' is empty"
                    )

        sample_cells = [
            ws.cell(row=r, column=cols["id"]),
            ws.cell(row=r, column=cols["module"]),
            ws.cell(row=r, column=cols["expected"]),
        ]
        if all(is_yellow(c) for c in sample_cells) and st not in ("NA", "Deferred"):
            failures.append(
                f"row {r} ({cid}): whole-row yellow but status={st!r} (expected NA or Deferred)"
            )

    return failures


def main(argv: list[str]) -> int:
    reconfigure_stdout()
    if len(argv) < 2:
        print(__doc__)
        return 2

    paths = expand_paths(argv[1:])
    if not paths:
        print(f"no xlsx matched: {argv[1:]}")
        return 2

    total_failures = 0
    for p in paths:
        print(f"=== {p.name}")
        fails = check_one(p)
        if not fails:
            print("  OK")
        else:
            for f in fails:
                print(f"  FAIL: {f}")
            total_failures += len(fails)

    return emit_result(checked=len(paths), failures=total_failures)


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv))
    except Exception as e:
        import traceback
        traceback.print_exc()
        sys.exit(2)
