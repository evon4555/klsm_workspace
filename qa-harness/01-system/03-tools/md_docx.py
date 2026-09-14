"""md_docx.py — bidirectional md ↔ docx converter for human-review hand-off.

Convention (see 01-system/12-md-docx-handoff.md):

  - `.md` is the canonical AI-readable form (lives in git, AI reads + writes).
  - `.docx` is a human-review derivative (generated from `.md` when the AI
    hands an artifact to a human for review).
  - When the user finishes reviewing, they hand the edited `.docx` back; the
    AI runs `to-md` to absorb edits back into the canonical `.md`.

Backed by pandoc (installed via `pypandoc-binary` — pandoc binary is
bundled inside the qa-harness venv, no system install needed).

Usage:

    python md_docx.py to-docx   <input.md>   [-o output.docx]
    python md_docx.py to-md     <input.docx> [-o output.md]
    python md_docx.py round-trip <input.md>            # fidelity self-test

Exit:
  0  success
  2  on usage error / file missing
  3  pandoc failure
"""
from __future__ import annotations

import argparse
import difflib
import sys
import tempfile
from pathlib import Path

try:
    import pypandoc
except ImportError:
    print(
        "ERROR: pypandoc not installed. Install with:\n"
        "  pip install pypandoc-binary python-docx",
        file=sys.stderr,
    )
    sys.exit(2)

# Force UTF-8 stdout so PowerShell on a gbk codepage doesn't choke on
# non-ASCII characters in printed paths or diff content.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


# Pandoc format choices: gfm (GitHub-flavored markdown) is the best pivot
# format for round-tripping tables, which the consolidation / review docs
# rely on heavily. --wrap=none keeps long lines intact instead of pandoc
# auto-wrapping them at column 72 (which would balloon the diff).
MD_FORMAT = "gfm"
DOCX_FORMAT = "docx"
MD_OUTPUT_EXTRA = ["--wrap=none"]


def md_to_docx(md_path: Path, docx_path: Path | None = None,
               reference: Path | None = None, tighten: bool = True,
               allow_raw_html_tables: bool = False) -> Path:
    """Convert markdown → docx. Returns output path.

    If tighten=True (default), post-process the docx to shrink fonts,
    set 微软雅黑 as the East-Asian font, compact paragraph spacing, and
    shrink table cell padding. Use --no-tighten to keep pandoc defaults
    (which are too loose / too large for review docs).
    """
    md_path = Path(md_path)
    if not md_path.exists():
        print(f"ERROR: input not found: {md_path}", file=sys.stderr)
        sys.exit(2)
    docx_path = Path(docx_path) if docx_path else md_path.with_suffix(".docx")

    extra_args: list[str] = []
    if reference:
        extra_args += ["--reference-doc", str(reference)]

    # Pre-process: pandoc's gfm reader silently DROPS <br> / <br/> inside
    # table cells (cells flatten to one line). We swap them out for a rare
    # placeholder character before pandoc, then post-process the docx to
    # turn placeholders into real w:br elements. This preserves the source
    # author's intent of "this cell has multiple lines / list items".
    BR_PLACEHOLDER = "␤"  # ␤  SYMBOL FOR NEWLINE  (rare in any real content)
    import re as _re
    import tempfile as _tempfile
    src_text = md_path.read_text(encoding="utf-8")
    if not allow_raw_html_tables and _re.search(r"(?im)^\s*<table\b", src_text):
        print(
            "ERROR: raw HTML <table> block found in markdown. "
            "Use markdown table syntax before generating docx; raw HTML "
            "tables render poorly or may be dropped by pandoc. "
            "Pass --allow-raw-html-tables only for legacy recovery.",
            file=sys.stderr,
        )
        print(f"ERROR: source: {md_path}", file=sys.stderr)
        sys.exit(2)
    src_processed = _re.sub(r"<br\s*/?>", BR_PLACEHOLDER, src_text)
    used_placeholder = BR_PLACEHOLDER in src_processed

    tmp_md: Path | None = None
    try:
        if used_placeholder:
            tmp = _tempfile.NamedTemporaryFile(
                mode="w", suffix=".md", delete=False, encoding="utf-8"
            )
            tmp.write(src_processed)
            tmp.close()
            tmp_md = Path(tmp.name)
            pandoc_input = str(tmp_md)
        else:
            pandoc_input = str(md_path)

        try:
            pypandoc.convert_file(
                pandoc_input,
                DOCX_FORMAT,
                format=MD_FORMAT,
                outputfile=str(docx_path),
                extra_args=extra_args,
            )
        except RuntimeError as e:
            print(f"ERROR: pandoc md->docx failed: {e}", file=sys.stderr)
            sys.exit(3)
    finally:
        if tmp_md is not None:
            tmp_md.unlink(missing_ok=True)

    if used_placeholder:
        try:
            _replace_br_placeholders(docx_path, BR_PLACEHOLDER)
        except Exception as e:
            print(f"[warn] br-placeholder restore failed: {e}", file=sys.stderr)

    if tighten:
        try:
            _tighten_docx_styles(docx_path)
        except Exception as e:
            print(f"[warn] tighten styles failed: {e}", file=sys.stderr)

    print(f"[OK] md -> docx: {docx_path}")
    return docx_path


def _replace_br_placeholders(docx_path: Path, placeholder: str) -> None:
    """Walk all paragraphs (in body + table cells) and convert occurrences
    of `placeholder` inside a run into actual `<w:br/>` line breaks.
    Pandoc's gfm reader drops <br> in table cells, so we route the
    line breaks through a placeholder character that pandoc preserves
    as plain text, then upgrade it here to real Word soft breaks.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))

    def _process_paragraph(para):
        for run in list(para.runs):
            if placeholder not in run.text:
                continue
            parts = run.text.split(placeholder)
            run.text = parts[0]
            for part in parts[1:]:
                br = OxmlElement("w:br")
                run._r.append(br)
                if part:
                    t = OxmlElement("w:t")
                    t.text = part
                    t.set(qn("xml:space"), "preserve")
                    run._r.append(t)

    for p in doc.paragraphs:
        _process_paragraph(p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _process_paragraph(p)

    doc.save(str(docx_path))


def _tighten_docx_styles(docx_path: Path,
                         body_pt: int = 10,
                         h1_pt: int = 14, h2_pt: int = 12, h3_pt: int = 11,
                         h4_pt: int = 10, h5_pt: int = 10,
                         table_pt: int = 9, code_pt: int = 9,
                         eastasia_font: str = "微软雅黑",
                         latin_font: str = "Calibri",
                         mono_font: str = "Consolas") -> None:
    """Post-process a pandoc-generated docx to make it readable for review:
       - shrink Normal + Heading font sizes
       - set 微软雅黑 as the East-Asian font (Chinese renders cleanly)
       - compact paragraph spacing (0pt before / 3pt after, 1.15 line)
       - shrink table cell margins from Word's ~115 to 40/80
    Idempotent — safe to re-run.
    """
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def _set_run_fonts(rPr, latin: str, ea: str) -> None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:ascii"), latin)
        rFonts.set(qn("w:hAnsi"), latin)
        rFonts.set(qn("w:eastAsia"), ea)
        rFonts.set(qn("w:cs"), latin)

    def _set_cell_margins(cell, top: int, bottom: int, left: int, right: int) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = tcPr.find(qn("w:tcMar"))
        if tcMar is None:
            tcMar = OxmlElement("w:tcMar")
            tcPr.append(tcMar)
        for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            elem = tcMar.find(qn(f"w:{side}"))
            if elem is None:
                elem = OxmlElement(f"w:{side}")
                tcMar.append(elem)
            elem.set(qn("w:w"), str(value))
            elem.set(qn("w:type"), "dxa")

    def _shade_cell(cell, fill: str) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tcPr.append(shd)
        shd.set(qn("w:fill"), fill)

    doc = Document(str(docx_path))

    style_targets = {
        "Normal":        (body_pt,    latin_font, eastasia_font, False),
        "Heading 1":     (h1_pt,      latin_font, eastasia_font, True),
        "Heading 2":     (h2_pt,      latin_font, eastasia_font, True),
        "Heading 3":     (h3_pt,      latin_font, eastasia_font, True),
        "Heading 4":     (h4_pt,      latin_font, eastasia_font, True),
        "Heading 5":     (h5_pt,      latin_font, eastasia_font, True),
        "Heading 6":     (h5_pt,      latin_font, eastasia_font, True),
        "Title":         (h1_pt + 2,  latin_font, eastasia_font, True),
        "Source Code":   (code_pt,    mono_font,  mono_font, False),
        "Verbatim Char": (code_pt,    mono_font,  mono_font, False),
    }
    for name, (pt, latin, ea, bold) in style_targets.items():
        try:
            s = doc.styles[name]
        except KeyError:
            continue
        try:
            s.font.size = Pt(pt)
            s.font.name = latin
            if bold:
                s.font.bold = True
            rPr = s.element.get_or_add_rPr()
            _set_run_fonts(rPr, latin, ea)
        except Exception:
            pass

    for p in doc.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.15
        for run in p.runs:
            if run.font.size is None:
                run.font.size = Pt(body_pt)
            rPr = run._r.get_or_add_rPr()
            _set_run_fonts(rPr, latin_font, eastasia_font)

    for tbl in doc.tables:
        try:
            tbl.style = "Table Grid"
        except Exception:
            pass
        for row_idx, row in enumerate(tbl.rows):
            is_header = row_idx == 0
            for cell in row.cells:
                _set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
                if is_header:
                    _shade_cell(cell, "1F4E79")
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    pf.line_spacing = 1.1
                    for run in p.runs:
                        run.font.size = Pt(table_pt)
                        if is_header:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        rPr = run._r.get_or_add_rPr()
                        _set_run_fonts(rPr, latin_font, eastasia_font)

    doc.save(str(docx_path))


def docx_to_md(docx_path: Path, md_path: Path | None = None) -> Path:
    """Convert docx → markdown. Returns output path."""
    docx_path = Path(docx_path)
    if not docx_path.exists():
        print(f"ERROR: input not found: {docx_path}", file=sys.stderr)
        sys.exit(2)
    md_path = Path(md_path) if md_path else docx_path.with_suffix(".md")

    try:
        pypandoc.convert_file(
            str(docx_path),
            MD_FORMAT,
            format=DOCX_FORMAT,
            outputfile=str(md_path),
            extra_args=MD_OUTPUT_EXTRA,
        )
    except RuntimeError as e:
        print(f"ERROR: pandoc docx->md failed: {e}", file=sys.stderr)
        sys.exit(3)

    # Pandoc emits a UTF-8 BOM at the start of the file; strip it.
    # Why: a BOM in a .md file silently breaks downstream tools that
    # assume UTF-8 without BOM (see memory: codex BOM auth 2026-05-28).
    raw = md_path.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        md_path.write_bytes(raw[3:])

    print(f"[OK] docx -> md: {md_path}")
    return md_path


def round_trip(md_path: Path, show_lines: int = 60) -> int:
    """Run md -> docx -> md' and print a unified diff vs the original.

    Round-trip is rarely byte-identical (pandoc may renormalize tables,
    list markers, blank lines). The check is informational: as long as
    the semantic content survives, the hand-off is safe.
    """
    md_path = Path(md_path)
    if not md_path.exists():
        print(f"ERROR: input not found: {md_path}", file=sys.stderr)
        return 2

    src = md_path.read_text(encoding="utf-8")

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        docx = tmp / (md_path.stem + ".docx")
        md2 = tmp / (md_path.stem + "_rt.md")
        md_to_docx(md_path, docx)
        docx_to_md(docx, md2)
        rt = md2.read_text(encoding="utf-8")

    if src == rt:
        print("[OK] round-trip byte-identical")
        return 0

    diff_lines = list(
        difflib.unified_diff(
            src.splitlines(),
            rt.splitlines(),
            fromfile=str(md_path),
            tofile="round-tripped",
            n=2,
            lineterm="",
        )
    )
    delta = sum(
        1 for d in diff_lines
        if d.startswith(("+", "-")) and not d.startswith(("+++", "---"))
    )

    src_lines = len(src.splitlines())
    pct = (delta / max(src_lines, 1)) * 100
    print(
        f"[INFO] round-trip differs: {delta} changed lines "
        f"out of {src_lines} source lines ({pct:.1f}%)."
    )
    print(f"[INFO] showing first {show_lines} diff lines:\n")
    for line in diff_lines[:show_lines]:
        print(line)
    if len(diff_lines) > show_lines:
        print(f"... ({len(diff_lines) - show_lines} more diff lines omitted)")

    # Round-trip drift is expected for non-trivial docs; informational only.
    return 0


def main() -> int:
    p = argparse.ArgumentParser(
        description="md ↔ docx converter for AI / human review hand-off.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    sub = p.add_subparsers(dest="cmd", required=True)

    p_to_docx = sub.add_parser("to-docx", help="markdown → docx")
    p_to_docx.add_argument("input", type=Path)
    p_to_docx.add_argument("-o", "--output", type=Path, default=None)
    p_to_docx.add_argument(
        "--reference",
        type=Path,
        default=None,
        help="Optional reference docx for styling (passed to pandoc).",
    )
    p_to_docx.add_argument(
        "--no-tighten",
        action="store_true",
        help="Skip the post-process style tightening (keep pandoc defaults).",
    )
    p_to_docx.add_argument(
        "--allow-raw-html-tables",
        action="store_true",
        help="Allow legacy raw HTML <table> blocks. Default is to fail fast.",
    )

    p_to_md = sub.add_parser("to-md", help="docx → markdown")
    p_to_md.add_argument("input", type=Path)
    p_to_md.add_argument("-o", "--output", type=Path, default=None)

    p_rt = sub.add_parser("round-trip", help="md -> docx -> md' fidelity self-test")
    p_rt.add_argument("input", type=Path)
    p_rt.add_argument(
        "--show-lines",
        type=int,
        default=60,
        help="Max diff lines to print (default 60).",
    )

    args = p.parse_args()

    if args.cmd == "to-docx":
        md_to_docx(args.input, args.output, args.reference,
                   tighten=not args.no_tighten,
                   allow_raw_html_tables=args.allow_raw_html_tables)
        return 0
    if args.cmd == "to-md":
        docx_to_md(args.input, args.output)
        return 0
    if args.cmd == "round-trip":
        return round_trip(args.input, args.show_lines)

    p.print_help()
    return 2


if __name__ == "__main__":
    sys.exit(main())
