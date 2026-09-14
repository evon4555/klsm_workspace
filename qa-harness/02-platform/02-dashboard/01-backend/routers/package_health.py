"""Package Health filesystem and execution-promotion routes."""

import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

from database import TestRun, TestScenario, get_db


router = APIRouter()


def configure_router(
    *,
    package_scanner_module,
    workspace_root,
    repo_root,
    project_key,
    project_workspace_root,
    case_id_from_scenario_fields,
):
    global package_scanner, _workspace_root, _repo_root
    global _project_key, _project_workspace_root, _case_id_from_scenario_fields
    package_scanner = package_scanner_module
    _workspace_root = workspace_root
    _repo_root = repo_root
    _project_key = project_key
    _project_workspace_root = project_workspace_root
    _case_id_from_scenario_fields = case_id_from_scenario_fields
    return router


# ---------------------------------------------------------------------------
# Quality System — Package Health (Phase 1)
#
# Reads the live filesystem state of the selected project's requirement
# packages and returns per-stage health. Filesystem IS the data source; no DB
# tables are involved.
# ---------------------------------------------------------------------------


@router.get("/api/quality-system/packages")
def list_quality_packages(project: str = "west-kowloon"):
    """Return a compact list of all packages under <project>/01-requirements
    with a per-stage summary. Recomputed on every request (manual refresh)."""
    if package_scanner is None:
        return {"error": "package_scanner not available", "packages": []}
    results = package_scanner.scan_workspace(_workspace_root, project=project)
    return {
        "project": project,
        "scanned_at": datetime.now(timezone.utc).isoformat(),
        "package_count": len(results),
        "stage_defs": package_scanner.STAGE_DEFS,
        "packages": [package_scanner.to_jsonable(r) for r in results],
    }


# ---- Phase 2: file actions (open / generate docx / run validator) ---------
#
# All three endpoints share the same security model:
#   - input path is a workspace-relative path with forward slashes
#   - resolve against _workspace_root and reject anything that escapes it
#   - reject if file/dir doesn't exist
#
# These endpoints have side effects (open file in OS, write new docx, run
# subprocess), so we use POST not GET.

import os as _os
import subprocess as _subprocess

# Force child Python processes to write UTF-8 to stdout/stderr so we can
# decode their output cleanly when it contains Chinese (validator filenames,
# md_docx.py paths, etc.). Otherwise the Windows console codec (gbk by
# default) mangles non-gbk characters before we ever see them.
_CHILD_ENV = {**_os.environ, "PYTHONIOENCODING": "utf-8"}


class _PathPayload(BaseModel):
    path: str  # workspace-relative path with forward slashes


def _resolve_safe(path: str) -> Optional[Path]:
    """Resolve a workspace-relative path; return None if it escapes the
    workspace or doesn't exist."""
    target = (_workspace_root / path).resolve()
    try:
        target.relative_to(_workspace_root)
    except ValueError:
        return None
    if not target.exists():
        return None
    return target


@router.post("/api/quality-system/open-file")
def open_file(payload: _PathPayload):
    """Open a file in the OS default application (Word/WPS for .docx,
    Notepad/VS Code for .md, etc.). Windows uses os.startfile.
    """
    target = _resolve_safe(payload.path)
    if target is None:
        return {"ok": False, "error": "path not found or outside workspace"}
    if not target.is_file():
        return {"ok": False, "error": "path is not a file"}
    try:
        if hasattr(_os, "startfile"):
            _os.startfile(str(target))  # Windows native
        else:
            # Fallback for non-Windows (xdg-open / open)
            opener = "open" if sys.platform == "darwin" else "xdg-open"
            _subprocess.Popen([opener, str(target)])
        return {"ok": True, "opened": str(target)}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@router.post("/api/quality-system/generate-docx")
def generate_docx(payload: _PathPayload):
    """Run md_docx.py to-docx on a workspace-relative .md path. Returns
    the generated .docx path on success."""
    target = _resolve_safe(payload.path)
    if target is None:
        return {"ok": False, "error": "path not found or outside workspace"}
    if target.suffix.lower() != ".md":
        return {"ok": False, "error": "path is not a .md file"}
    tool = _repo_root / "01-system" / "03-tools" / "md_docx.py"
    out_docx = target.with_suffix(".docx")
    try:
        result = _subprocess.run(
            [sys.executable, str(tool), "to-docx", str(target)],
            capture_output=True, text=True, timeout=60,
            encoding="utf-8", errors="replace", env=_CHILD_ENV,
        )
        if result.returncode != 0:
            return {
                "ok": False,
                "error": "md_docx.py failed",
                "stdout": result.stdout[-500:],
                "stderr": result.stderr[-500:],
            }
        rel = str(out_docx.relative_to(_workspace_root)).replace("\\", "/")
        return {"ok": True, "docx_path": rel, "abs_path": str(out_docx)}
    except _subprocess.TimeoutExpired:
        return {"ok": False, "error": "md_docx.py timed out"}
    except Exception as e:
        return {"ok": False, "error": str(e)}


def _extract_case_ids(xlsx_path: Path) -> set[str]:
    """Read the Test Case ID column from a test-cases-*.xlsx and return the
    set of valid case IDs. The column position is discovered from the header
    row — West Kowloon's template has a leading 'Label' column that pushes
    'Test Case ID' to column B, so a fixed column index would miss it.
    Silently returns empty set on any failure — non-fatal."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
        if "Test Cases" not in wb.sheetnames:
            return set()
        ws = wb["Test Cases"]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        case_col = None
        for i, h in enumerate(header):
            if h and isinstance(h, str) and "Test Case ID" in h:
                case_col = i  # 0-based
                break
        if case_col is None:
            return set()
        ids: set[str] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if case_col >= len(row):
                continue
            v = row[case_col]
            if v and isinstance(v, str):
                v = v.strip()
                # Match e.g. SIT-TC-WEB-AUTH-001 or SIT-TC-AMW-017
                if re.match(r"^[A-Z]+(-[A-Z]+)+-\d+$", v):
                    ids.add(v)
        return ids
    except Exception:
        return set()


def _execution_scope_from_package(pkg: Path) -> str:
    test_design = pkg / "03-test-design"
    if test_design.is_dir():
        for xlsx in sorted(test_design.glob("test-cases-*.xlsx")):
            return xlsx.stem.removeprefix("test-cases-")
    result_dir = pkg / "05-execution" / "04-execution-results"
    if result_dir.is_dir():
        for js in sorted(result_dir.glob("execution-results-*.json")):
            return js.stem.removeprefix("execution-results-")
    return pkg.name


def _artifact_scope_from_snapshot(snapshot: dict | None, scope: str) -> str:
    if snapshot:
        root = str(snapshot.get("artifact_root") or "").replace("\\", "/").rstrip("/")
        if root:
            return root.rsplit("/", 1)[-1]
    return scope.replace("-", "_")


def _manual_scope_from_snapshot(snapshot: dict | None) -> str:
    if snapshot and snapshot.get("manual_scope"):
        return str(snapshot["manual_scope"])
    return "See 05-execution/05-manual-execution."


def _read_package_execution_snapshots(pkg: Path) -> list[dict]:
    out: list[dict] = []
    result_dir = pkg / "05-execution" / "04-execution-results"
    if not result_dir.is_dir():
        return out
    for path in sorted(result_dir.glob("execution-results-*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            out.append({
                "path": str(path).replace("\\", "/"),
                "error": str(exc),
            })
            continue
        out.append({
            "path": str(path).replace("\\", "/"),
            "relative_path": str(path.relative_to(_workspace_root)).replace("\\", "/"),
            "scope": data.get("scope") or path.stem.removeprefix("execution-results-"),
            "dashboard_run_id": data.get("dashboard_run_id"),
            "exported_at": data.get("exported_at"),
            "summary": data.get("summary") or {},
            "manual_scope": data.get("manual_scope"),
            "artifact_root": data.get("artifact_root"),
        })
    out.sort(key=lambda row: str(row.get("exported_at") or row.get("path") or ""), reverse=True)
    return out


def _scope_kind(stats: dict, case_count: int, run_total: int | None) -> str:
    if case_count <= 0:
        return "unknown"
    matched = stats.get("matched", 0)
    if matched < case_count:
        return "partial"
    if run_total is not None and run_total > case_count:
        return "package_superset"
    return "package_complete"


def _scope_status(stats: dict, case_count: int, run_total: int | None) -> str:
    if case_count <= 0 or stats.get("matched", 0) <= 0:
        return "no-match"
    if stats.get("matched", 0) < case_count:
        return "partial"
    if stats.get("failed", 0) or stats.get("errored", 0) or stats.get("skipped", 0):
        return "not-passed"
    if stats.get("passed", 0) >= case_count:
        return "passed"
    return "unknown"


@router.get("/api/quality-system/packages/{package_id:path}/executions")
def get_package_executions(package_id: str):
    """Pull dashboard.db test_scenarios whose name or legacy tag matches any case ID
    from the package's xlsx files. Returns the most recent runs that
    exercised this package's scope, with per-scope pass/fail counts."""
    pkg = _resolve_safe(package_id)
    if pkg is None or not pkg.is_dir():
        return {"error": "package not found"}
    package_project_key = None
    first_part = package_id.replace("\\", "/").split("/", 1)[0]
    if first_part:
        try:
            package_project_key = _project_key(first_part)
        except HTTPException:
            package_project_key = None

    case_ids: set[str] = set()
    test_design = pkg / "03-test-design"
    if test_design.is_dir():
        for x in test_design.rglob("test-cases-*.xlsx"):
            case_ids.update(_extract_case_ids(x))

    scope = _execution_scope_from_package(pkg)
    snapshots = _read_package_execution_snapshots(pkg)
    current_snapshot = snapshots[0] if snapshots else None

    if not case_ids:
        return {
            "package_id": package_id,
            "scope": scope,
            "case_ids": [],
            "case_count": 0,
            "snapshots": snapshots,
            "current_snapshot": current_snapshot,
            "runs": [],
            "note": "no case IDs extracted from xlsx",
        }

    from collections import defaultdict
    from sqlalchemy import or_

    db = get_db()
    try:
        # Build OR conditions: scenario tag OR name contains any of the IDs.
        # Use case-insensitive LIKE for robustness.
        conditions = []
        for cid in case_ids:
            conditions.append(TestScenario.tags.like(f"%{cid}%"))
            conditions.append(TestScenario.name.like(f"%{cid}%"))
        scenario_query = (
            db.query(TestScenario)
            .join(TestRun, TestScenario.run_id == TestRun.id)
            .filter(TestRun.run_kind == "full", or_(*conditions))
        )
        if package_project_key:
            scenario_query = scenario_query.filter(TestRun.project_key == package_project_key)
        scenarios = scenario_query.all()

        # Group by run_id, count per-scope stats.
        per_run: dict[int, dict] = defaultdict(
            lambda: {"matched": 0, "passed": 0, "failed": 0, "skipped": 0, "errored": 0}
        )
        per_run_cases: dict[int, set[str]] = defaultdict(set)
        for s in scenarios:
            cid = _case_id_from_scenario_fields(s.name, s.tags)
            if cid:
                per_run_cases[s.run_id].add(cid)
            stats = per_run[s.run_id]
            stats["matched"] += 1
            key = s.status if s.status in ("passed", "failed", "skipped", "errored") else "errored"
            stats[key] += 1

        run_ids = list(per_run.keys())
        if not run_ids:
            return {
                "package_id": package_id,
                "scope": scope,
                "case_ids": sorted(case_ids),
                "case_count": len(case_ids),
                "snapshots": snapshots,
                "current_snapshot": current_snapshot,
                "runs": [],
                "note": "no matching runs in dashboard.db",
            }

        run_query = (
            db.query(TestRun)
            .filter(TestRun.id.in_(run_ids), TestRun.run_kind == "full")
        )
        if package_project_key:
            run_query = run_query.filter(TestRun.project_key == package_project_key)
        runs = run_query.order_by(TestRun.started_at.desc()).limit(20).all()

        result_runs = []
        for r in runs:
            stats = per_run[r.id]
            matched_case_ids = sorted(per_run_cases.get(r.id, set()))
            missing_case_ids = sorted(case_ids - set(matched_case_ids))
            scope_kind = _scope_kind(stats, len(case_ids), r.total)
            scope_status = _scope_status(stats, len(case_ids), r.total)
            result_runs.append({
                "id": r.id,
                "started_at": r.started_at.isoformat() if r.started_at else None,
                "finished_at": r.finished_at.isoformat() if r.finished_at else None,
                "status": r.status,
                "env": r.env,
                "tags": r.tags,
                "run_kind": r.run_kind or "full",
                "scope_matched": stats["matched"],
                "scope_passed": stats["passed"],
                "scope_failed": stats["failed"],
                "scope_skipped": stats["skipped"],
                "scope_errored": stats["errored"],
                "scope_kind": scope_kind,
                "scope_status": scope_status,
                "scope_coverage": round((stats["matched"] / len(case_ids)) * 100, 1) if case_ids else 0,
                "matched_case_ids": matched_case_ids,
                "missing_case_ids": missing_case_ids,
                "package_complete": scope_kind == "package_complete" and scope_status == "passed",
                "promotable": scope_kind == "package_complete" and scope_status == "passed",
                "total_in_run": r.total,
            })
        latest_live_run = result_runs[0] if result_runs else None
        latest_package_complete_run = next(
            (r for r in result_runs if r.get("package_complete")),
            None,
        )
        snapshot_run_id = current_snapshot.get("dashboard_run_id") if current_snapshot else None
        snapshot_run = next(
            (r for r in result_runs if r.get("id") == snapshot_run_id),
            None,
        )
        snapshot_status = "missing"
        if current_snapshot:
            if not snapshot_run:
                snapshot_status = "not-linked"
            elif not snapshot_run.get("package_complete"):
                snapshot_status = "not-package-complete"
            elif latest_package_complete_run and latest_package_complete_run.get("id") != snapshot_run_id:
                snapshot_status = "stale"
            else:
                snapshot_status = "current"

        execution_gate = {
            "status": (
                "ready"
                if snapshot_status in {"current", "stale"} else
                "promote-available"
                if latest_package_complete_run else
                "not-ready"
            ),
            "snapshot_status": snapshot_status,
            "current_snapshot_run_id": snapshot_run_id,
            "latest_package_complete_run_id": (
                latest_package_complete_run.get("id")
                if latest_package_complete_run else None
            ),
            "latest_live_run_id": latest_live_run.get("id") if latest_live_run else None,
        }
        return {
            "package_id": package_id,
            "scope": scope,
            "case_ids": sorted(case_ids),
            "case_count": len(case_ids),
            "snapshots": snapshots,
            "current_snapshot": current_snapshot,
            "current_snapshot_run": snapshot_run,
            "latest_live_run": latest_live_run,
            "latest_package_complete_run": latest_package_complete_run,
            "execution_gate": execution_gate,
            "runs": result_runs,
        }
    finally:
        db.close()


@router.post("/api/quality-system/packages/{package_id:path}/executions/{run_id}/promote")
def promote_package_execution(package_id: str, run_id: int):
    """Promote a package-complete passed dashboard run into 05-execution.

    This keeps Test Run as the live execution surface while Package Health owns
    the audit handoff. Partial runs stay visible but cannot become the package
    execution snapshot used by 06-execution-review.
    """
    pkg = _resolve_safe(package_id)
    if pkg is None or not pkg.is_dir():
        return {"ok": False, "error": "package not found"}

    linked = get_package_executions(package_id)
    if linked.get("error"):
        return {"ok": False, "error": linked["error"]}

    run_info = next((r for r in linked.get("runs", []) if r.get("id") == run_id), None)
    if not run_info:
        return {"ok": False, "error": f"run #{run_id} is not linked to this package"}
    if not run_info.get("promotable"):
        return {
            "ok": False,
            "error": (
                f"run #{run_id} is {run_info.get('scope_matched', 0)}/"
                f"{linked.get('case_count', 0)} for this package and cannot be promoted"
            ),
            "run": run_info,
        }

    first_part = package_id.replace("\\", "/").split("/", 1)[0]
    try:
        project_key = _project_key(first_part)
    except HTTPException:
        project_key = first_part or "west-kowloon"

    scope = linked.get("scope") or _execution_scope_from_package(pkg)
    snapshot = linked.get("current_snapshot")
    artifact_scope = _artifact_scope_from_snapshot(snapshot, scope)
    manual_scope = _manual_scope_from_snapshot(snapshot)
    export_tool = _repo_root / "01-system" / "03-tools" / "export_execution_results.py"
    if not export_tool.is_file():
        return {"ok": False, "error": f"export tool not found: {export_tool}"}

    project_root = _project_workspace_root(project_key)
    cmd = [
        sys.executable,
        str(export_tool),
        "--run-id", str(run_id),
        "--package", str(pkg),
        "--scope", scope,
        "--project", project_key,
        "--project-root", str(project_root),
        "--artifact-scope", artifact_scope,
        "--manual-scope", manual_scope,
    ]
    try:
        result = _subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=120,
            encoding="utf-8",
            errors="replace",
            env=_CHILD_ENV,
        )
    except _subprocess.TimeoutExpired:
        return {"ok": False, "error": "execution snapshot export timed out"}
    if result.returncode != 0:
        return {
            "ok": False,
            "error": "execution snapshot export failed",
            "stdout": result.stdout[-1000:],
            "stderr": result.stderr[-1000:],
        }

    out_dir = pkg / "05-execution" / "04-execution-results"
    md_path = out_dir / f"execution-results-{scope}.md"
    json_path = out_dir / f"execution-results-{scope}.json"
    docx_path = md_path.with_suffix(".docx")
    docx_ok = False
    docx_error = None
    md_docx_tool = _repo_root / "01-system" / "03-tools" / "md_docx.py"
    if md_path.is_file() and md_docx_tool.is_file():
        try:
            docx_result = _subprocess.run(
                [sys.executable, str(md_docx_tool), "to-docx", str(md_path)],
                capture_output=True,
                text=True,
                timeout=60,
                encoding="utf-8",
                errors="replace",
                env=_CHILD_ENV,
            )
            docx_ok = docx_result.returncode == 0 and docx_path.is_file()
            if not docx_ok:
                docx_error = (docx_result.stdout + docx_result.stderr).strip()[-1000:]
        except _subprocess.TimeoutExpired:
            docx_error = "docx generation timed out"

    return {
        "ok": True,
        "package_id": package_id,
        "scope": scope,
        "run_id": run_id,
        "md_path": str(md_path.relative_to(_workspace_root)).replace("\\", "/") if md_path.exists() else None,
        "json_path": str(json_path.relative_to(_workspace_root)).replace("\\", "/") if json_path.exists() else None,
        "docx_path": str(docx_path.relative_to(_workspace_root)).replace("\\", "/") if docx_path.exists() else None,
        "docx_ok": docx_ok,
        "docx_error": docx_error,
        "stdout": result.stdout[-1000:],
    }


@router.post("/api/quality-system/packages/{package_id:path}/validate-xlsx")
def validate_package_xlsx(package_id: str):
    """Run validate_testcase_xlsx.py against every test-cases-*.xlsx in
    this package's 03-test-design tree. Returns one result per xlsx."""
    pkg = _resolve_safe(package_id)
    if pkg is None or not pkg.is_dir():
        return {"ok": False, "error": "package not found"}
    test_design = pkg / "03-test-design"
    xlsx_files = sorted(test_design.rglob("test-cases-*.xlsx")) if test_design.is_dir() else []
    if not xlsx_files:
        return {"ok": True, "package_id": package_id, "results": [], "summary": "no xlsx to validate"}
    tool = _repo_root / "01-system" / "03-tools" / "validate_testcase_xlsx.py"
    results = []
    for x in xlsx_files:
        try:
            r = _subprocess.run(
                [sys.executable, str(tool), str(x)],
                capture_output=True, text=True, timeout=30,
                encoding="utf-8", errors="replace", env=_CHILD_ENV,
            )
            results.append({
                "file": str(x.relative_to(_workspace_root)).replace("\\", "/"),
                "ok": r.returncode == 0,
                "exit_code": r.returncode,
                "output": (r.stdout + r.stderr).strip()[-1200:],
            })
        except _subprocess.TimeoutExpired:
            results.append({
                "file": str(x.relative_to(_workspace_root)).replace("\\", "/"),
                "ok": False,
                "exit_code": -1,
                "output": "timeout",
            })
    passed = sum(1 for r in results if r["ok"])
    return {
        "ok": True,
        "package_id": package_id,
        "summary": f"{passed}/{len(results)} xlsx passed",
        "results": results,
    }


# Catch-all package detail route MUST come last in the quality-system block —
# {package_id:path} would otherwise swallow `/executions` and `/validate-xlsx`
# suffixes (FastAPI matches GET routes in source order).
@router.get("/api/quality-system/packages/{package_id:path}")
def get_quality_package(package_id: str):
    """Return full detail for one package. package_id is the workspace-
    relative path with forward slashes, e.g.
    `west-kowloon/01-requirements/02-subprojects/02-website/02-modules/08-discount/2026-06-25`.
    """
    if package_scanner is None:
        return {"error": "package_scanner not available"}
    target = (_workspace_root / package_id).resolve()
    try:
        target.relative_to(_workspace_root)
    except ValueError:
        return {"error": "package_id outside workspace"}
    if not target.is_dir():
        return {"error": "package not found", "package_id": package_id}
    detail = package_scanner.scan_package(target)
    return package_scanner.to_jsonable(detail)

