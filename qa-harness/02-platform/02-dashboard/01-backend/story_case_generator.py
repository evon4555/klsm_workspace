"""Generate West Kowloon story-scoped test-case workbooks.

The generator intentionally reuses the validated login/registration workbook as
the style and evidence source. It copies row values, formatting, row heights,
and images anchored to selected test-case rows, then adds story-specific rows
for requirements that were not present in the source workbook.
"""

from __future__ import annotations

import copy
import os
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter


WORKSPACE_ROOT = Path(os.environ.get("QA_WORKSPACE_ROOT", r"D:\Workspace")).resolve()
WESTK_ROOT = Path(os.environ.get("QA_WESTK_ROOT", WORKSPACE_ROOT / "west-kowloon")).resolve()

LOGIN_MODULE_ROOT = (
    WESTK_ROOT
    / "01-requirements"
    / "02-subprojects"
    / "02-website"
    / "02-modules"
    / "01-login-registration"
)
LOGIN_REFERENCE_ROOT = LOGIN_MODULE_ROOT / "2026-06-16"
CURRENT_STORY_ROOT = LOGIN_MODULE_ROOT / "2026-06-23"

SOURCE_WORKBOOK = (
    LOGIN_REFERENCE_ROOT
    / "03-test-design"
    / "test-cases-registration-login_2026-06-12.scope-clean-2026-06-16.env-fixed-2026-06-18.xlsx"
)

OUTPUT_DIR = CURRENT_STORY_ROOT / "03-test-design" / "story-splits"

TEST_SHEET = "Test Cases"
TODAY = datetime.now().strftime("%Y-%m-%d")

SCREENSHOT_COLUMN_WIDTH = 55
SCREENSHOT_TARGET_WIDTH_PX = int(SCREENSHOT_COLUMN_WIDTH * 7 + 5) - 4
EMU_PER_PX = 9525
MIN_SCREENSHOT_ROW_HEIGHT_PT = 50
SCREENSHOT_BOTTOM_PADDING_PX = 4


@dataclass(frozen=True)
class StoryRule:
    story_id: int
    slug: str
    title: str
    source: str
    case_ids: tuple[str, ...]
    extra_rows: tuple[dict[str, str], ...] = ()


STORY_4161_CASES = (
    "SIT-TC-WEB-AUTH-021",
    "SIT-TC-WEB-AUTH-022",
    "SIT-TC-WEB-AUTH-023",
    "SIT-TC-WEB-AUTH-024",
    "SIT-TC-WEB-AUTH-025",
    "SIT-TC-WEB-AUTH-038",
    "SIT-TC-WEB-AUTH-056",
    "SIT-TC-WEB-AUTH-068",
    "SIT-TC-WEB-AUTH-069",
)

STORY_4266_CASES = (
    "SIT-TC-WEB-AUTH-100",
    "SIT-TC-WEB-AUTH-101",
    "SIT-TC-WEB-AUTH-102",
    "SIT-TC-WEB-AUTH-103",
    "SIT-TC-WEB-AUTH-104",
    "SIT-TC-WEB-AUTH-105",
    "SIT-TC-WEB-AUTH-106",
    "SIT-TC-WEB-AUTH-107",
    "SIT-TC-WEB-AUTH-108",
    "SIT-TC-WEB-AUTH-109",
    "SIT-TC-WEB-AUTH-113",
    "SIT-TC-WEB-AUTH-114",
    "SIT-TC-WEB-AUTH-115",
    "SIT-TC-WEB-AUTH-116",
    "SIT-TC-WEB-AUTH-117",
    "SIT-TC-WEB-AUTH-123",
    "SIT-TC-WEB-AUTH-124",
    "SIT-TC-WEB-AUTH-125",
)


STORY_RULES: dict[int, StoryRule] = {
    4161: StoryRule(
        story_id=4161,
        slug="story-4161-guest-mode-optimization",
        title="游客模式几个问题优化",
        source=(
            "ZenTao STORY-4161; 国际版标准官网 - 在线购票.md "
            "§游客模式 / 游客模式的几个问题"
        ),
        case_ids=STORY_4161_CASES,
        extra_rows=(
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-129",
                "Module/Feature": "Website / Guest Purchase",
                "Priority": "High",
                "Severity": "High",
                "Collected from": "ZenTao STORY-4161; 在线购票 §游客模式的几个问题 item 1",
                "Test Scenario": "Guest purchase limit uses single-order validation only",
                "Test Case Description": (
                    "Verify guest purchase limitation follows the clarified rule: "
                    "guest mode performs only single-order limit validation; other "
                    "registered-user purchase-limit checks are not applied."
                ),
                "Preconditions": "Guest checkout is enabled; a product with purchase-limit rules is available.",
                "Test Steps": (
                    "1. Continue as guest and enter checkout.\n"
                    "2. Fill contact email information on the order.\n"
                    "3. Submit quantities that should trigger single-order limit validation.\n"
                    "4. Compare against registered-user-only limit rules."
                ),
                "Test Data": "Guest contact email; product with configured purchase limits.",
                "Expected Result": (
                    "1. Single-order purchase limit is enforced for the guest order.\n"
                    "2. Other registered-user limit checks are not applied to the guest flow.\n"
                    "3. Error message, if triggered, is clear and localized."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
            },
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-130",
                "Module/Feature": "Website / Guest Session",
                "Priority": "High",
                "Severity": "High",
                "Collected from": "ZenTao STORY-4161; 在线购票 §游客模式的几个问题 item 2",
                "Test Scenario": "Guest checkout countdown uses the smaller value between guest validity and order countdown",
                "Test Case Description": (
                    "Verify frontend countdown behavior when guest-session validity and "
                    "order countdown are different."
                ),
                "Preconditions": "Guest session active; order countdown is available.",
                "Test Steps": (
                    "1. Login as guest.\n"
                    "2. Start an order with its own countdown.\n"
                    "3. Observe the displayed remaining time when guest validity and order countdown differ."
                ),
                "Test Data": "Guest session; order countdown shorter or longer than guest-session validity.",
                "Expected Result": (
                    "1. Displayed countdown uses the smaller remaining time.\n"
                    "2. When the chosen countdown expires, guest checkout can no longer continue.\n"
                    "3. Cart/order draft is handled according to guest timeout rules."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
            },
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-131",
                "Module/Feature": "Website / Registered Session",
                "Priority": "Medium",
                "Severity": "Medium",
                "Collected from": "ZenTao STORY-4161; 在线购票 §游客模式的几个问题 item 3",
                "Test Scenario": "Registered-user session validity follows 30-day configuration",
                "Test Case Description": (
                    "Verify non-guest login mode follows the configured registered-user "
                    "validity period, currently 30 days per requirement note."
                ),
                "Preconditions": "Registered user can login successfully.",
                "Test Steps": (
                    "1. Login as registered user.\n"
                    "2. Inspect session expiry configuration or cookie/session metadata.\n"
                    "3. Reopen the site within and beyond the configured validity window."
                ),
                "Test Data": "Registered test account.",
                "Expected Result": (
                    "1. Registered-user session validity is configured as 30 days.\n"
                    "2. Session remains valid within the configured period.\n"
                    "3. Expired session requires re-login."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
            },
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-132",
                "Module/Feature": "Website / Guest Verification",
                "Priority": "High",
                "Severity": "Medium",
                "Collected from": "ZenTao STORY-4161; 在线购票 §游客模式的几个问题 item 4",
                "Test Scenario": "Guest verification method aligns with regular login and registration",
                "Test Case Description": (
                    "Verify the guest-mode verification method uses the same approach as "
                    "regular login/registration per the clarified requirement."
                ),
                "Preconditions": "Guest login, regular login, and registration are available.",
                "Test Steps": (
                    "1. Open guest login verification.\n"
                    "2. Open regular login/registration verification.\n"
                    "3. Compare verification type, failure handling, and retry behavior."
                ),
                "Test Data": "Guest browser and normal login/registration pages.",
                "Expected Result": (
                    "1. Guest verification method is consistent with regular login/registration.\n"
                    "2. Failure and retry handling are consistent unless a documented exception exists."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
            },
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-133",
                "Module/Feature": "Website / Guest Ticket Access",
                "Priority": "High",
                "Severity": "High",
                "Collected from": "ZenTao STORY-4161; 在线购票 §游客模式的几个问题 item 5",
                "Test Scenario": "Guest can obtain admission code after successful purchase",
                "Test Case Description": (
                    "Verify how a guest obtains admission code information after purchase: "
                    "order detail, ticket-code detail, or cart-order handling."
                ),
                "Preconditions": "Guest purchase can be completed successfully.",
                "Test Steps": (
                    "1. Complete a guest purchase.\n"
                    "2. Open the success page and order detail route available to the guest.\n"
                    "3. Verify admission code or ticket credential access.\n"
                    "4. Repeat for cart-order flow if applicable."
                ),
                "Test Data": "Guest order with payable ticket.",
                "Expected Result": (
                    "1. Guest can access the admission code or ticket credential through the documented route.\n"
                    "2. The route matches the implementation decision for order detail / ticket-code detail / cart order.\n"
                    "3. No registered-only page is required unless the user was converted."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
                "Comments/Remarks": "Confirm final route with product/dev if still pending.",
            },
        ),
    ),
    4266: StoryRule(
        story_id=4266,
        slug="story-4266-guest-purchase-registration",
        title="游客购票功能升级，引导注册功能",
        source=(
            "ZenTao STORY-4266; 国际版标准官网 V1.2 - 游客购票引导注册功能.md"
        ),
        case_ids=STORY_4266_CASES,
        extra_rows=(
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-126",
                "Module/Feature": "Guest-to-Member",
                "Priority": "Medium",
                "Severity": "Medium",
                "Collected from": "ZenTao STORY-4266; V1.2 §礼包发放后台配置",
                "Test Scenario": "Reward configuration supports guest-purchase conversion registration scene",
                "Test Case Description": (
                    "Verify the back-office reward configuration adds a scenario for "
                    "guest-purchase converted registered users."
                ),
                "Preconditions": "Back-office reward configuration is available.",
                "Test Steps": (
                    "1. Open reward / benefit configuration in back office.\n"
                    "2. Locate reward scenario options.\n"
                    "3. Configure coupon, points, or growth-value reward for guest purchase conversion."
                ),
                "Test Data": "Reward configuration with coupon / points / growth value.",
                "Expected Result": (
                    "1. Guest-purchase conversion registration is available as a reward scene.\n"
                    "2. Coupon, points, and growth value can be configured where supported.\n"
                    "3. Configuration can be saved and later used by the conversion flow."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
                "Comments/Remarks": "Section is marked 待讨论 in V1.2; execute after product decision.",
            },
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-127",
                "Module/Feature": "Guest-to-Member",
                "Priority": "High",
                "Severity": "Medium",
                "Collected from": "ZenTao STORY-4266; V1.2 §提交订单 / 自动发放福利",
                "Test Scenario": "Converted guest receives configured member benefit after payment success",
                "Test Case Description": (
                    "Verify that after opt-in conversion and successful payment, the "
                    "converted account receives the configured exclusive benefit."
                ),
                "Preconditions": "Guest conversion reward configured; guest checkout available.",
                "Test Steps": (
                    "1. Configure a reward for guest-purchase conversion.\n"
                    "2. Checkout as guest with an unregistered contact and opt in.\n"
                    "3. Complete payment.\n"
                    "4. Inspect the converted account's coupon / points / growth-value records."
                ),
                "Test Data": "Unregistered email or mobile; configured reward.",
                "Expected Result": (
                    "1. Guest account is converted to registered account.\n"
                    "2. Configured benefit is issued exactly once to the converted account.\n"
                    "3. Benefit issue record is traceable to the conversion scene."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
            },
            {
                "Label": "SIT/UAT",
                "Test Case ID": "SIT-TC-WEB-AUTH-128",
                "Module/Feature": "Guest-to-Member",
                "Priority": "High",
                "Severity": "High",
                "Collected from": "ZenTao STORY-4266; V1.2 §其他补充说明 item 2/4",
                "Test Scenario": "Payment callback after conversion is handled as registered-user state",
                "Test Case Description": (
                    "Verify the critical payment timing rule: account creation happens before "
                    "payment, so callback and post-payment flow use registered-user state even "
                    "if the original guest validity would otherwise expire."
                ),
                "Preconditions": "Guest conversion opt-in; controllable payment callback or delayed payment test path.",
                "Test Steps": (
                    "1. Start guest checkout with an unregistered contact.\n"
                    "2. Tick the conversion opt-in checkbox and submit order.\n"
                    "3. Confirm account is created before payment completion.\n"
                    "4. Complete or simulate payment callback after a delay.\n"
                    "5. Query order and user state after callback."
                ),
                "Test Data": "Delayed payment or callback test order.",
                "Expected Result": (
                    "1. Account has already been converted before payment callback.\n"
                    "2. Callback and post-payment logic use registered-user state.\n"
                    "3. Order remains linked to the converted account.\n"
                    "4. User info is refreshed through the upgraded login state."
                ),
                "Test Case Owner": "Antank QA Team",
                "Environment": "SIT",
                "Status": "Not Run",
            },
        ),
    ),
}


STORY_4266_SOURCE = (
    "ZenTao STORY-4266; 国际版标准官网 V1.2 - 游客购票引导注册功能.md"
)

STORY_4266_OVERRIDES: dict[str, dict[str, str]] = {
    "SIT-TC-WEB-AUTH-105": {
        "Collected from": f"{STORY_4266_SOURCE} §账号校验",
        "Expected Result": (
            "1. Checkbox label per language:\n"
            "   - zh-CN: 同意同步注册为用户\n"
            "   - zh-HK: 同意同步註冊為用戶\n"
            "   - en: I agree to register as a registered user.\n"
            "2. Description text appears below per V1.2 in zh-CN / zh-HK / en.\n"
            "3. Checkbox is unchecked by default and is not required for order submission."
        ),
    },
    "SIT-TC-WEB-AUTH-109": {
        "Collected from": f"{STORY_4266_SOURCE} §支付成功页",
        "Test Case Description": (
            "Verify the payment-success page when conversion happened via email: "
            "the 3-language registration-success message renders with the selected "
            "email placeholder filled in, and the Back to Home button routes to the homepage."
        ),
        "Expected Result": (
            "1. Registration-success message appears in zh-CN / zh-HK / en per V1.2.\n"
            "2. The email placeholder is replaced by the selected email.\n"
            "3. Button label per language: 返回首页 / 返回首頁 / Back to Home.\n"
            "4. Clicking the button returns to the homepage."
        ),
        "Comments/Remarks": "Updated from latest V1.2 source package on 2026-06-23.",
    },
    "SIT-TC-WEB-AUTH-114": {
        "Collected from": f"{STORY_4266_SOURCE} §账号校验",
        "Expected Result": (
            "1. Same 3-language checkbox and description as AUTH-105 are shown.\n"
            "2. Trigger is unregistered mobile.\n"
            "3. Checkbox is unchecked by default and is not required for order submission."
        ),
    },
    "SIT-TC-WEB-AUTH-117": {
        "Collected from": f"{STORY_4266_SOURCE} §支付成功页",
        "Test Case Description": (
            "Verify the payment-success page when conversion happened via mobile: "
            "the 3-language registration-success message renders with the selected "
            "mobile placeholder filled in, and the Back to Home button routes to the homepage."
        ),
        "Expected Result": (
            "1. Same 3-language registration-success message as AUTH-109 is shown.\n"
            "2. The mobile placeholder is replaced by the selected mobile number.\n"
            "3. Button label per language: 返回首页 / 返回首頁 / Back to Home.\n"
            "4. Clicking the button returns to the homepage."
        ),
        "Comments/Remarks": "Updated from latest V1.2 source package on 2026-06-23.",
    },
}


SPLIT_CASE_IDS = set(STORY_4161_CASES) | set(STORY_4266_CASES)


def _copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy.copy(src.protection)
    if src.hyperlink:
        dst._hyperlink = copy.copy(src.hyperlink)
    if src.comment:
        dst.comment = copy.copy(src.comment)


def _copy_sheet_dimensions(src_ws, dst_ws) -> None:
    for key, dim in src_ws.column_dimensions.items():
        dst = dst_ws.column_dimensions[key]
        dst.width = dim.width
        dst.hidden = dim.hidden
        dst.bestFit = dim.bestFit
    dst_ws.freeze_panes = src_ws.freeze_panes
    dst_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines


def _image_bytes(img) -> bytes:
    data = img._data()
    if isinstance(data, bytes):
        return data
    return bytes(data)


def _copy_images(src_ws, dst_ws, row_map: dict[int, int] | None = None) -> int:
    copied = 0
    for img in getattr(src_ws, "_images", []):
        try:
            source_row = img.anchor._from.row + 1
            source_col = img.anchor._from.col + 1
        except Exception:
            continue
        target_row = row_map.get(source_row) if row_map is not None else source_row
        if not target_row:
            continue
        new_img = XLImage(BytesIO(_image_bytes(img)))
        new_img.width = img.width
        new_img.height = img.height
        dst_ws.add_image(new_img, f"{get_column_letter(source_col)}{target_row}")
        copied += 1
    return copied


def _normalize_screenshot_images(ws) -> dict[str, int]:
    headers = _headers(ws)
    if not headers.get("Screenshots"):
        return {"dropped": 0, "resized": 0, "rowsAdjusted": 0}

    # Do not resize, collapse, or otherwise alter embedded screenshots here.
    # Users may have manually adjusted image sizes in the source workbook, and a
    # story split must preserve that display size exactly unless the user asks
    # for a resize. Keep this helper limited to workbook-range housekeeping.
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    return {"dropped": 0, "resized": 0, "rowsAdjusted": 0}


def _case_id_for_row(ws, row: int) -> str:
    for col in (2, 1):
        value = ws.cell(row, col).value
        if value and str(value).startswith("SIT-TC-"):
            return str(value).strip()
    return ""


def _copy_rows(
    src_ws,
    dst_ws,
    source_rows: Iterable[int],
    overrides: dict[str, dict[str, str]] | None = None,
) -> dict[int, int]:
    overrides = overrides or {}
    row_map: dict[int, int] = {}

    for col in range(1, src_ws.max_column + 1):
        _copy_cell(src_ws.cell(1, col), dst_ws.cell(1, col))
    dst_ws.row_dimensions[1].height = src_ws.row_dimensions[1].height

    headers = {
        str(dst_ws.cell(1, col).value).split("\n")[0].strip(): col
        for col in range(1, dst_ws.max_column + 1)
        if dst_ws.cell(1, col).value
    }

    target_row = 2
    for source_row in source_rows:
        row_map[source_row] = target_row
        for col in range(1, src_ws.max_column + 1):
            _copy_cell(src_ws.cell(source_row, col), dst_ws.cell(target_row, col))
        dst_ws.row_dimensions[target_row].height = src_ws.row_dimensions[source_row].height

        cid = _case_id_for_row(src_ws, source_row)
        for header, value in overrides.get(cid, {}).items():
            col = headers.get(header)
            if col:
                dst_ws.cell(target_row, col).value = value
        target_row += 1

    return row_map


def _append_extra_rows(dst_ws, extra_rows: Iterable[dict[str, str]]) -> int:
    headers = {
        str(dst_ws.cell(1, col).value).split("\n")[0].strip(): col
        for col in range(1, dst_ws.max_column + 1)
        if dst_ws.cell(1, col).value
    }
    if dst_ws.max_row >= 2:
        style_row = 2
    else:
        style_row = 1
    count = 0
    for data in extra_rows:
        row = dst_ws.max_row + 1
        for col in range(1, dst_ws.max_column + 1):
            _copy_cell(dst_ws.cell(style_row, col), dst_ws.cell(row, col))
            dst_ws.cell(row, col).value = None
        for header, value in data.items():
            col = headers.get(header)
            if col:
                dst_ws.cell(row, col).value = value
        dst_ws.row_dimensions[row].height = dst_ws.row_dimensions[style_row].height
        count += 1
    return count


def _headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, col).value).split("\n")[0].strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(1, col).value
    }


def _cell_text(ws, row: int, col: int | None) -> str:
    if not col:
        return ""
    value = ws.cell(row, col).value
    return "" if value is None else str(value).strip()


def _normalize_execution_metadata(ws) -> dict[str, int]:
    """Keep copied execution data auditable without overwriting source values."""
    headers = _headers(ws)
    status_col = headers.get("Status")
    actual_col = headers.get("Actual Result")
    comments_col = headers.get("Comments/Remarks")
    if not status_col:
        return {"actualFilled": 0, "commentsFilled": 0}

    actual_filled = 0
    comments_filled = 0
    for row in range(2, ws.max_row + 1):
        status = _cell_text(ws, row, status_col)
        if status not in {"Pass", "Fail"}:
            continue

        comments = _cell_text(ws, row, comments_col)
        if actual_col and not _cell_text(ws, row, actual_col):
            if status == "Fail" and comments:
                ws.cell(row, actual_col).value = f"Failure noted in source comments: {comments}"
            elif status == "Pass":
                ws.cell(row, actual_col).value = (
                    "Passed in source workbook; source actual detail was not provided."
                )
            else:
                ws.cell(row, actual_col).value = (
                    "Failure recorded in source workbook; source actual detail was not provided."
                )
            actual_filled += 1

        if comments_col and not _cell_text(ws, row, comments_col):
            ws.cell(row, comments_col).value = (
                f"Preserved from source workbook split on {TODAY}; "
                "execution screenshots retained where present."
            )
            comments_filled += 1

    return {"actualFilled": actual_filled, "commentsFilled": comments_filled}


def _clone_entire_sheet(src_ws, dst_ws) -> int:
    row_map = {row: row for row in range(1, src_ws.max_row + 1)}
    _copy_sheet_dimensions(src_ws, dst_ws)
    for row in range(1, src_ws.max_row + 1):
        dst_ws.row_dimensions[row].height = src_ws.row_dimensions[row].height
        for col in range(1, src_ws.max_column + 1):
            _copy_cell(src_ws.cell(row, col), dst_ws.cell(row, col))
    copied = _copy_images(src_ws, dst_ws, row_map)
    _normalize_screenshot_images(dst_ws)
    return copied


def _new_workbook_from_source(src_wb, include_support_sheets: bool = False):
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    src_ws = src_wb[TEST_SHEET]
    ws = wb.create_sheet(TEST_SHEET)
    _copy_sheet_dimensions(src_ws, ws)

    if include_support_sheets:
        for sheet_name in src_wb.sheetnames:
            if sheet_name == TEST_SHEET:
                continue
            support = wb.create_sheet(sheet_name)
            _clone_entire_sheet(src_wb[sheet_name], support)

    return wb, ws


def _story_output_path(rule: StoryRule, output_dir: Path = OUTPUT_DIR) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir / f"{rule.slug}-test-cases_{TODAY}.xlsx"


def generate_story_workbook(
    story_id: int,
    output_dir: Path = OUTPUT_DIR,
    source_workbook: Path = SOURCE_WORKBOOK,
) -> dict:
    rule = STORY_RULES.get(int(story_id))
    if not rule:
        raise ValueError(f"No local story testcase rule configured for story {story_id}")
    if not source_workbook.exists():
        raise FileNotFoundError(f"Source workbook not found: {source_workbook}")

    src_wb = load_workbook(source_workbook)
    src_ws = src_wb[TEST_SHEET]
    selected_rows = [
        row for row in range(2, src_ws.max_row + 1)
        if _case_id_for_row(src_ws, row) in rule.case_ids
    ]

    wb, ws = _new_workbook_from_source(src_wb, include_support_sheets=False)
    overrides = STORY_4266_OVERRIDES if story_id == 4266 else {}
    row_map = _copy_rows(src_ws, ws, selected_rows, overrides=overrides)
    copied_images = _copy_images(src_ws, ws, row_map)
    _normalize_screenshot_images(ws)
    added = _append_extra_rows(ws, rule.extra_rows)
    metadata_fills = _normalize_execution_metadata(ws)

    out = _story_output_path(rule, output_dir=output_dir)
    wb.save(out)
    case_count = len(selected_rows) + added
    source_case_ids = {_case_id_for_row(src_ws, row) for row in selected_rows}
    return {
        "storyId": rule.story_id,
        "title": rule.title,
        "source": rule.source,
        "path": str(out),
        "filename": out.name,
        "cases": case_count,
        "copiedRows": len(selected_rows),
        "addedRows": added,
        "images": copied_images,
        "missingSourceCaseIds": sorted(set(rule.case_ids) - source_case_ids),
        "executionMetadataFilled": metadata_fills,
    }


def generate_login_registration_core_workbook(
    output_dir: Path = OUTPUT_DIR,
    source_workbook: Path = SOURCE_WORKBOOK,
) -> dict:
    if not source_workbook.exists():
        raise FileNotFoundError(f"Source workbook not found: {source_workbook}")

    src_wb = load_workbook(source_workbook)
    src_ws = src_wb[TEST_SHEET]
    source_case_count = sum(
        1 for row in range(2, src_ws.max_row + 1)
        if _case_id_for_row(src_ws, row)
    )
    selected_rows = [
        row for row in range(2, src_ws.max_row + 1)
        if _case_id_for_row(src_ws, row)
        and _case_id_for_row(src_ws, row) not in SPLIT_CASE_IDS
    ]

    wb, ws = _new_workbook_from_source(src_wb, include_support_sheets=False)
    row_map = _copy_rows(src_ws, ws, selected_rows)
    copied_images = _copy_images(src_ws, ws, row_map)
    _normalize_screenshot_images(ws)
    metadata_fills = _normalize_execution_metadata(ws)

    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / f"test-cases-login-registration-core_{TODAY}.xlsx"
    wb.save(out)
    return {
        "path": str(out),
        "filename": out.name,
        "cases": len(selected_rows),
        "removedCases": source_case_count - len(selected_rows),
        "configuredSplitCases": len(SPLIT_CASE_IDS),
        "images": copied_images,
        "executionMetadataFilled": metadata_fills,
    }


def generated_file_path(filename: str, output_dir: Path = OUTPUT_DIR) -> Path:
    if not re.match(r"^[A-Za-z0-9_.\-]+$", filename):
        raise ValueError("Invalid generated filename")
    path = (output_dir / filename).resolve()
    root = output_dir.resolve()
    if root not in path.parents and path != root:
        raise ValueError("Generated filename escapes output directory")
    return path
